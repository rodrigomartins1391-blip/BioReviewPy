# -*- coding: utf-8 -*-
"""
BioReviewPy v0.23.1

Open-source software for bibliographic data harmonization, duplicate detection,
screening support, audit tracking, and evidence-synthesis workflows.

Developed by: Rodrigo Martins dos Santos
Copyright (c) 2026 Rodrigo Martins dos Santos
License: MIT
Repository: https://github.com/rodrigomartins1391-blip/BioReviewPy
Citation metadata: see CITATION.cff in the project repository.

Fluxo:
1) Detectar automaticamente (ou escolher manualmente) a base e carregar o(s) arquivo(s).
2) Processar.
3) O programa seleciona automaticamente a maior base como referência.
4) Compara a referência com as demais de forma sequencial.
5) Mostra duplicatas confirmadas e possíveis duplicatas lado a lado.
6) Permite confirmar/rejeitar casos duvidosos.
7) Exporta:
   - um Excel LIMPO separado para cada base;
   - um Excel com a comparação das duplicatas;
   - um Excel FINAL para triagem;
   - relatório TXT.

Campos principais:
Autores | Título | Revista | Ano | DOI

Bases:
PubMed, Scopus, Web of Science, Embase, LILACS/BVS,
Cochrane Library, SciELO e Outra.

O parser Embase reconhece também o formato:
RECORD N
TITLE
  ...
AUTHOR NAMES
  ...
SOURCE
  ...
VOLUME
  ...
DATE OF PUBLICATION
  ...
"""

import csv
import html
import json
import hashlib
import math
import os
import sys
import queue
import threading
import re
import shutil
import subprocess
import tempfile
import tkinter as tk
import unicodedata
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime
from collections import defaultdict
from pathlib import Path
from tkinter import ttk, filedialog, messagebox, simpledialog

import pandas as pd
from rapidfuzz import fuzz, process

APP_TITLE = "BioReviewPy"
VERSION = "0.23"
AUTHOR_NAME = "Rodrigo Martins dos Santos"
AUTHOR_EMAIL = "rodrigoms13@hotmail.com"
COPYRIGHT_NOTICE = "Copyright © 2026 Rodrigo Martins dos Santos"
LICENSE_NAME = "MIT"
PROJECT_URL = "https://github.com/rodrigomartins1391-blip/BioReviewPy"

__author__ = AUTHOR_NAME
__version__ = VERSION
__license__ = LICENSE_NAME
__copyright__ = COPYRIGHT_NOTICE


# ============================================================
# INTERFACE MULTILÍNGUE
# ============================================================

LANGUAGE_OPTIONS = {
    "Português": "pt",
    "English": "en",
    "Español": "es",
}

UI_TEXT = {
    "project": {"pt": "Projeto", "en": "Project", "es": "Proyecto"},
    "new_project": {"pt": "Novo projeto", "en": "New project", "es": "Nuevo proyecto"},
    "open_project": {"pt": "Abrir projeto...", "en": "Open project...", "es": "Abrir proyecto..."},
    "save_project": {"pt": "Salvar projeto", "en": "Save project", "es": "Guardar proyecto"},
    "save_project_as": {"pt": "Salvar projeto como...", "en": "Save project as...", "es": "Guardar proyecto como..."},
    "tools": {"pt": "Ferramentas", "en": "Tools", "es": "Herramientas"},
    "help": {"pt": "Ajuda", "en": "Help", "es": "Ayuda"},
    "about": {"pt": "Sobre o BioReviewPy", "en": "About BioReviewPy", "es": "Acerca de BioReviewPy"},
    "dashboard": {"pt": "Painel do projeto", "en": "Project dashboard", "es": "Panel del proyecto"},
    "second_screen": {"pt": "2ª triagem – Marcar revisões sistemáticas/meta-análises", "en": "Second screening – Mark systematic reviews/meta-analyses", "es": "Segundo cribado – Marcar revisiones sistemáticas/metaanálisis"},
    "search_article": {"pt": "Buscar artigo", "en": "Search article", "es": "Buscar artículo"},
    "format_refs": {"pt": "Formatar referências (ABNT / Vancouver / APA)...", "en": "Format references (ABNT / Vancouver / APA)...", "es": "Formatear referencias (ABNT / Vancouver / APA)..."},
    "prisma_auto": {"pt": "PRISMA automático", "en": "Automatic PRISMA", "es": "PRISMA automático"},
    "audit_report": {"pt": "Exportar relatório de auditoria...", "en": "Export audit report...", "es": "Exportar informe de auditoría..."},
    "bibliometrics": {"pt": "Bibliometria", "en": "Bibliometrics", "es": "Bibliometría"},
    "biblio_merge": {"pt": "Mesclar, deduplicar e padronizar bases...", "en": "Merge, deduplicate and standardize databases...", "es": "Combinar, deduplicar y estandarizar bases..."},
    "language": {"pt": "Idioma", "en": "Language", "es": "Idioma"},
    "subtitle": {"pt": "Deduplicação simples, auditável e orientada por bases", "en": "Simple, auditable database-oriented deduplication", "es": "Deduplicación simple, auditable y orientada por bases"},
    "credits": {
        "pt": "Desenvolvido por Rodrigo Martins dos Santos   |   © 2026   |   Licença MIT",
        "en": "Developed by Rodrigo Martins dos Santos   |   © 2026   |   MIT License",
        "es": "Desarrollado por Rodrigo Martins dos Santos   |   © 2026   |   Licencia MIT",
    },
    "load_bases": {"pt": "1. Carregar bases", "en": "1. Load databases", "es": "1. Cargar bases"},
    "database": {"pt": "Base:", "en": "Database:", "es": "Base:"},
    "add_files": {"pt": "Adicionar arquivo(s)", "en": "Add file(s)", "es": "Añadir archivo(s)"},
    "remove_selected": {"pt": "Remover selecionado", "en": "Remove selected", "es": "Eliminar seleccionado"},
    "clear_all": {"pt": "Limpar tudo", "en": "Clear all", "es": "Limpiar todo"},
    "process": {"pt": "PROCESSAR", "en": "PROCESS", "es": "PROCESAR"},
    "auto_detection_note": {
        "pt": "Detecção automática de base e formato. Bases suportadas: PubMed, Web of Science, Scopus, Embase, LILACS/BVS, Cochrane Library e SciELO. Formatos comuns: RIS, BibTeX, NBIB, TXT/CIW, CSV, XML, DOCX/XLSX, HTML e PDF.",
        "en": "Automatic database and format detection. Supported databases: PubMed, Web of Science, Scopus, Embase, LILACS/BVS, Cochrane Library and SciELO. Common formats: RIS, BibTeX, NBIB, TXT/CIW, CSV, XML, DOCX/XLSX, HTML and PDF.",
        "es": "Detección automática de base y formato. Bases compatibles: PubMed, Web of Science, Scopus, Embase, LILACS/BVS, Cochrane Library y SciELO. Formatos comunes: RIS, BibTeX, NBIB, TXT/CIW, CSV, XML, DOCX/XLSX, HTML y PDF."
    },
    "file": {"pt": "Arquivo", "en": "File", "es": "Archivo"},
    "extension": {"pt": "Extensão", "en": "Extension", "es": "Extensión"},
    "detected_format": {"pt": "Formato detectado", "en": "Detected format", "es": "Formato detectado"},
    "comparison": {"pt": "2. Comparação", "en": "2. Comparison", "es": "2. Comparación"},
    "auto_reference": {"pt": "Base principal automática:", "en": "Automatic reference database:", "es": "Base principal automática:"},
    "review_similarity": {"pt": "Similaridade para revisar:", "en": "Similarity for manual review:", "es": "Similitud para revisar:"},
    "update_comparison": {"pt": "ATUALIZAR COMPARAÇÃO", "en": "UPDATE COMPARISON", "es": "ACTUALIZAR COMPARACIÓN"},
    "comparison_note": {"pt": "A maior base é usada como principal; as demais são comparadas sequencialmente.", "en": "The largest database is used as the reference; the others are compared sequentially.", "es": "La base más grande se usa como referencia; las demás se comparan secuencialmente."},
    "waiting_processing": {"pt": "Aguardando processamento", "en": "Waiting for processing", "es": "Esperando procesamiento"},
    "summary": {"pt": "Resumo", "en": "Summary", "es": "Resumen"},
    "raw": {"pt": "Bruto", "en": "Raw", "es": "Bruto"},
    "duplicates_removed": {"pt": "Duplicatas removidas", "en": "Duplicates removed", "es": "Duplicados eliminados"},
    "pending": {"pt": "Pendentes", "en": "Pending", "es": "Pendientes"},
    "final": {"pt": "Final", "en": "Final", "es": "Final"},
    "no_files_processed": {"pt": "Nenhum arquivo processado.", "en": "No files processed.", "es": "Ningún archivo procesado."},
    "duplicate_section": {"pt": "3. Comparação entre bases — duplicatas e casos para revisar", "en": "3. Cross-database comparison — duplicates and cases to review", "es": "3. Comparación entre bases — duplicados y casos para revisar"},
    "confirm_duplicate": {"pt": "✓ CONFIRMAR DUPLICATA", "en": "✓ CONFIRM DUPLICATE", "es": "✓ CONFIRMAR DUPLICADO"},
    "not_duplicate": {"pt": "✗ NÃO É DUPLICATA", "en": "✗ NOT A DUPLICATE", "es": "✗ NO ES DUPLICADO"},
    "all_duplicate": {"pt": "✓ TODAS = DUPLICATA", "en": "✓ ALL = DUPLICATE", "es": "✓ TODAS = DUPLICADO"},
    "all_different": {"pt": "✗ TODAS = DIFERENTES", "en": "✗ ALL = DIFFERENT", "es": "✗ TODAS = DIFERENTES"},
    "evidence": {"pt": "🔎 EVIDÊNCIAS", "en": "🔎 EVIDENCE", "es": "🔎 EVIDENCIAS"},
    "color_legend": {"pt": "Amarelo = revisar • Verde = duplicata confirmada • Cinza = não duplicata.", "en": "Yellow = review • Green = confirmed duplicate • Gray = not a duplicate.", "es": "Amarillo = revisar • Verde = duplicado confirmado • Gris = no duplicado."},
    "status": {"pt": "Status", "en": "Status", "es": "Estado"},
    "criterion": {"pt": "Critério", "en": "Criterion", "es": "Criterio"},
    "existing_base": {"pt": "Já existente - Base", "en": "Existing record - Database", "es": "Registro existente - Base"},
    "existing_title": {"pt": "Já existente - Título", "en": "Existing record - Title", "es": "Registro existente - Título"},
    "new_base": {"pt": "Novo registro - Base", "en": "New record - Database", "es": "Nuevo registro - Base"},
    "new_title": {"pt": "Novo registro - Título", "en": "New record - Title", "es": "Nuevo registro - Título"},
    "decision": {"pt": "Decisão", "en": "Decision", "es": "Decisión"},
    "export": {"pt": "4. Exportar", "en": "4. Export", "es": "4. Exportar"},
    "export_all": {"pt": "EXPORTAR TUDO PARA UMA PASTA", "en": "EXPORT ALL TO A FOLDER", "es": "EXPORTAR TODO A UNA CARPETA"},
    "export_note": {"pt": "Gera um Excel limpo para cada base + comparação de duplicatas + banco final para triagem + relatório.", "en": "Creates a clean Excel file for each database + duplicate comparison + final screening dataset + report.", "es": "Genera un Excel limpio para cada base + comparación de duplicados + base final para cribado + informe."},
    "ready": {"pt": "Pronto.", "en": "Ready.", "es": "Listo."},
    "records": {"pt": "registros", "en": "records", "es": "registros"},
    "main_database": {"pt": "Base principal", "en": "Reference database", "es": "Base principal"},
    "order": {"pt": "Ordem", "en": "Order", "es": "Orden"},
    "auto_duplicates": {"pt": "Duplicatas automáticas", "en": "Automatic duplicates", "es": "Duplicados automáticos"},
    "manual_duplicates": {"pt": "Duplicatas confirmadas manualmente", "en": "Manually confirmed duplicates", "es": "Duplicados confirmados manualmente"},
    "pending_check": {"pt": "Pendentes para conferir", "en": "Pending manual review", "es": "Pendientes de revisión"},
    "total_after": {"pt": "TOTAL APÓS DEDUPLICAÇÃO", "en": "TOTAL AFTER DEDUPLICATION", "es": "TOTAL DESPUÉS DE DEDUPLICACIÓN"},
    "second_not_run": {"pt": "2ª triagem (revisões/meta): ainda não executada", "en": "Second screening (reviews/meta-analyses): not yet run", "es": "Segundo cribado (revisiones/metaanálisis): aún no ejecutado"},
    "second_summary": {"pt": "2ª triagem (classificatória): {confirmed} confirmado(s) como revisão/meta | {pending} para revisar | Nenhum registro removido", "en": "Second screening (classification): {confirmed} confirmed as review/meta-analysis | {pending} to review | No records removed", "es": "Segundo cribado (clasificación): {confirmed} confirmado(s) como revisión/metaanálisis | {pending} para revisar | Ningún registro eliminado"},
}

DATABASE_UI_NAMES = {
    "Automático": {"pt": "Automático", "en": "Automatic", "es": "Automático"},
    "PubMed": {"pt": "PubMed", "en": "PubMed", "es": "PubMed"},
    "Scopus": {"pt": "Scopus", "en": "Scopus", "es": "Scopus"},
    "Web of Science": {"pt": "Web of Science", "en": "Web of Science", "es": "Web of Science"},
    "Embase": {"pt": "Embase", "en": "Embase", "es": "Embase"},
    "LILACS/BVS": {"pt": "LILACS/BVS", "en": "LILACS/BVS", "es": "LILACS/BVS"},
    "Cochrane Library": {"pt": "Cochrane Library", "en": "Cochrane Library", "es": "Cochrane Library"},
    "SciELO": {"pt": "SciELO", "en": "SciELO", "es": "SciELO"},
    "Outra": {"pt": "Outra", "en": "Other", "es": "Otra"},
}

DISPLAY_VALUE_TRANSLATIONS = {
    "REVISAR": {"pt": "REVISAR", "en": "REVIEW", "es": "REVISAR"},
    "DUPLICATA CONFIRMADA": {"pt": "DUPLICATA CONFIRMADA", "en": "CONFIRMED DUPLICATE", "es": "DUPLICADO CONFIRMADO"},
    "DUPLICATA CONFIRMADA MANUAL": {"pt": "DUPLICATA CONFIRMADA MANUAL", "en": "MANUALLY CONFIRMED DUPLICATE", "es": "DUPLICADO CONFIRMADO MANUALMENTE"},
    "NÃO DUPLICATA": {"pt": "NÃO DUPLICATA", "en": "NOT A DUPLICATE", "es": "NO DUPLICADO"},
    "DUPLICATA": {"pt": "DUPLICATA", "en": "DUPLICATE", "es": "DUPLICADO"},
    "DIFERENTE": {"pt": "DIFERENTE", "en": "DIFFERENT", "es": "DIFERENTE"},
    "DOI idêntico": {"pt": "DOI idêntico", "en": "Identical DOI", "es": "DOI idéntico"},
    "Título idêntico": {"pt": "Título idêntico", "en": "Identical title", "es": "Título idéntico"},
    "Título semelhante": {"pt": "Título semelhante", "en": "Similar title", "es": "Título similar"},
    "CONFIRMADO": {"pt": "CONFIRMADO", "en": "CONFIRMED", "es": "CONFIRMADO"},
    "NÃO É REVISÃO/META": {"pt": "NÃO É REVISÃO/META", "en": "NOT A REVIEW/META-ANALYSIS", "es": "NO ES REVISIÓN/METAANÁLISIS"},
    "REVISÃO SISTEMÁTICA": {"pt": "REVISÃO SISTEMÁTICA", "en": "SYSTEMATIC REVIEW", "es": "REVISIÓN SISTEMÁTICA"},
    "META-ANÁLISE": {"pt": "META-ANÁLISE", "en": "META-ANALYSIS", "es": "METAANÁLISIS"},
    "REVISÃO SISTEMÁTICA + META-ANÁLISE": {"pt": "REVISÃO SISTEMÁTICA + META-ANÁLISE", "en": "SYSTEMATIC REVIEW + META-ANALYSIS", "es": "REVISIÓN SISTEMÁTICA + METAANÁLISIS"},
    "REVISÃO DE REVISÕES": {"pt": "REVISÃO DE REVISÕES", "en": "OVERVIEW OF REVIEWS", "es": "REVISIÓN DE REVISIONES"},
    "POSSÍVEL REVISÃO SISTEMÁTICA/META-ANÁLISE": {"pt": "POSSÍVEL REVISÃO SISTEMÁTICA/META-ANÁLISE", "en": "POSSIBLE SYSTEMATIC REVIEW/META-ANALYSIS", "es": "POSIBLE REVISIÓN SISTEMÁTICA/METAANÁLISIS"},
    "AUTOMÁTICA": {"pt": "AUTOMÁTICA", "en": "AUTOMATIC", "es": "AUTOMÁTICA"},
    "MANUAL": {"pt": "MANUAL", "en": "MANUAL", "es": "MANUAL"},
    "Título": {"pt": "Título", "en": "Title", "es": "Título"},
    "Resumo": {"pt": "Resumo", "en": "Abstract", "es": "Resumen"},
}

LITERAL_UI_TEXT = {
    "Painel do projeto": {"en":"Project dashboard", "es":"Panel del proyecto"},
    "PAINEL DO PROJETO": {"en":"PROJECT DASHBOARD", "es":"PANEL DEL PROYECTO"},
    "Projeto ainda não salvo": {"en":"Project not saved yet", "es":"Proyecto aún no guardado"},
    "REGISTROS BRUTOS": {"en":"RAW RECORDS", "es":"REGISTROS BRUTOS"},
    "DUPLICATAS REMOVIDAS": {"en":"DUPLICATES REMOVED", "es":"DUPLICADOS ELIMINADOS"},
    "APÓS DEDUPLICAÇÃO": {"en":"AFTER DEDUPLICATION", "es":"DESPUÉS DE DEDUPLICACIÓN"},
    "DUPLICATAS PENDENTES": {"en":"PENDING DUPLICATES", "es":"DUPLICADOS PENDIENTES"},
    "BASES IMPORTADAS": {"en":"DATABASES IMPORTED", "es":"BASES IMPORTADAS"},
    "REVISÕES/META CONFIRMADAS": {"en":"REVIEWS/META-ANALYSES CONFIRMED", "es":"REVISIONES/META CONFIRMADAS"},
    "2ª TRIAGEM PENDENTE": {"en":"SECOND SCREENING PENDING", "es":"SEGUNDO CRIBADO PENDIENTE"},
    "Resumo por base": {"en":"Summary by database", "es":"Resumen por base"},
    "ATUALIZAR": {"en":"UPDATE", "es":"ACTUALIZAR"},
    "BIBLIOMETRIA": {"en":"BIBLIOMETRICS", "es":"BIBLIOMETRÍA"},
    "SALVAR PROJETO": {"en":"SAVE PROJECT", "es":"GUARDAR PROYECTO"},
    "Formatador de referências — ABNT / Vancouver / APA": {"en":"Reference formatter — ABNT / Vancouver / APA", "es":"Formateador de referencias — ABNT / Vancouver / APA"},
    "Objetivo": {"en":"Objective", "es":"Objetivo"},
    "Arquivo de entrada": {"en":"Input file", "es":"Archivo de entrada"},
    "Saída": {"en":"Output", "es":"Salida"},
    "ESCOLHER...": {"en":"CHOOSE...", "es":"ELEGIR..."},
    "Estilo:": {"en":"Style:", "es":"Estilo:"},
    "Tentar localizar metadados pelo texto quando não houver DOI": {"en":"Try to find metadata from the reference text when no DOI is available", "es":"Intentar localizar metadatos por el texto cuando no haya DOI"},
    "FORMATAR REFERÊNCIAS": {"en":"FORMAT REFERENCES", "es":"FORMATEAR REFERENCIAS"},
    "Iniciando...": {"en":"Starting...", "es":"Iniciando..."},
    "Mesclagem e arquivo pronto para Biblioshiny": {"en":"Merge and Biblioshiny-ready file", "es":"Combinación y archivo listo para Biblioshiny"},
    "Bases para a mesclagem principal": {"en":"Databases for the main merge", "es":"Bases para la combinación principal"},
    "ABRIR BIBLIOSHINY": {"en":"OPEN BIBLIOSHINY", "es":"ABRIR BIBLIOSHINY"},
    "REALIZAR MESCLAGEM": {"en":"RUN MERGE", "es":"REALIZAR COMBINACIÓN"},
    "Abrir o Biblioshiny ao terminar (requer R + pacote bibliometrix)": {"en":"Open Biblioshiny when finished (requires R + bibliometrix package)", "es":"Abrir Biblioshiny al finalizar (requiere R + paquete bibliometrix)"},
    "2ª triagem — Marcação de revisões sistemáticas e meta-análises": {"en":"Second screening — Systematic review and meta-analysis marking", "es":"Segundo cribado — Marcación de revisiones sistemáticas y metaanálisis"},
    "Prévia classificatória da 2ª triagem": {"en":"Second-screening classification preview", "es":"Vista previa de clasificación del segundo cribado"},
    "Classificação": {"en":"Classification", "es":"Clasificación"},
    "Detectado em": {"en":"Detected in", "es":"Detectado en"},
    "Buscar:": {"en":"Search:", "es":"Buscar:"},
    "CONFIRMAR TIPO DETECTADO": {"en":"CONFIRM DETECTED TYPE", "es":"CONFIRMAR TIPO DETECTADO"},
    "Classificar como:": {"en":"Classify as:", "es":"Clasificar como:"},
    "APLICAR CLASSIFICAÇÃO": {"en":"APPLY CLASSIFICATION", "es":"APLICAR CLASIFICACIÓN"},
    "RESTAURAR AUTOMÁTICO": {"en":"RESTORE AUTOMATIC", "es":"RESTAURAR AUTOMÁTICO"},
    "Verde = confirmado | Amarelo = revisar | Cinza = falso positivo": {"en":"Green = confirmed | Yellow = review | Gray = false positive", "es":"Verde = confirmado | Amarillo = revisar | Gris = falso positivo"},
    "CONFIRMAR MARCAÇÕES DA 2ª TRIAGEM": {"en":"CONFIRM SECOND-SCREENING MARKS", "es":"CONFIRMAR MARCAS DEL SEGUNDO CRIBADO"},
    "FECHAR SEM CONFIRMAR": {"en":"CLOSE WITHOUT CONFIRMING", "es":"CERRAR SIN CONFIRMAR"},
    "As marcações serão exportadas; nenhum artigo será excluído.": {"en":"Marks will be exported; no article will be excluded.", "es":"Las marcas se exportarán; ningún artículo será excluido."},
    "Buscar artigo nas bases": {"en":"Search article across databases", "es":"Buscar artículo en las bases"},
    "Título, DOI, PMID/ID ou autor:": {"en":"Title, DOI, PMID/ID or author:", "es":"Título, DOI, PMID/ID o autor:"},
    "Situação": {"en":"Status", "es":"Situación"},
    "PRISMA automático": {"en":"Automatic PRISMA", "es":"PRISMA automático"},
    "Contagens automáticas": {"en":"Automatic counts", "es":"Recuentos automáticos"},
    "Etapas posteriores — números agregados opcionais": {"en":"Later stages — optional aggregate counts", "es":"Etapas posteriores — recuentos agregados opcionales"},
    "Excluídos na triagem externa de título/resumo:": {"en":"Excluded during external title/abstract screening:", "es":"Excluidos en el cribado externo de título/resumen:"},
    "Relatórios não recuperados:": {"en":"Reports not retrieved:", "es":"Informes no recuperados:"},
    "Relatórios excluídos após texto completo:": {"en":"Reports excluded after full-text assessment:", "es":"Informes excluidos tras texto completo:"},
    "Estudos incluídos na revisão:": {"en":"Studies included in the review:", "es":"Estudios incluidos en la revisión:"},
    "Motivos das exclusões em texto completo:": {"en":"Reasons for full-text exclusions:", "es":"Motivos de exclusión a texto completo:"},
    "ATUALIZAR DADOS": {"en":"UPDATE DATA", "es":"ACTUALIZAR DATOS"},
    "GERAR IMAGEM PNG": {"en":"GENERATE PNG IMAGE", "es":"GENERAR IMAGEN PNG"},
    "EXPORTAR DADOS EXCEL": {"en":"EXPORT EXCEL DATA", "es":"EXPORTAR DATOS EXCEL"},
    "Título": {"en":"Title", "es":"Título"},
    "Base": {"en":"Database", "es":"Base"},
    "Bruto": {"en":"Raw", "es":"Bruto"},
    "Duplicatas": {"en":"Duplicates", "es":"Duplicados"},
    "Pendentes": {"en":"Pending", "es":"Pendientes"},
    "Final": {"en":"Final", "es":"Final"},
    "Aguardando": {"en":"Waiting", "es":"Esperando"},
    "Selecione um arquivo e clique em FORMATAR REFERÊNCIAS.": {"en":"Select a file and click FORMAT REFERENCES.", "es":"Seleccione un archivo y haga clic en FORMATEAR REFERENCIAS."},
    "Carregue um DOCX ou TXT com uma referência por parágrafo/linha. O módulo extrai o DOI, consulta os metadados no Crossref e cria um NOVO DOCX formatado. Referências sem correspondência segura são preservadas e marcadas para revisão manual.": {"en":"Load a DOCX or TXT file with one reference per paragraph/line. The module extracts the DOI, queries Crossref metadata, and creates a NEW formatted DOCX. References without a reliable match are preserved and flagged for manual review.", "es":"Cargue un archivo DOCX o TXT con una referencia por párrafo/línea. El módulo extrae el DOI, consulta los metadatos de Crossref y crea un NUEVO DOCX formateado. Las referencias sin coincidencia segura se conservan y se marcan para revisión manual."},
    "A busca por texto é opcional porque pode encontrar trabalhos parecidos; o programa só aceita automaticamente correspondências com escore alto. DOI explícito é sempre priorizado.": {"en":"Text-based searching is optional because it may retrieve similar records; the software only accepts high-score matches automatically. An explicit DOI always has priority.", "es":"La búsqueda por texto es opcional porque puede encontrar trabajos similares; el programa solo acepta automáticamente coincidencias con puntuación alta. Un DOI explícito siempre tiene prioridad."},
    "Mescla Web of Science, Embase e/ou Scopus, remove duplicatas por DOI idêntico ou título + ano idênticos e gera CSV, Excel e BibTeX. Também tenta criar BIBLIOSHINY_PRONTO.RData para você carregar diretamente no Biblioshiny sem refazer a mesclagem.": {"en":"Merges Web of Science, Embase and/or Scopus, removes duplicates by identical DOI or identical title + year, and generates CSV, Excel and BibTeX files. It also attempts to create BIBLIOSHINY_PRONTO.RData for direct loading into Biblioshiny without repeating the merge.", "es":"Combina Web of Science, Embase y/o Scopus, elimina duplicados por DOI idéntico o título + año idénticos y genera archivos CSV, Excel y BibTeX. También intenta crear BIBLIOSHINY_PRONTO.RData para cargarlo directamente en Biblioshiny sin repetir la combinación."},
    "Importante: o programa não inventa metadados. Palavras-chave, referências citadas, afiliações e citações só estarão disponíveis se esses campos tiverem sido importados do arquivo original.": {"en":"Important: the software does not invent metadata. Keywords, cited references, affiliations and citation counts are available only when those fields were imported from the original file.", "es":"Importante: el programa no inventa metadatos. Las palabras clave, referencias citadas, afiliaciones y citas solo estarán disponibles si esos campos fueron importados del archivo original."},
    "O caminho do RData será copiado e o arquivo será selecionado no Explorador para facilitar a última etapa de carga.": {"en":"The RData path will be copied and the file selected in Explorer to simplify the final loading step.", "es":"La ruta del RData se copiará y el archivo se seleccionará en el Explorador para facilitar el último paso de carga."},
    "2ª triagem classificatória de revisões/meta-análises: ainda não executada": {"en":"Second screening for review/meta-analysis classification: not yet run", "es":"Segundo cribado para clasificación de revisiones/metaanálisis: aún no ejecutado"},
    "Revisões/meta-análises marcadas na 2ª triagem:": {"en":"Reviews/meta-analyses marked in the second screening:", "es":"Revisiones/metaanálisis marcados en el segundo cribado:"},
    "Identificados:": {"en":"Identified:", "es":"Identificados:"},
    "Duplicatas removidas:": {"en":"Duplicates removed:", "es":"Duplicados eliminados:"},
    "Após deduplicação:": {"en":"After deduplication:", "es":"Después de deduplicación:"},
    "Disponíveis para triagem de título/resumo:": {"en":"Available for title/abstract screening:", "es":"Disponibles para cribado de título/resumen:"},
    "Ano": {"en":"Year", "es":"Año"},
    "Triagem": {"en":"Screening", "es":"Cribado"},
    "Triagem de título e resumo": {"en":"Title and abstract screening", "es":"Cribado de título y resumen"},
    "✓ INCLUIR": {"en":"✓ INCLUDE", "es":"✓ INCLUIR"},
    "✗ EXCLUIR": {"en":"✗ EXCLUDE", "es":"✗ EXCLUIR"},
    "? TALVEZ": {"en":"? MAYBE", "es":"? QUIZÁS"},
    "LIMPAR DECISÃO": {"en":"CLEAR DECISION", "es":"BORRAR DECISIÓN"},
    "As decisões ficam salvas no projeto e entram no PRISMA automático.": {"en":"Decisions are saved in the project and included in the automatic PRISMA counts.", "es":"Las decisiones se guardan en el proyecto y se incluyen en los recuentos automáticos de PRISMA."},
    "Selecione o arquivo com as referências": {"en":"Select the file containing the references", "es":"Seleccione el archivo que contiene las referencias"},
}


# Textos dinâmicos exibidos em caixas de diálogo, barra de status,
# barra de progresso e seletores de arquivo. Os dados internos continuam
# em português para manter compatibilidade com projetos antigos.
MESSAGE_UI_TEXT = {
    # Títulos / seções curtas
    "Processamento": {"en":"Processing", "es":"Procesamiento"},
    "Sem arquivos": {"en":"No files", "es":"Sin archivos"},
    "Falha": {"en":"Failure", "es":"Fallo"},
    "Processado com avisos": {"en":"Processed with warnings", "es":"Procesado con avisos"},
    "Processado": {"en":"Processed", "es":"Procesado"},
    "Comparação": {"en":"Comparison", "es":"Comparación"},
    "Comparação atualizada": {"en":"Comparison updated", "es":"Comparación actualizada"},
    "Erro de comparação": {"en":"Comparison error", "es":"Error de comparación"},
    "Seleção": {"en":"Selection", "es":"Selección"},
    "Sem alteração": {"en":"No change", "es":"Sin cambios"},
    "Duplicatas": {"en":"Duplicates", "es":"Duplicados"},
    "Decisão em lote": {"en":"Batch decision", "es":"Decisión por lote"},
    "Decisão em lote concluída": {"en":"Batch decision completed", "es":"Decisión por lote completada"},
    "Novo projeto": {"en":"New project", "es":"Nuevo proyecto"},
    "Painel": {"en":"Dashboard", "es":"Panel"},
    "Auditoria": {"en":"Audit", "es":"Auditoría"},
    "Formatador": {"en":"Formatter", "es":"Formateador"},
    "Formatador de referências": {"en":"Reference formatter", "es":"Formateador de referencias"},
    "Referências formatadas": {"en":"References formatted", "es":"Referencias formateadas"},
    "Biblioshiny iniciado": {"en":"Biblioshiny started", "es":"Biblioshiny iniciado"},
    "Bibliometria": {"en":"Bibliometrics", "es":"Bibliometría"},
    "Projeto": {"en":"Project", "es":"Proyecto"},
    "Evidência": {"en":"Evidence", "es":"Evidencia"},
    "Evidências da comparação": {"en":"Comparison evidence", "es":"Evidencias de la comparación"},
    "2ª triagem": {"en":"Second screening", "es":"Segundo cribado"},
    "Confirmar marcações da 2ª triagem": {"en":"Confirm second-screening marks", "es":"Confirmar marcas del segundo cribado"},
    "2ª triagem confirmada": {"en":"Second screening confirmed", "es":"Segundo cribado confirmado"},
    "Fechar sem confirmar": {"en":"Close without confirming", "es":"Cerrar sin confirmar"},
    "Triagem": {"en":"Screening", "es":"Cribado"},
    "Motivo da exclusão": {"en":"Reason for exclusion", "es":"Motivo de exclusión"},
    "Buscar artigo": {"en":"Search article", "es":"Buscar artículo"},
    "PRISMA": {"en":"PRISMA", "es":"PRISMA"},
    "Sem dados": {"en":"No data", "es":"Sin datos"},
    "Erro de exportação": {"en":"Export error", "es":"Error de exportación"},
    "Exportação concluída": {"en":"Export completed", "es":"Exportación completada"},

    # Barra de status / progresso
    "Arquivos adicionados com detecção automática de base/formato.": {"en":"Files added with automatic database/format detection.", "es":"Archivos añadidos con detección automática de base/formato."},
    "Preparando arquivos": {"en":"Preparing files", "es":"Preparando archivos"},
    "Importando arquivos": {"en":"Importing files", "es":"Importando archivos"},
    "Arquivos importados": {"en":"Files imported", "es":"Archivos importados"},
    "Comparando as bases...": {"en":"Comparing databases...", "es":"Comparando bases..."},
    "Iniciando comparação": {"en":"Starting comparison", "es":"Iniciando comparación"},
    "Concluído": {"en":"Completed", "es":"Completado"},
    "Comparação concluída": {"en":"Comparison completed", "es":"Comparación completada"},
    "registros no banco atual": {"en":"records in the current dataset", "es":"registros en la base actual"},
    "Atualizando decisão": {"en":"Updating decision", "es":"Actualizando decisión"},
    "Decisão atualizada": {"en":"Decision updated", "es":"Decisión actualizada"},
    "decisão(ões) atualizada(s)": {"en":"decision(s) updated", "es":"decisión(es) actualizada(s)"},
    "Decisão em lote aplicada a": {"en":"Batch decision applied to", "es":"Decisión por lote aplicada a"},
    "caso(s)": {"en":"case(s)", "es":"caso(s)"},
    "decisão(ões) em lote atualizada(s)": {"en":"batch decision(s) updated", "es":"decisión(es) por lote actualizada(s)"},
    "Erro durante a comparação": {"en":"Error during comparison", "es":"Error durante la comparación"},
    "Projeto salvo:": {"en":"Project saved:", "es":"Proyecto guardado:"},
    "Projeto aberto:": {"en":"Project opened:", "es":"Proyecto abierto:"},
    "Exportado para:": {"en":"Exported to:", "es":"Exportado a:"},
    "2ª triagem confirmada:": {"en":"Second screening confirmed:", "es":"Segundo cribado confirmado:"},
    "revisão(ões)/meta-análise(s) marcada(s)": {"en":"review(s)/meta-analysis(es) marked", "es":"revisión(es)/metaanálisis marcado(s)"},
    "nenhum registro removido": {"en":"no records removed", "es":"ningún registro eliminado"},

    # Mensagens principais
    "A comparação ainda está em andamento. Aguarde a conclusão antes de processar novamente.": {"en":"The comparison is still running. Wait for it to finish before processing again.", "es":"La comparación todavía está en curso. Espere a que finalice antes de procesar nuevamente."},
    "Adicione pelo menos um arquivo.": {"en":"Add at least one file.", "es":"Añada al menos un archivo."},
    "Nenhum registro pôde ser importado.": {"en":"No records could be imported.", "es":"No se pudo importar ningún registro."},
    "Alguns arquivos não foram lidos:": {"en":"Some files could not be read:", "es":"Algunos archivos no pudieron leerse:"},
    "A ferramenta aceita várias variantes por base. Se um novo formato aparecer, envie o arquivo para adicionarmos mais um detector sem alterar os anteriores.": {"en":"The tool accepts multiple export variants per database. If a new format appears, provide the file so another detector can be added without changing the existing ones.", "es":"La herramienta acepta varias variantes de exportación por base. Si aparece un formato nuevo, proporcione el archivo para añadir otro detector sin modificar los existentes."},
    "registros foram importados de": {"en":"records were imported from", "es":"registros fueron importados de"},
    "base(s).": {"en":"database(s).", "es":"base(s)."},
    "Já existe uma comparação em andamento.": {"en":"A comparison is already running.", "es":"Ya hay una comparación en curso."},
    "Base principal:": {"en":"Reference database:", "es":"Base principal:"},
    "Registros finais atuais:": {"en":"Current final records:", "es":"Registros finales actuales:"},
    "Comparações registradas:": {"en":"Recorded comparisons:", "es":"Comparaciones registradas:"},
    "Selecione pelo menos uma linha amarela (REVISAR).": {"en":"Select at least one yellow row (REVIEW).", "es":"Seleccione al menos una fila amarilla (REVISAR)."},
    "A seleção não contém um caso de similaridade que possa ser revisado.": {"en":"The selection does not contain a similarity case that can be reviewed.", "es":"La selección no contiene un caso de similitud que pueda revisarse."},
    "comparação(ões) confirmada(s) como duplicata.": {"en":"comparison(s) confirmed as duplicate.", "es":"comparación(es) confirmada(s) como duplicado."},
    "comparação(ões) mantida(s) como diferente.": {"en":"comparison(s) kept as different.", "es":"comparación(es) mantenida(s) como diferente."},
    "Não há comparações disponíveis.": {"en":"No comparisons are available.", "es":"No hay comparaciones disponibles."},
    "Não há duplicatas pendentes para decidir.": {"en":"There are no pending duplicates to decide.", "es":"No hay duplicados pendientes por decidir."},
    "Confirmar TODOS os": {"en":"Confirm ALL", "es":"Confirmar TODOS los"},
    "casos amarelos como duplicata?": {"en":"yellow cases as duplicates?", "es":"casos amarillos como duplicados?"},
    "Esses registros serão removidos do banco final.": {"en":"These records will be removed from the final dataset.", "es":"Estos registros se eliminarán de la base final."},
    "Marcar TODOS os": {"en":"Mark ALL", "es":"Marcar TODOS los"},
    "casos amarelos como não duplicata?": {"en":"yellow cases as not duplicates?", "es":"casos amarillos como no duplicados?"},
    "Esses registros permanecerão no banco final.": {"en":"These records will remain in the final dataset.", "es":"Estos registros permanecerán en la base final."},
    "confirmado(s) como duplicata.": {"en":"confirmed as duplicate(s).", "es":"confirmado(s) como duplicado(s)."},
    "mantido(s) como não duplicata.": {"en":"kept as not duplicate(s).", "es":"mantenido(s) como no duplicado(s)."},
    "Iniciar um novo projeto apagará da tela os dados atuais.": {"en":"Starting a new project will clear the current data from the screen.", "es":"Iniciar un nuevo proyecto borrará de la pantalla los datos actuales."},
    "Salve o projeto antes, se quiser continuar depois. Deseja prosseguir?": {"en":"Save the project first if you want to continue later. Do you want to proceed?", "es":"Guarde primero el proyecto si desea continuar después. ¿Desea continuar?"},
    "Processe pelo menos uma base para abrir o painel do projeto.": {"en":"Process at least one database before opening the project dashboard.", "es":"Procese al menos una base antes de abrir el panel del proyecto."},
    "Processe pelo menos uma base antes de exportar a auditoria.": {"en":"Process at least one database before exporting the audit.", "es":"Procese al menos una base antes de exportar la auditoría."},
    "Relatórios criados:": {"en":"Reports created:", "es":"Informes creados:"},
    "Selecione um arquivo DOCX ou TXT válido.": {"en":"Select a valid DOCX or TXT file.", "es":"Seleccione un archivo DOCX o TXT válido."},
    "Concluído.": {"en":"Completed.", "es":"Completado."},
    "Formatadas:": {"en":"Formatted:", "es":"Formateadas:"},
    "Para revisão manual:": {"en":"For manual review:", "es":"Para revisión manual:"},
    "Pasta:": {"en":"Folder:", "es":"Carpeta:"},
    "Rscript não foi localizado. Instale o R e o pacote bibliometrix ou adicione o R ao PATH.": {"en":"Rscript was not found. Install R and the bibliometrix package or add R to PATH.", "es":"No se encontró Rscript. Instale R y el paquete bibliometrix o añada R al PATH."},
    "O arquivo BIBLIOSHINY_PRONTO.RData ainda não existe.": {"en":"The BIBLIOSHINY_PRONTO.RData file does not exist yet.", "es":"El archivo BIBLIOSHINY_PRONTO.RData aún no existe."},
    "O Biblioshiny foi iniciado.": {"en":"Biblioshiny has been started.", "es":"Biblioshiny se ha iniciado."},
    "O arquivo pronto está selecionado no Explorador e o caminho foi copiado.": {"en":"The ready-to-use file is selected in File Explorer and its path has been copied.", "es":"El archivo listo está seleccionado en el Explorador y su ruta se ha copiado."},
    "No Biblioshiny use Data > Import or Load > Load files e escolha BIBLIOSHINY_PRONTO.RData.": {"en":"In Biblioshiny, use Data > Import or Load > Load files and select BIBLIOSHINY_PRONTO.RData.", "es":"En Biblioshiny, use Data > Import or Load > Load files y seleccione BIBLIOSHINY_PRONTO.RData."},
    "O Biblioshiny oficial não oferece um argumento para pré-carregar um arquivo ao iniciar; por isso essa última seleção é necessária.": {"en":"The official Biblioshiny interface does not provide an argument to preload a file at startup; therefore this final selection is required.", "es":"La interfaz oficial de Biblioshiny no ofrece un argumento para precargar un archivo al iniciar; por eso esta selección final es necesaria."},
    "Processe as bases antes de realizar a mesclagem bibliométrica.": {"en":"Process the databases before running the bibliometric merge.", "es":"Procese las bases antes de realizar la combinación bibliométrica."},
    "Selecione pelo menos uma base ou o PubMed separado.": {"en":"Select at least one database or the separate PubMed export.", "es":"Seleccione al menos una base o la exportación separada de PubMed."},
    "Processe pelo menos uma base antes de salvar o projeto.": {"en":"Process at least one database before saving the project.", "es":"Procese al menos una base antes de guardar el proyecto."},
    "Não foi possível salvar o projeto.": {"en":"The project could not be saved.", "es":"No se pudo guardar el proyecto."},
    "Projeto salvo com sucesso.": {"en":"Project saved successfully.", "es":"Proyecto guardado correctamente."},
    "Arquivo de projeto inválido.": {"en":"Invalid project file.", "es":"Archivo de proyecto no válido."},
    "Não foi possível restaurar o projeto.": {"en":"The project could not be restored.", "es":"No se pudo restaurar el proyecto."},
    "Projeto restaurado com sucesso.": {"en":"Project restored successfully.", "es":"Proyecto restaurado correctamente."},
    "Selecione uma comparação na grade.": {"en":"Select a comparison in the table.", "es":"Seleccione una comparación en la tabla."},
    "Processe e deduplique as bases antes de executar a 2ª triagem.": {"en":"Process and deduplicate the databases before running the second screening.", "es":"Procese y deduplique las bases antes de ejecutar el segundo cribado."},
    "Ainda existem": {"en":"There are still", "es":"Todavía hay"},
    "possível(is) duplicata(s) aguardando decisão.": {"en":"possible duplicate(s) awaiting a decision.", "es":"posible(s) duplicado(s) esperando una decisión."},
    "Resolva primeiro a conferência de duplicatas; depois execute a 2ª triagem.": {"en":"Resolve duplicate review first; then run the second screening.", "es":"Resuelva primero la revisión de duplicados; después ejecute el segundo cribado."},
    "Selecione um ou mais registros.": {"en":"Select one or more records.", "es":"Seleccione uno o más registros."},
    "Escolha uma classificação válida.": {"en":"Choose a valid classification.", "es":"Elija una clasificación válida."},
    "Salvar definitivamente estas marcações?": {"en":"Permanently save these marks?", "es":"¿Guardar definitivamente estas marcas?"},
    "Registros detectados:": {"en":"Detected records:", "es":"Registros detectados:"},
    "Classificados como revisão/meta:": {"en":"Classified as review/meta-analysis:", "es":"Clasificados como revisión/metaanálisis:"},
    "Ainda para revisar:": {"en":"Still to review:", "es":"Aún por revisar:"},
    "Marcados como falso positivo:": {"en":"Marked as false positives:", "es":"Marcados como falsos positivos:"},
    "Registros no banco final:": {"en":"Records in the final dataset:", "es":"Registros en la base final:"},
    "Nenhum artigo será excluído. As classificações serão incluídas nas exportações.": {"en":"No article will be excluded. Classifications will be included in the exports.", "es":"No se excluirá ningún artículo. Las clasificaciones se incluirán en las exportaciones."},
    "Marcações salvas com sucesso.": {"en":"Marks saved successfully.", "es":"Marcas guardadas correctamente."},
    "Confirmados:": {"en":"Confirmed:", "es":"Confirmados:"},
    "Para revisar:": {"en":"To review:", "es":"Para revisar:"},
    "Falsos positivos:": {"en":"False positives:", "es":"Falsos positivos:"},
    "Registros mantidos no banco:": {"en":"Records retained in the dataset:", "es":"Registros mantenidos en la base:"},
    "Fechar esta janela sem salvar as alterações desta conferência?": {"en":"Close this window without saving the changes from this review?", "es":"¿Cerrar esta ventana sin guardar los cambios de esta revisión?"},
    "As marcações confirmadas anteriormente, se houver, permanecerão inalteradas.": {"en":"Previously confirmed marks, if any, will remain unchanged.", "es":"Las marcas confirmadas anteriormente, si las hubiera, permanecerán sin cambios."},
    "Nenhuma revisão sistemática ou meta-análise explícita foi detectada. Você pode confirmar esta etapa para registrar zero marcações.": {"en":"No explicit systematic review or meta-analysis was detected. You can confirm this step to record zero marks.", "es":"No se detectó ninguna revisión sistemática o metaanálisis explícito. Puede confirmar esta etapa para registrar cero marcas."},
    "Processe e deduplique as bases antes de iniciar a triagem.": {"en":"Process and deduplicate the databases before starting screening.", "es":"Procese y deduplique las bases antes de iniciar el cribado."},
    "Selecione um artigo.": {"en":"Select an article.", "es":"Seleccione un artículo."},
    "Informe o motivo da exclusão (opcional):": {"en":"Enter the reason for exclusion (optional):", "es":"Indique el motivo de exclusión (opcional):"},
    "Nenhuma base foi processada.": {"en":"No database has been processed.", "es":"No se ha procesado ninguna base."},
    "Use apenas números inteiros iguais ou maiores que zero.": {"en":"Use only integer numbers greater than or equal to zero.", "es":"Use solo números enteros mayores o iguales a cero."},
    "Contagens agregadas atualizadas.": {"en":"Aggregate counts updated.", "es":"Recuentos agregados actualizados."},
    "Fluxograma gerado em:": {"en":"Flowchart generated at:", "es":"Diagrama de flujo generado en:"},
    "Dados salvos em:": {"en":"Data saved at:", "es":"Datos guardados en:"},
    "Processe os arquivos antes de exportar.": {"en":"Process the files before exporting.", "es":"Procese los archivos antes de exportar."},
    "arquivo(s) foram criados em:": {"en":"file(s) were created in:", "es":"archivo(s) fueron creados en:"},

    # Diálogos de arquivo
    "Adicionar arquivos -": {"en":"Add files -", "es":"Añadir archivos -"},
    "Arquivos bibliográficos": {"en":"Bibliographic files", "es":"Archivos bibliográficos"},
    "Texto": {"en":"Text", "es":"Texto"},
    "Todos": {"en":"All files", "es":"Todos los archivos"},
    "Selecionar pasta": {"en":"Select folder", "es":"Seleccionar carpeta"},

    # Logs e erros técnicos que também podem chegar às caixas de aviso
    "Consultando referência": {"en":"Querying reference", "es":"Consultando referencia"},
    "Salvando arquivos...": {"en":"Saving files...", "es":"Guardando archivos..."},
    "Mesclando e deduplicando:": {"en":"Merging and deduplicating:", "es":"Combinando y deduplicando:"},
    "Gerando BIBLIOSHINY_PRONTO.RData...": {"en":"Generating BIBLIOSHINY_PRONTO.RData...", "es":"Generando BIBLIOSHINY_PRONTO.RData..."},
    "Preparando bases selecionadas...": {"en":"Preparing selected databases...", "es":"Preparando bases seleccionadas..."},
    "Salvando CSV mesclado...": {"en":"Saving merged CSV...", "es":"Guardando CSV combinado..."},
    "Salvando BibTeX mesclado e deduplicado...": {"en":"Saving merged and deduplicated BibTeX...", "es":"Guardando BibTeX combinado y deduplicado..."},
    "Salvando Excel de auditoria...": {"en":"Saving audit Excel file...", "es":"Guardando Excel de auditoría..."},
    "Preparando arquivo direto para Biblioshiny...": {"en":"Preparing direct Biblioshiny file...", "es":"Preparando archivo directo para Biblioshiny..."},
    "Salvando PubMed separadamente...": {"en":"Saving PubMed separately...", "es":"Guardando PubMed por separado..."},
    "Registros de entrada na mesclagem:": {"en":"Input records for merge:", "es":"Registros de entrada en la combinación:"},
    "Duplicatas exatas removidas:": {"en":"Exact duplicates removed:", "es":"Duplicados exactos eliminados:"},
    "Registros finais salvos:": {"en":"Final records saved:", "es":"Registros finales guardados:"},
    "BibTeX pronto:": {"en":"BibTeX ready:", "es":"BibTeX listo:"},
    "RData pronto para Biblioshiny:": {"en":"RData ready for Biblioshiny:", "es":"RData listo para Biblioshiny:"},
    "RData não foi criado automaticamente.": {"en":"RData was not created automatically.", "es":"RData no se creó automáticamente."},
    "AVISO:": {"en":"WARNING:", "es":"AVISO:"},
    "PubMed salvo separadamente:": {"en":"PubMed saved separately:", "es":"PubMed guardado por separado:"},
    "Mesclagem concluída.": {"en":"Merge completed.", "es":"Combinación completada."},
    "Entrada:": {"en":"Input:", "es":"Entrada:"},
    "Duplicatas removidas:": {"en":"Duplicates removed:", "es":"Duplicados eliminados:"},
    "Registros finais:": {"en":"Final records:", "es":"Registros finales:"},
    "Formatadas automaticamente:": {"en":"Automatically formatted:", "es":"Formateadas automáticamente:"},
    "Arquivos criados:": {"en":"Files created:", "es":"Archivos creados:"},
    "SALVO": {"en":"SAVED", "es":"GUARDADO"},
    "NÃO SALVO": {"en":"NOT SAVED", "es":"NO GUARDADO"},
    "Ordem de comparação:": {"en":"Comparison order:", "es":"Orden de comparación:"},
    "Limiar fuzzy:": {"en":"Fuzzy threshold:", "es":"Umbral fuzzy:"},
    "Versão": {"en":"Version", "es":"Versión"},

    "Use um arquivo .docx ou .txt com uma referência por parágrafo/linha.": {"en":"Use a .docx or .txt file with one reference per paragraph/line.", "es":"Use un archivo .docx o .txt con una referencia por párrafo/línea."},
    "Não foi possível extrair texto do PDF.": {"en":"Text could not be extracted from the PDF.", "es":"No se pudo extraer texto del PDF."},
    "Instale 'pypdf' ou use outro formato de exportação do Embase.": {"en":"Install 'pypdf' or use another Embase export format.", "es":"Instale 'pypdf' o use otro formato de exportación de Embase."},
    "Detalhes:": {"en":"Details:", "es":"Detalles:"},
    "Não foi possível ler o Excel do Web of Science.": {"en":"The Web of Science Excel file could not be read.", "es":"No se pudo leer el archivo Excel de Web of Science."},
    "Não foi possível ler o Excel .xls do Web of Science.": {"en":"The Web of Science .xls file could not be read.", "es":"No se pudo leer el archivo .xls de Web of Science."},
    "Nenhum registro Embase foi reconhecido.": {"en":"No Embase records were recognized.", "es":"No se reconoció ningún registro de Embase."},
    "Nenhum registro foi reconhecido no layout vertical do Embase.": {"en":"No records were recognized in the Embase vertical layout.", "es":"No se reconoció ningún registro en el diseño vertical de Embase."},
    "XML não reconhecido como exportação do Embase.": {"en":"XML was not recognized as an Embase export.", "es":"El XML no se reconoció como exportación de Embase."},
    "Arquivo não reconhecido como PubMed NBIB/MEDLINE.": {"en":"File was not recognized as PubMed NBIB/MEDLINE.", "es":"El archivo no se reconoció como PubMed NBIB/MEDLINE."},
    "Formato PubMed Summary (text) não reconhecido.": {"en":"PubMed Summary (text) format was not recognized.", "es":"No se reconoció el formato PubMed Summary (text)."},
    "Nenhum registro PubMed Summary foi reconhecido.": {"en":"No PubMed Summary records were recognized.", "es":"No se reconoció ningún registro PubMed Summary."},
    "Formato PubMed PMID não reconhecido.": {"en":"PubMed PMID format was not recognized.", "es":"No se reconoció el formato PubMed PMID."},
    "Formato PubMed Abstract (text) não reconhecido.": {"en":"PubMed Abstract (text) format was not recognized.", "es":"No se reconoció el formato PubMed Abstract (text)."},
    "Nenhum registro PubMed Abstract foi reconhecido.": {"en":"No PubMed Abstract records were recognized.", "es":"No se reconoció ningún registro PubMed Abstract."},
    "CSV não reconhecido como exportação do PubMed.": {"en":"CSV was not recognized as a PubMed export.", "es":"El CSV no se reconoció como exportación de PubMed."},
    "Arquivo não reconhecido como Web of Science Plain Text.": {"en":"File was not recognized as Web of Science Plain Text.", "es":"El archivo no se reconoció como Web of Science Plain Text."},
    "Nenhum registro WoS foi encontrado.": {"en":"No WoS records were found.", "es":"No se encontró ningún registro de WoS."},
    "Nenhum registro Web of Science foi encontrado na tabela.": {"en":"No Web of Science records were found in the table.", "es":"No se encontró ningún registro de Web of Science en la tabla."},
    "HTML não reconhecido como exportação do Web of Science.": {"en":"HTML was not recognized as a Web of Science export.", "es":"El HTML no se reconoció como exportación de Web of Science."},
    "Nenhum registro foi encontrado no HTML do Web of Science.": {"en":"No records were found in the Web of Science HTML file.", "es":"No se encontró ningún registro en el HTML de Web of Science."},
    "Nenhum registro BibTeX foi reconhecido.": {"en":"No BibTeX records were recognized.", "es":"No se reconoció ningún registro BibTeX."},
    "TXT não reconhecido como exportação textual do Scopus.": {"en":"TXT was not recognized as a Scopus text export.", "es":"El TXT no se reconoció como exportación textual de Scopus."},
    "Nenhum registro foi reconhecido no TXT do Scopus.": {"en":"No records were recognized in the Scopus TXT file.", "es":"No se reconoció ningún registro en el TXT de Scopus."},
    "CSV não reconhecido como exportação do Scopus.": {"en":"CSV was not recognized as a Scopus export.", "es":"El CSV no se reconoció como exportación de Scopus."},
    "Nenhum registro Scopus foi encontrado no CSV.": {"en":"No Scopus records were found in the CSV file.", "es":"No se encontró ningún registro de Scopus en el CSV."},
    "Nenhuma referência em linhas do LILACS/BVS foi reconhecida.": {"en":"No line-based LILACS/BVS references were recognized.", "es":"No se reconoció ninguna referencia por líneas de LILACS/BVS."},
    "Nenhuma referência numerada LILACS/BVS foi reconhecida.": {"en":"No numbered LILACS/BVS references were recognized.", "es":"No se reconoció ninguna referencia numerada de LILACS/BVS."},
    "Formato Cochrane 'Record #N of TOTAL' não reconhecido.": {"en":"Cochrane 'Record #N of TOTAL' format was not recognized.", "es":"No se reconoció el formato Cochrane 'Record #N of TOTAL'."},
    "Nenhum registro RIS foi reconhecido.": {"en":"No RIS records were recognized.", "es":"No se reconoció ningún registro RIS."},
    "RIS não reconhecido como exportação LILACS/BVS.": {"en":"RIS was not recognized as a LILACS/BVS export.", "es":"El RIS no se reconoció como exportación de LILACS/BVS."},
    "CSV não reconhecido como exportação LILACS/BVS.": {"en":"CSV was not recognized as a LILACS/BVS export.", "es":"El CSV no se reconoció como exportación de LILACS/BVS."},
    "Nenhum registro LILACS/BVS foi encontrado no CSV.": {"en":"No LILACS/BVS records were found in the CSV file.", "es":"No se encontró ningún registro LILACS/BVS en el CSV."},
    "Arquivo não reconhecido como exportação estruturada do Embase.": {"en":"File was not recognized as a structured Embase export.", "es":"El archivo no se reconoció como exportación estructurada de Embase."},
    "Nenhuma planilha foi encontrada no arquivo XLSX.": {"en":"No worksheet was found in the XLSX file.", "es":"No se encontró ninguna hoja en el archivo XLSX."},
    "Nenhum registro foi encontrado nas bases selecionadas.": {"en":"No records were found in the selected databases.", "es":"No se encontró ningún registro en las bases seleccionadas."},
    "Para gerar o PRISMA em PNG, instale a biblioteca Pillow (pip install pillow).": {"en":"To generate PRISMA as PNG, install the Pillow library (pip install pillow).", "es":"Para generar PRISMA en PNG, instale la biblioteca Pillow (pip install pillow)."},
    "Nenhuma referência foi reconhecida no arquivo.": {"en":"No references were recognized in the file.", "es":"No se reconoció ninguna referencia en el archivo."},
    "XML atualmente suportado para exportações do Embase.": {"en":"XML is currently supported for Embase exports.", "es":"XML es compatible actualmente con exportaciones de Embase."},
    "MS Word atualmente suportado para exportações do Embase.": {"en":"MS Word is currently supported for Embase exports.", "es":"MS Word es compatible actualmente con exportaciones de Embase."},
    "MS Excel .xlsx atualmente suportado para Embase e Web of Science.": {"en":"MS Excel .xlsx is currently supported for Embase and Web of Science.", "es":"MS Excel .xlsx es compatible actualmente con Embase y Web of Science."},
    "Excel .xls atualmente suportado para exportações do Web of Science.": {"en":"Excel .xls is currently supported for Web of Science exports.", "es":"Excel .xls es compatible actualmente con exportaciones de Web of Science."},
    "HTML atualmente suportado para exportações do Web of Science.": {"en":"HTML is currently supported for Web of Science exports.", "es":"HTML es compatible actualmente con exportaciones de Web of Science."},
    "PDF atualmente suportado para exportações do Embase.": {"en":"PDF is currently supported for Embase exports.", "es":"PDF es compatible actualmente con exportaciones de Embase."},
    "Formato não suportado:": {"en":"Unsupported format:", "es":"Formato no compatible:"},
    "TXT não reconhecido para": {"en":"TXT not recognized for", "es":"TXT no reconocido para"},
    "Formato detectado:": {"en":"Detected format:", "es":"Formato detectado:"},
}


DATABASES = [
    "Automático",
    "PubMed",
    "Scopus",
    "Web of Science",
    "Embase",
    "LILACS/BVS",
    "Cochrane Library",
    "SciELO",
    "Outra",
]

# Cores suaves usadas somente na planilha final para identificar rapidamente
# a base de origem. As colunas de classificação de revisão/meta mantêm
# sinalização própria e, portanto, não perdem sua cor informativa.
DATABASE_EXPORT_COLORS = {
    "PubMed": "DDEBF7",          # azul claro
    "Web of Science": "E4DFEC", # lilás claro
    "Scopus": "FCE4D6",          # laranja claro
    "Embase": "F4CCCC",          # rosa claro
    "LILACS/BVS": "FFF2CC",      # amarelo claro
    "Cochrane Library": "D0E0E3",# azul/cinza claro
    "SciELO": "D9EAD3",          # verde claro
    "Outra": "E7E6E6",           # cinza claro
}

FIELDS = [
    "uid",
    "base",
    "arquivo",
    "id_origem",
    "autores",
    "titulo",
    "resumo",
    "revista",
    "ano",
    "doi",
]

TRIAGE_FIELDS = FIELDS + [
    "triagem_status",
    "motivo_exclusao",
]

SECOND_SCREEN_FIELDS = FIELDS + [
    "tipo_detectado",
    "fonte_deteccao",
    "evidencia_2triagem",
    "decisao_2triagem",
    "origem_decisao_2triagem",
]

COMPARISON_FIELDS = [
    "comparacao_id",
    "status",
    "criterio",
    "similaridade",
    "evidencia",
    "uid_referencia",
    "base_referencia",
    "id_referencia",
    "autores_referencia",
    "ano_referencia",
    "titulo_referencia",
    "doi_referencia",
    "uid_comparado",
    "base_comparada",
    "id_comparado",
    "autores_comparado",
    "ano_comparado",
    "titulo_comparado",
    "doi_comparado",
    "decisao_manual",
]


# ============================================================
# NORMALIZAÇÃO
# ============================================================

def clean(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        value = "; ".join(str(v) for v in value if v is not None)
    value = str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", value).strip()


def normalize_doi(value):
    value = clean(value).lower()
    value = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", value)
    value = re.sub(r"^doi:\s*", "", value)
    m = re.search(r"(10\.\d{4,9}/[-._;()/:a-z0-9]+)", value, flags=re.I)
    if m:
        value = m.group(1)
    return value.rstrip(" .;,")


# ============================================================
# REFERÊNCIAS / DOCX / CROSSREF
# ============================================================

def extract_doi_from_text(value):
    """Extrai um DOI de uma referência textual, quando disponível."""
    text = clean(value)
    match = re.search(r"10\.\d{4,9}/[-._;()/:A-Za-z0-9]+", text, flags=re.I)
    return normalize_doi(match.group(0)) if match else ""


def _crossref_request(url, timeout=18):
    request = urllib.request.Request(
        url,
        headers={
            "User-Agent": f"BioReviewPy/{VERSION} (mailto:{AUTHOR_EMAIL}; reference formatter)",
            "Accept": "application/json",
        },
    )
    with urllib.request.urlopen(request, timeout=timeout) as response:
        payload = response.read().decode("utf-8", errors="replace")
    return json.loads(payload)


def crossref_metadata_from_reference(reference, allow_text_search=True):
    """
    Busca metadados no Crossref. DOI explícito tem prioridade.
    Sem DOI, a busca textual só é aceita com escore alto para reduzir falsos matches.
    Retorna (metadata, método, observação).
    """
    doi = extract_doi_from_text(reference)
    try:
        if doi:
            url = "https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="")
            payload = _crossref_request(url)
            message = payload.get("message", {})
            if isinstance(message, dict) and message:
                return message, "DOI", ""

        if allow_text_search:
            query = urllib.parse.urlencode({
                "query.bibliographic": clean(reference),
                "rows": 1,
                "select": "DOI,title,author,container-title,published-print,published-online,issued,volume,issue,page,publisher,type,URL",
            })
            payload = _crossref_request("https://api.crossref.org/works?" + query)
            items = payload.get("message", {}).get("items", [])
            if items:
                item = items[0]
                score = float(item.get("score", 0) or 0)
                if score >= 80:
                    return item, "BUSCA POR TEXTO", f"score={score:.1f}"
                return None, "", f"Correspondência Crossref incerta (score={score:.1f})"
    except Exception as exc:
        return None, "", f"Crossref indisponível: {exc}"

    return None, "", "DOI/metadados não localizados"


def _crossref_year(meta):
    for key in ("published-print", "published-online", "issued", "created"):
        part = meta.get(key, {}) if isinstance(meta, dict) else {}
        date_parts = part.get("date-parts", []) if isinstance(part, dict) else []
        if date_parts and date_parts[0]:
            try:
                return str(int(date_parts[0][0]))
            except Exception:
                pass
    return ""


def _crossref_first(meta, key):
    value = meta.get(key, "") if isinstance(meta, dict) else ""
    if isinstance(value, list):
        return clean(value[0]) if value else ""
    return clean(value)


def _author_initials(given):
    tokens = re.findall(r"[A-Za-zÀ-ÖØ-öø-ÿ]+", clean(given))
    return " ".join((token[0].upper() + ".") for token in tokens if token)


def _format_authors_apa(authors):
    parts = []
    for author in authors or []:
        family = clean(author.get("family", ""))
        given = clean(author.get("given", ""))
        name = family
        initials = _author_initials(given)
        if family and initials:
            name = f"{family}, {initials}"
        elif given and not family:
            name = given
        if name:
            parts.append(name)
    if not parts:
        return ""
    if len(parts) == 1:
        return parts[0]
    if len(parts) <= 20:
        return ", ".join(parts[:-1]) + ", & " + parts[-1]
    return ", ".join(parts[:19]) + ", … " + parts[-1]


def _format_authors_vancouver(authors):
    parts = []
    for author in (authors or [])[:6]:
        family = clean(author.get("family", ""))
        given = clean(author.get("given", ""))
        initials = re.sub(r"[^A-Za-zÀ-ÖØ-öø-ÿ]", "", _author_initials(given))
        name = (family + (" " + initials if initials else "")).strip()
        if name:
            parts.append(name)
    suffix = " et al" if len(authors or []) > 6 else ""
    return ", ".join(parts) + suffix


def _format_authors_abnt(authors):
    parts = []
    for author in authors or []:
        family = clean(author.get("family", ""))
        given = clean(author.get("given", ""))
        if family:
            parts.append(f"{family.upper()}, {given}".rstrip(", "))
        elif given:
            parts.append(given.upper())
    return "; ".join(parts)


def format_reference_from_crossref(meta, style):
    style_key = clean(style).upper()
    authors = meta.get("author", []) if isinstance(meta, dict) else []
    title = _crossref_first(meta, "title")
    journal = _crossref_first(meta, "container-title")
    year = _crossref_year(meta)
    volume = clean(meta.get("volume", ""))
    issue = clean(meta.get("issue", ""))
    pages = clean(meta.get("page", ""))
    publisher = clean(meta.get("publisher", ""))
    doi = normalize_doi(meta.get("DOI", ""))
    doi_url = f"https://doi.org/{doi}" if doi else clean(meta.get("URL", ""))

    if style_key.startswith("APA"):
        au = _format_authors_apa(authors)
        start = f"{au} ({year})." if au else (f"({year})." if year else "")
        pieces = [start, f"{title}." if title else ""]
        if journal:
            source = journal
            if volume:
                source += f", {volume}"
            if issue:
                source += f"({issue})"
            if pages:
                source += f", {pages}"
            source += "."
            pieces.append(source)
        elif publisher:
            pieces.append(f"{publisher}.")
        if doi_url:
            pieces.append(doi_url)
        return clean(" ".join(p for p in pieces if p))

    if style_key.startswith("VAN"):
        au = _format_authors_vancouver(authors)
        pieces = [f"{au}." if au else "", f"{title}." if title else ""]
        if journal:
            source = journal + "."
            if year:
                source += f" {year}"
            if volume:
                source += f";{volume}"
            if issue:
                source += f"({issue})"
            if pages:
                source += f":{pages}"
            source += "."
            pieces.append(source)
        else:
            if publisher:
                pieces.append(f"{publisher}; {year}." if year else f"{publisher}.")
            elif year:
                pieces.append(f"{year}.")
        if doi:
            pieces.append(f"doi: {doi}.")
        return clean(" ".join(p for p in pieces if p))

    au = _format_authors_abnt(authors)
    pieces = [(au.rstrip(".") + ".") if au else "", f"{title}." if title else ""]
    if journal:
        source = journal
        if volume:
            source += f", v. {volume}"
        if issue:
            source += f", n. {issue}"
        if pages:
            source += f", p. {pages}"
        if year:
            source += f", {year}"
        source += "."
        pieces.append(source)
    elif publisher:
        pieces.append(f"{publisher}, {year}." if year else f"{publisher}.")
    elif year:
        pieces.append(f"{year}.")
    if doi_url:
        pieces.append(f"DOI: {doi_url}.")
    return clean(" ".join(p for p in pieces if p))


def extract_reference_paragraphs(path):
    path = Path(path)
    if path.suffix.lower() == ".docx":
        with zipfile.ZipFile(path) as zf:
            if "word/document.xml" not in zf.namelist():
                raise RuntimeError("DOCX sem word/document.xml.")
            root = ET.fromstring(zf.read("word/document.xml"))
            paragraphs = []
            for paragraph in root.findall(".//{*}p"):
                text = "".join(node.text or "" for node in paragraph.findall(".//{*}t"))
                text = clean(text)
                if text:
                    paragraphs.append(text)
            return paragraphs
    if path.suffix.lower() == ".txt":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        return [clean(line) for line in text.splitlines() if clean(line)]
    raise RuntimeError("Use um arquivo .docx ou .txt com uma referência por parágrafo/linha.")


def write_simple_docx(path, title, paragraphs):
    """Gera um DOCX simples sem depender do pacote python-docx."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    def xml_text(text):
        return html.escape(str(text), quote=False)

    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    paras = []
    paras.append(
        '<w:p><w:pPr><w:spacing w:after="200"/></w:pPr><w:r><w:rPr><w:b/><w:sz w:val="28"/></w:rPr>'
        f'<w:t xml:space="preserve">{xml_text(title)}</w:t></w:r></w:p>'
    )
    for paragraph in paragraphs:
        paras.append(
            '<w:p><w:pPr><w:spacing w:after="160"/><w:jc w:val="left"/></w:pPr><w:r>'
            f'<w:t xml:space="preserve">{xml_text(paragraph)}</w:t></w:r></w:p>'
        )
    document = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>""" + "".join(paras) + """<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bottom="1440" w:left="1440"/></w:sectPr></w:body></w:document>"""

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("word/document.xml", document)
    return path


GREEK_MAP = str.maketrans({
    "α": " alpha ", "β": " beta ", "γ": " gamma ", "δ": " delta ",
    "κ": " kappa ", "λ": " lambda ", "μ": " mu ", "ω": " omega ",
    "Α": " alpha ", "Β": " beta ", "Γ": " gamma ", "Δ": " delta ",
    "Κ": " kappa ", "Λ": " lambda ", "Μ": " mu ", "Ω": " omega ",
})


def normalize_title(value):
    """Normalização forte, mas conservadora, para comparação de títulos."""
    value = html.unescape(clean(value))
    value = re.sub(r"<[^>]+>", " ", value)
    value = value.translate(GREEK_MAP)
    value = value.replace("&", " and ")
    value = value.replace("–", "-").replace("—", "-").replace("‑", "-").replace("−", "-")
    value = value.replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.casefold()
    value = re.sub(r"[\u200b-\u200d\ufeff]", "", value)
    value = re.sub(r"[_/\\:;,.!?()\[\]{}<>+*=|~`^'\"-]+", " ", value)
    value = re.sub(r"[^\w\s]", " ", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()


def normalize_author_token(value):
    value = html.unescape(clean(value))
    value = unicodedata.normalize("NFKD", value)
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = value.casefold()
    value = re.sub(r"[^a-z0-9\s,;]", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def first_author_key(value):
    text = normalize_author_token(value)
    if not text:
        return ""
    first = re.split(r";|\band\b", text, maxsplit=1)[0].strip()
    # "Surname, Name" -> surname; "Name Surname" -> last token.
    if "," in first:
        return first.split(",", 1)[0].strip()
    parts = first.split()
    if not parts:
        return ""
    # Exportações biomédicas frequentemente usam "Sobrenome Iniciais" (ex.: Cui Y).
    if len(parts) >= 2 and len(parts[-1]) <= 2:
        return parts[0]
    return parts[-1]


def build_evidence(reference, candidate, title_similarity=None):
    title_ref = normalize_title(reference.get("titulo", ""))
    title_cand = normalize_title(candidate.get("titulo", ""))
    if title_similarity is None:
        title_similarity = (
            fuzz.token_set_ratio(title_ref, title_cand)
            if title_ref and title_cand else 0
        )

    doi_ref = normalize_doi(reference.get("doi", ""))
    doi_cand = normalize_doi(candidate.get("doi", ""))
    doi_same = bool(doi_ref and doi_cand and doi_ref == doi_cand)

    year_ref = normalize_year(reference.get("ano", ""))
    year_cand = normalize_year(candidate.get("ano", ""))
    year_state = (
        "igual" if year_ref and year_cand and year_ref == year_cand
        else "diferente" if year_ref and year_cand
        else "ausente"
    )

    author_ref = first_author_key(reference.get("autores", ""))
    author_cand = first_author_key(candidate.get("autores", ""))
    author_state = (
        "igual" if author_ref and author_cand and author_ref == author_cand
        else "diferente" if author_ref and author_cand
        else "ausente"
    )

    compact = [f"T{round(float(title_similarity))}"]
    if doi_same:
        compact.insert(0, "DOI✓")
    if author_state == "igual":
        compact.append("A✓")
    elif author_state == "diferente":
        compact.append("A≠")
    if year_state == "igual":
        compact.append("Y✓")
    elif year_state == "diferente":
        compact.append("Y≠")

    return {
        "compact": " • ".join(compact),
        "title_similarity": round(float(title_similarity), 1),
        "doi_same": doi_same,
        "year_state": year_state,
        "author_state": author_state,
        "first_author_reference": author_ref,
        "first_author_candidate": author_cand,
    }


def normalize_year(value):
    m = re.search(r"\b(?:19|20)\d{2}\b", clean(value))
    return m.group(0) if m else ""


def safe_filename(name):
    name = re.sub(r'[<>:"/\\|?*]+', "_", clean(name))
    return name.strip(" ._") or "BASE"


# ============================================================
# PARSERS
# ============================================================

def _embase_first(fields, *names):
    """Retorna o primeiro valor não vazio entre campos estruturados do Embase."""
    for name in names:
        values = fields.get(name, [])
        if not isinstance(values, list):
            values = [values]
        for value in values:
            value = clean(value)
            if value:
                return value
    return ""


EMBASE_STRUCTURED_HEADINGS = {
    "TITLE",
    "AUTHOR NAMES",
    "AUTHOR",
    "AUTHOR ADDRESSES",
    "CORRESPONDENCE ADDRESS",
    "AiP/IP ENTRY DATE",
    "FULL RECORD ENTRY DATE",
    "SOURCE",
    "SOURCE TITLE",
    "PUBLICATION YEAR",
    "VOLUME",
    "ISSUE",
    "PAGES",
    "DATE OF PUBLICATION",
    "PUBLICATION TYPE",
    "ISSN",
    "ISBN",
    "BOOK PUBLISHER",
    "LANGUAGE OF ARTICLE",
    "LANGUAGE OF SUMMARY",
    "MEDLINE PMID",
    "PUI",
    "DOI",
    "ABSTRACT",
    "SUMMARY",
    "ABSTRACT/SUMMARY",
    "AUTHOR KEYWORDS",
    "EMTREE DRUG INDEX TERMS",
    "EMTREE MEDICAL INDEX TERMS",
    "EMTREE DEVICE INDEX TERMS",
}


def _clean_embase_export_text(text):
    """Remove cabeçalhos/rodapés adicionados pelos exports PDF/Word do Embase."""
    output = []
    for raw_line in str(text).replace("\r", "").split("\n"):
        line = raw_line.rstrip()
        stripped = line.strip()
        if re.match(r"^Record downloaded\s+-\s+.*\s+Page\s+\d+\s*$", stripped, flags=re.I):
            continue
        # Alguns extratores quebram o rodapé em duas partes.
        if re.fullmatch(r"Page\s+\d+", stripped, flags=re.I):
            continue
        output.append(line)
    return "\n".join(output)


def _parse_embase_structured_text(text, source_name):
    """Converte o layout vertical RECORD/TITLE usado por Plain Text, Word e PDF."""
    text = _clean_embase_export_text(text)
    parts = re.split(r"(?m)^\s*RECORD\s+(\d+)\s*$", text)

    # Fallback para um único registro quando o export não preserva a linha RECORD N.
    if len(parts) < 3:
        if re.search(r"(?m)^\s*TITLE\s*$", text) and re.search(r"(?m)^\s*SOURCE(?: TITLE)?\s*$", text):
            parts = ["", "1", text]
        else:
            raise RuntimeError("Arquivo não reconhecido como exportação estruturada do Embase.")

    records = []

    for i in range(1, len(parts), 2):
        record_number = parts[i]
        block = parts[i + 1]
        fields = defaultdict(list)
        current = None
        buffer = []

        def flush_current():
            nonlocal current, buffer
            if current is not None:
                value = clean(" ".join(x.strip() for x in buffer if x.strip()))
                if value:
                    fields[current].append(value)
            current = None
            buffer = []

        for raw_line in block.splitlines():
            line = raw_line.rstrip()
            stripped = line.strip()

            if not stripped:
                continue

            if stripped in EMBASE_STRUCTURED_HEADINGS:
                flush_current()
                current = stripped
                continue

            if current is not None:
                buffer.append(line)

        flush_current()

        title = _embase_first(fields, "TITLE")
        authors = _embase_first(fields, "AUTHOR NAMES", "AUTHOR")
        source = _embase_first(fields, "SOURCE")
        journal = _embase_first(fields, "SOURCE TITLE")
        year = _embase_first(fields, "PUBLICATION YEAR")
        doi = _embase_first(fields, "DOI")
        abstract = _embase_first(fields, "ABSTRACT", "SUMMARY", "ABSTRACT/SUMMARY")

        if not journal and source:
            source_match = re.match(r"^(.*?)\s*\(((?:19|20)\d{2})\)\s*(.*)$", source)
            if source_match:
                journal = source_match.group(1).strip()
                if not year:
                    year = source_match.group(2)
            else:
                journal = source

        if not year:
            year = normalize_year(_embase_first(fields, "DATE OF PUBLICATION") or source)

        if not doi:
            doi_match = re.search(r"(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)", source, flags=re.I)
            if doi_match:
                doi = doi_match.group(1)

        pui = _embase_first(fields, "PUI")
        medline_pmid = _embase_first(fields, "MEDLINE PMID")
        if medline_pmid:
            pmid_match = re.search(r"\b\d{5,10}\b", medline_pmid)
            medline_pmid = pmid_match.group(0) if pmid_match else medline_pmid

        record_id = pui or medline_pmid or record_number

        if title or authors or journal or doi or record_id:
            records.append({
                "base": "Embase",
                "arquivo": Path(source_name).name,
                "id_origem": record_id,
                "autores": authors,
                "titulo": title,
                "resumo": abstract,
                "revista": journal,
                "ano": year,
                "doi": doi,
            })

    if not records:
        raise RuntimeError("Nenhum registro Embase foi reconhecido.")

    return records


def parse_embase_custom(path):
    """Lê o Plain Text do Embase (layout RECORD N / campos verticais)."""
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    return _parse_embase_structured_text(text, path)


def _parse_embase_vertical_rows(rows, source_name):
    """Lê o layout vertical usado pelos exports CSV e MS Excel do Embase."""
    records_fields = []
    fields = defaultdict(list)

    def flush_record():
        nonlocal fields
        if fields:
            records_fields.append(fields)
        fields = defaultdict(list)

    for row in rows:
        values = [clean(v) for v in row]
        if not any(values):
            continue

        key = values[0]
        rest = [v for v in values[1:] if v]

        if re.fullmatch(r"RECORD\s+\d+", key, flags=re.I):
            if fields:
                flush_record()
            continue

        if key == "TITLE" and fields.get("TITLE"):
            # Em exports com vários registros, TITLE reinicia o bloco.
            flush_record()

        if key in EMBASE_STRUCTURED_HEADINGS:
            if rest:
                if key == "AUTHOR NAMES":
                    fields[key].append("; ".join(rest))
                elif key in {"ISSN", "AUTHOR ADDRESSES"}:
                    fields[key].append("; ".join(rest))
                else:
                    fields[key].append(" ".join(rest))

    flush_record()

    records = []
    for n, f in enumerate(records_fields, 1):
        title = _embase_first(f, "TITLE")
        authors = _embase_first(f, "AUTHOR NAMES", "AUTHOR")
        source = _embase_first(f, "SOURCE")
        journal = _embase_first(f, "SOURCE TITLE")
        year = _embase_first(f, "PUBLICATION YEAR") or normalize_year(_embase_first(f, "DATE OF PUBLICATION") or source)
        doi = _embase_first(f, "DOI")
        abstract = _embase_first(f, "ABSTRACT", "SUMMARY", "ABSTRACT/SUMMARY")

        if not journal and source:
            m = re.match(r"^(.*?)\s*\(((?:19|20)\d{2})\)", source)
            if m:
                journal = m.group(1).strip()
                if not year:
                    year = m.group(2)
            else:
                journal = source

        pui = _embase_first(f, "PUI")
        pmid = _embase_first(f, "MEDLINE PMID")
        pm = re.search(r"\b\d{5,10}\b", pmid) if pmid else None
        record_id = pui or (pm.group(0) if pm else pmid) or str(n)

        if title or authors or journal or doi:
            records.append({
                "base": "Embase",
                "arquivo": Path(source_name).name,
                "id_origem": record_id,
                "autores": authors,
                "titulo": title,
                "resumo": abstract,
                "revista": journal,
                "ano": year,
                "doi": doi,
            })

    if not records:
        raise RuntimeError("Nenhum registro foi reconhecido no layout vertical do Embase.")
    return records


def parse_embase_csv(path):
    """Lê o CSV vertical exportado pelo Embase."""
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        rows = list(csv.reader(f))
    return _parse_embase_vertical_rows(rows, path)


def _xlsx_column_index(cell_ref):
    letters = re.match(r"([A-Z]+)", cell_ref or "")
    if not letters:
        return 0
    value = 0
    for ch in letters.group(1):
        value = value * 26 + (ord(ch) - ord("A") + 1)
    return value - 1


def _read_xlsx_rows_basic(path):
    """Leitor XLSX simples via OOXML, evitando dependência obrigatória de openpyxl."""
    with zipfile.ZipFile(path) as zf:
        shared_strings = []
        if "xl/sharedStrings.xml" in zf.namelist():
            root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall(".//{*}si"):
                shared_strings.append("".join(t.text or "" for t in si.findall(".//{*}t")))

        sheet_path = "xl/worksheets/sheet1.xml"
        if sheet_path not in zf.namelist():
            candidates = sorted(name for name in zf.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"))
            if not candidates:
                raise RuntimeError("Nenhuma planilha foi encontrada no arquivo XLSX.")
            sheet_path = candidates[0]

        root = ET.fromstring(zf.read(sheet_path))
        rows = []
        for row_el in root.findall(".//{*}sheetData/{*}row"):
            row_values = {}
            max_col = -1
            for cell in row_el.findall("{*}c"):
                idx = _xlsx_column_index(cell.attrib.get("r", ""))
                max_col = max(max_col, idx)
                cell_type = cell.attrib.get("t", "")
                value = ""
                if cell_type == "inlineStr":
                    value = "".join(t.text or "" for t in cell.findall(".//{*}t"))
                else:
                    v = cell.find("{*}v")
                    raw = v.text if v is not None and v.text is not None else ""
                    if cell_type == "s" and raw.isdigit():
                        pos = int(raw)
                        value = shared_strings[pos] if 0 <= pos < len(shared_strings) else raw
                    else:
                        value = raw
                row_values[idx] = value
            if max_col >= 0:
                rows.append([row_values.get(i, "") for i in range(max_col + 1)])
        return rows


def parse_embase_xlsx(path):
    return _parse_embase_vertical_rows(_read_xlsx_rows_basic(path), path)


def _extract_docx_text_basic(path):
    with zipfile.ZipFile(path) as zf:
        if "word/document.xml" not in zf.namelist():
            raise RuntimeError("DOCX sem word/document.xml.")
        root = ET.fromstring(zf.read("word/document.xml"))
        paragraphs = []
        for paragraph in root.findall(".//{*}p"):
            text = "".join(node.text or "" for node in paragraph.findall(".//{*}t"))
            if text.strip():
                paragraphs.append(text)
        return "\n".join(paragraphs)


def parse_embase_docx(path):
    return _parse_embase_structured_text(_extract_docx_text_basic(path), path)


def _extract_pdf_text_basic(path, max_pages=None):
    errors = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = reader.pages if max_pages is None else reader.pages[:max_pages]
        text = "\n".join(page.extract_text() or "" for page in pages)
        if text.strip():
            return text
    except Exception as exc:
        errors.append(str(exc))

    try:
        import fitz
        doc = fitz.open(path)
        page_count = len(doc) if max_pages is None else min(len(doc), max_pages)
        text = "\n".join(doc[i].get_text("text") for i in range(page_count))
        if text.strip():
            return text
    except Exception as exc:
        errors.append(str(exc))

    raise RuntimeError(
        "Não foi possível extrair texto do PDF. Instale 'pypdf' ou use outro formato de exportação do Embase."
        + (f" Detalhes: {' | '.join(errors[:2])}" if errors else "")
    )


def parse_embase_pdf(path):
    return _parse_embase_structured_text(_extract_pdf_text_basic(path), path)


def _xml_text(element):
    if element is None:
        return ""
    return clean(" ".join(part for part in element.itertext() if part))


def parse_embase_xml(path):
    """Lê o XML nativo do Embase/Elsevier, inclusive resumo."""
    root = ET.parse(path).getroot()
    items = root.findall(".//{*}item")
    if not items and root.tag.endswith("item"):
        items = [root]

    records = []
    for n, item in enumerate(items, 1):
        bibrecord = item.find("{*}bibrecord")
        if bibrecord is None:
            bibrecord = item.find(".//{*}bibrecord")
        if bibrecord is None:
            continue

        title_el = bibrecord.find(".//{*}citation-title/{*}titletext")
        title = _xml_text(title_el)

        author_by_seq = {}
        authors_no_seq = []
        for author_el in bibrecord.findall(".//{*}author-group/{*}author"):
            name = _xml_text(author_el.find("{*}indexed-name"))
            if not name:
                surname = _xml_text(author_el.find("{*}surname"))
                given = _xml_text(author_el.find("{*}given-name"))
                name = clean(f"{surname}, {given}".strip(" ,"))
            if not name:
                continue
            seq = author_el.attrib.get("seq", "")
            if seq.isdigit():
                author_by_seq.setdefault(int(seq), name)
            elif name not in authors_no_seq:
                authors_no_seq.append(name)
        authors = [author_by_seq[k] for k in sorted(author_by_seq)] + authors_no_seq

        journal = _xml_text(bibrecord.find(".//{*}source/{*}sourcetitle"))
        publicationyear = bibrecord.find(".//{*}source/{*}publicationyear")
        year = ""
        if publicationyear is not None:
            year = publicationyear.attrib.get("first", "") or _xml_text(publicationyear)
        if not year:
            year = _xml_text(bibrecord.find(".//{*}source/{*}publicationdate/{*}year"))

        doi = _xml_text(bibrecord.find(".//{*}item-info/{*}itemidlist/{*}doi"))
        if not doi:
            # O namespace CE pode aparecer, mas wildcard cobre ambos os casos.
            doi = _xml_text(bibrecord.find(".//{*}doi"))

        pui = ""
        medline = ""
        for itemid in bibrecord.findall(".//{*}item-info/{*}itemidlist/{*}itemid"):
            idtype = itemid.attrib.get("idtype", "").upper()
            value = _xml_text(itemid)
            if idtype == "PUI" and value:
                pui = value if value.upper().startswith("L") else f"L{value}"
            elif idtype == "MEDL" and value:
                medline = value

        abstract_parts = []
        for abstract in bibrecord.findall(".//{*}abstracts/{*}abstract"):
            for para in abstract.findall(".//{*}para"):
                value = _xml_text(para)
                if value:
                    abstract_parts.append(value)
        abstract = clean(" ".join(abstract_parts))

        if title or doi or journal:
            records.append({
                "base": "Embase",
                "arquivo": Path(path).name,
                "id_origem": pui or medline or str(n),
                "autores": "; ".join(authors),
                "titulo": title,
                "resumo": abstract,
                "revista": journal,
                "ano": year,
                "doi": doi,
            })

    if not records:
        raise RuntimeError("XML não reconhecido como exportação do Embase.")
    return records

def parse_pubmed(path):
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    starts = [m.start() for m in re.finditer(r"(?m)^PMID-\s*", text)]

    if not starts:
        raise RuntimeError("Arquivo não reconhecido como PubMed NBIB/MEDLINE.")

    records = []

    for i, start in enumerate(starts):
        end = starts[i + 1] if i + 1 < len(starts) else len(text)
        block = text[start:end]

        fields = defaultdict(list)
        current = None

        for line in block.splitlines():
            m = re.match(r"^([A-Z0-9]{2,4})\s*-\s?(.*)$", line)

            if m:
                current = m.group(1)
                fields[current].append(m.group(2).strip())
            elif current and line.startswith("  "):
                fields[current][-1] += " " + line.strip()

        authors = fields.get("FAU", []) or fields.get("AU", [])

        doi = ""
        for item in fields.get("AID", []) + fields.get("LID", []):
            if "[doi]" in item.lower() or item.lower().startswith("10."):
                doi = item.split()[0]
                break

        pmid = (fields.get("PMID") or [str(i + 1)])[0]

        records.append({
            "base": "PubMed",
            "arquivo": Path(path).name,
            "id_origem": pmid,
            "autores": "; ".join(authors),
            "titulo": " ".join(fields.get("TI", [])),
            "resumo": " ".join(fields.get("AB", [])),
            "revista": (fields.get("JT") or fields.get("TA") or [""])[0],
            "ano": (fields.get("DP") or [""])[0],
            "doi": doi,
        })

    return records



def _pubmed_common_from_citation(citation):
    """Extrai campos bibliográficos básicos de uma citação textual do PubMed."""
    citation = clean(citation)

    doi_match = re.search(
        r"\bdoi:\s*(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
        citation,
        flags=re.I,
    )
    doi = doi_match.group(1) if doi_match else ""

    pmid_match = re.search(r"\bPMID:\s*(\d+)\b", citation, flags=re.I)
    pmid = pmid_match.group(1) if pmid_match else ""

    year = normalize_year(citation)

    # No formato textual do PubMed, a fonte aparece imediatamente antes do ano.
    journal = ""
    source_match = re.search(
        r"(?:^|\.\s+)([^.;]+?)\.\s+((?:19|20)\d{2})(?:\s|;|$)",
        citation,
    )
    if source_match:
        journal = clean(source_match.group(1))

    return {
        "pmid": pmid,
        "year": year,
        "journal": journal,
        "doi": doi,
    }


def parse_pubmed_summary(path):
    """Reconhece o formato PubMed 'Summary (text)'."""
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")

    # Cada registro começa por "N:". Linhas seguintes pertencem ao mesmo registro.
    parts = re.split(r"(?m)^\s*(\d+)\s*:\s*", content)
    if len(parts) < 3:
        raise RuntimeError("Formato PubMed Summary (text) não reconhecido.")

    records = []

    for i in range(1, len(parts), 2):
        sequence = parts[i]
        citation = clean(parts[i + 1])
        if not citation:
            continue

        common = _pubmed_common_from_citation(citation)

        # Estrutura típica:
        # Autores. Título. Revista. ANO ... PMID: N.
        # O bloco de autores termina no primeiro ponto seguido de espaço;
        # a fonte é localizada pelo ano, permitindo que o título contenha pontuação.
        authors = ""
        title = ""
        journal = common["journal"]

        first_dot = re.search(r"\.\s+", citation)
        if first_dot:
            authors = citation[:first_dot.start()].strip()
            after_authors = citation[first_dot.end():]

            if journal:
                journal_token = f". {journal}."
                idx = after_authors.find(journal_token)
                if idx >= 0:
                    title = after_authors[:idx].strip().rstrip(".")
                else:
                    # Fallback: primeiro ponto antes da fonte/ano.
                    source_pos = re.search(
                        rf"\.\s+{re.escape(journal)}\.\s+(?:19|20)\d{{2}}",
                        after_authors,
                    )
                    if source_pos:
                        title = after_authors[:source_pos.start()].strip().rstrip(".")
            else:
                # Fallback conservador quando a fonte não pôde ser identificada.
                next_dot = re.search(r"\.\s+", after_authors)
                if next_dot:
                    title = after_authors[:next_dot.start()].strip()

        records.append({
            "base": "PubMed",
            "arquivo": Path(path).name,
            "id_origem": common["pmid"] or sequence,
            "autores": authors,
            "titulo": title,
            "resumo": "",
            "revista": journal,
            "ano": common["year"],
            "doi": common["doi"],
        })

    if not records:
        raise RuntimeError("Nenhum registro PubMed Summary foi reconhecido.")

    return records


def parse_pubmed_pmid_only(path):
    """Reconhece exportação PubMed 'PMID': uma lista de PMIDs, um por linha."""
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    pmids = [line.strip() for line in content.splitlines() if line.strip()]

    if not pmids or any(not re.fullmatch(r"\d{5,10}", item) for item in pmids):
        raise RuntimeError("Formato PubMed PMID não reconhecido.")

    return [
        {
            "base": "PubMed",
            "arquivo": Path(path).name,
            "id_origem": pmid,
            "autores": "",
            "titulo": "",
            "resumo": "",
            "revista": "",
            "ano": "",
            "doi": "",
        }
        for pmid in pmids
    ]


def parse_pubmed_abstract_text(path):
    """Reconhece o formato PubMed 'Abstract (text)'."""
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")

    starts = list(re.finditer(r"(?m)^\s*(\d+)\.\s+", content))
    if not starts:
        raise RuntimeError("Formato PubMed Abstract (text) não reconhecido.")

    records = []

    for idx, start_match in enumerate(starts):
        start = start_match.start()
        end = starts[idx + 1].start() if idx + 1 < len(starts) else len(content)
        block = content[start:end].strip()
        sequence = start_match.group(1)

        # Parágrafos separados por linha em branco. O export padrão traz:
        # citação; título; autores; Author information; resumo; DOI/PMID etc.
        paragraphs = [
            clean(p)
            for p in re.split(r"\n\s*\n", block)
            if clean(p)
        ]
        if len(paragraphs) < 3:
            continue

        citation = re.sub(r"^\s*\d+\.\s+", "", paragraphs[0]).strip()
        title = paragraphs[1]
        authors = paragraphs[2]

        # Remove indicadores de afiliação, ex.: Cui Y(1), Hong S(1).
        authors = re.sub(r"\(\d+(?:,\s*\d+)*\)", "", authors)
        authors = re.sub(r"\s*,\s*", "; ", authors).strip(" ;")

        common = _pubmed_common_from_citation(citation)

        # PMID e DOI também aparecem no rodapé; usa-os como fonte preferencial.
        full_pmid = re.search(r"(?im)^PMID:\s*(\d+)\b", block)
        full_doi = re.search(
            r"(?im)^DOI:\s*(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
            block,
        )

        # O resumo fica, em regra, após "Author information:" e antes do rodapé.
        abstract_parts = []
        author_info_seen = False
        for paragraph in paragraphs[3:]:
            low = paragraph.casefold()
            if low.startswith("author information:"):
                author_info_seen = True
                continue
            if re.match(r"^(?:©|doi:|pmcid:|pmid:|conflict of interest)", paragraph, flags=re.I):
                break
            if author_info_seen:
                abstract_parts.append(paragraph)
        abstract_text = " ".join(abstract_parts)

        records.append({
            "base": "PubMed",
            "arquivo": Path(path).name,
            "id_origem": full_pmid.group(1) if full_pmid else (common["pmid"] or sequence),
            "autores": authors,
            "titulo": title,
            "resumo": abstract_text,
            "revista": common["journal"],
            "ano": common["year"],
            "doi": full_doi.group(1) if full_doi else common["doi"],
        })

    if not records:
        raise RuntimeError("Nenhum registro PubMed Abstract foi reconhecido.")

    return records


def parse_pubmed_csv(path):
    """Reconhece o CSV exportado diretamente pelo PubMed."""
    df = pd.read_csv(
        path,
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
        encoding_errors="replace",
    )

    columns = {str(c).strip().lower(): c for c in df.columns}
    required = {"pmid", "title", "authors"}
    if not required.issubset(columns):
        raise RuntimeError("CSV não reconhecido como exportação do PubMed.")

    def col(name):
        return columns.get(name.lower())

    records = []
    for n, row in df.iterrows():
        pmid = clean(row[col("PMID")]) if col("PMID") else str(n + 1)
        records.append({
            "base": "PubMed",
            "arquivo": Path(path).name,
            "id_origem": pmid or str(n + 1),
            "autores": row[col("Authors")] if col("Authors") else "",
            "titulo": row[col("Title")] if col("Title") else "",
            "resumo": row[col("Abstract")] if col("Abstract") else "",
            "revista": row[col("Journal/Book")] if col("Journal/Book") else "",
            "ano": row[col("Publication Year")] if col("Publication Year") else "",
            "doi": row[col("DOI")] if col("DOI") else "",
        })

    return records


def parse_wos_plain(path):
    text = Path(path).read_text(encoding="utf-8-sig", errors="replace")

    if not (
        "FN Clarivate" in text
        or re.search(r"(?m)^PT\s+", text)
        or re.search(r"(?m)^VR\s+", text)
    ):
        raise RuntimeError("Arquivo não reconhecido como Web of Science Plain Text.")

    records = []

    for n, chunk in enumerate(re.split(r"(?m)^ER\s*$", text), 1):
        if not re.search(r"(?m)^TI\s+", chunk):
            continue

        fields = defaultdict(list)
        current = None

        for line in chunk.splitlines():
            m = re.match(r"^([A-Z0-9]{2})\s+(.*)$", line)

            if m:
                current = m.group(1)
                fields[current].append(m.group(2).strip())
            elif current and line.startswith("   "):
                # Em AU/AF cada linha indentada normalmente representa outro autor;
                # nos demais campos trata-se de continuação do mesmo valor.
                if current in {"AU", "AF"}:
                    fields[current].append(line.strip())
                else:
                    fields[current][-1] += " " + line.strip()

        # AF = nomes completos no WoS, quando exportado
        authors = fields.get("AF", []) or fields.get("AU", [])

        records.append({
            "base": "Web of Science",
            "arquivo": Path(path).name,
            "id_origem": (fields.get("UT") or [str(n)])[0],
            "autores": "; ".join(authors),
            "titulo": " ".join(fields.get("TI", [])),
            "resumo": " ".join(fields.get("AB", [])),
            "revista": " ".join(fields.get("SO", [])),
            "ano": (fields.get("PY") or [""])[0],
            "doi": (fields.get("DI") or [""])[0],
        })

    if not records:
        raise RuntimeError("Nenhum registro WoS foi encontrado.")

    return records


def _records_from_wos_dataframe(df, source_name):
    """Converte tabelas exportadas pelo Web of Science para os campos internos."""
    df = df.fillna("")
    columns = {str(c).strip().casefold(): c for c in df.columns}

    def value_from(row, *names):
        for name in names:
            col = columns.get(name.casefold())
            if col is not None:
                value = clean(row[col])
                if value:
                    return value
        return ""

    records = []
    for n, row in df.iterrows():
        records.append({
            "base": "Web of Science",
            "arquivo": Path(source_name).name,
            "id_origem": value_from(
                row,
                "UT",
                "UT (Unique WOS ID)",
                "Accession Number",
                "Unique-ID",
            ) or str(n + 1),
            "autores": value_from(
                row,
                "Author Full Names",
                "AF",
                "Authors",
                "AU",
                "Author(s)",
            ),
            "titulo": value_from(row, "Article Title", "TI", "Title"),
            "resumo": value_from(row, "Abstract", "AB"),
            "revista": value_from(row, "Source Title", "SO", "Journal"),
            "ano": value_from(row, "Publication Year", "PY", "Year"),
            "doi": value_from(row, "DOI", "DI", "Digital Object Identifier"),
        })

    if not records:
        raise RuntimeError("Nenhum registro Web of Science foi encontrado na tabela.")
    return records


def parse_wos_tabular(path):
    """Lê o formato Web of Science 'Tab delimited file'."""
    df = pd.read_csv(
        path,
        sep="\t",
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
        encoding_errors="replace",
    )
    return _records_from_wos_dataframe(df, path)


def parse_wos_csv(path):
    """Lê CSV do Web of Science quando disponível."""
    df = pd.read_csv(
        path,
        sep=None,
        engine="python",
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
        encoding_errors="replace",
    )
    return _records_from_wos_dataframe(df, path)


def parse_wos_excel(path):
    """
    Lê o Excel exportado pelo Web of Science (.xls ou .xlsx).

    Arquivos .xls clássicos dependem de xlrd. Quando xlrd não está instalado,
    tenta usar LibreOffice/soffice como conversor local, se disponível.
    """
    ext = Path(path).suffix.lower()
    first_error = None

    try:
        engine = "xlrd" if ext == ".xls" else None
        df = pd.read_excel(
            path,
            dtype=str,
            keep_default_na=False,
            engine=engine,
        )
        return _records_from_wos_dataframe(df, path)
    except Exception as exc:
        first_error = exc

    # Fallback útil em máquinas com LibreOffice instalado.
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if office:
        try:
            with tempfile.TemporaryDirectory(prefix="wos_xls_") as tmpdir:
                result = subprocess.run(
                    [
                        office,
                        "--headless",
                        "--convert-to",
                        "csv",
                        "--outdir",
                        tmpdir,
                        str(path),
                    ],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    timeout=90,
                    check=False,
                )
                candidates = list(Path(tmpdir).glob("*.csv"))
                if result.returncode == 0 and candidates:
                    df = pd.read_csv(
                        candidates[0],
                        dtype=str,
                        keep_default_na=False,
                        encoding="utf-8-sig",
                        encoding_errors="replace",
                    )
                    return _records_from_wos_dataframe(df, path)
        except Exception:
            pass

    detail = clean(str(first_error))
    if ext == ".xls":
        raise RuntimeError(
            "Não foi possível ler o Excel .xls do Web of Science. "
            "Instale a dependência 'xlrd' (pip install xlrd)."
            + (f" Detalhes: {detail}" if detail else "")
        )
    raise RuntimeError(
        "Não foi possível ler o Excel do Web of Science."
        + (f" Detalhes: {detail}" if detail else "")
    )


def _html_fragment_text(value):
    value = re.sub(r"(?is)<br\s*/?>", "; ", str(value))
    value = re.sub(r"(?is)<[^>]+>", " ", value)
    value = html.unescape(value).replace("\xa0", " ")
    return clean(value)


def _wos_html_field(block, label):
    """Extrai o texto após um rótulo <b>...</b> no Printable HTML do WoS."""
    match = re.search(
        rf"(?is)<b>\s*{re.escape(label)}\s*</b>\s*(.*?)"
        rf"(?=(?:</value>)|(?:\s*(?:&nbsp;)*\s*<b>)|(?:</td>))",
        block,
    )
    if not match:
        return ""
    value = match.group(1)
    value = re.sub(r"(?is)^\s*<value>", "", value)
    return _html_fragment_text(value)


def parse_wos_html(path):
    """Lê o 'Printable HTML file' exportado pelo Web of Science."""
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")

    if not (
        "Clarivate" in content
        and "Web of Science" in content
        and re.search(r"(?is)<b>\s*Record\s+\d+\s+of\s+\d+\s*</b>", content)
    ):
        raise RuntimeError("HTML não reconhecido como exportação do Web of Science.")

    markers = list(
        re.finditer(
            r"(?is)<b>\s*Record\s+(\d+)\s+of\s+\d+\s*</b>",
            content,
        )
    )
    records = []

    for idx, marker in enumerate(markers):
        start = marker.end()
        end = markers[idx + 1].start() if idx + 1 < len(markers) else len(content)
        block = content[start:end]
        sequence = marker.group(1)

        title = _wos_html_field(block, "Title:")
        authors = _wos_html_field(block, "Author(s):")
        journal = _wos_html_field(block, "Source:")
        abstract = _wos_html_field(block, "Abstract:")
        record_id = _wos_html_field(block, "Accession Number:") or sequence
        doi = _wos_html_field(block, "DOI:")
        published = _wos_html_field(block, "Published Date:")
        year = normalize_year(published)

        # Fallback para layouts em que o ano não aparece como Published Date.
        if not year:
            year = normalize_year(block)

        if title or authors or journal or record_id:
            records.append({
                "base": "Web of Science",
                "arquivo": Path(path).name,
                "id_origem": record_id,
                "autores": authors,
                "titulo": title,
                "resumo": abstract,
                "revista": journal,
                "ano": year,
                "doi": doi,
            })

    if not records:
        raise RuntimeError("Nenhum registro foi encontrado no HTML do Web of Science.")
    return records


def strip_bib_braces(value):
    """Remove chaves externas redundantes de valores BibTeX."""
    value = clean(value)

    # Alguns exports SciELO usam {{Título}}.
    # O fallback pode capturar uma das chaves; retiramos apenas
    # chaves nas extremidades, sem alterar o texto interno.
    while value.startswith("{"):
        value = value[1:].strip()

    while value.endswith("}"):
        value = value[:-1].strip()

    return value


def parse_bibtex_fallback(path):
    """
    Leitor BibTeX básico usado como fallback.
    É suficiente para os campos usados pela ferramenta:
    author, title, journal, year e doi.
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    entries = []
    pos = 0

    while True:
        match = re.search(r"@\w+\s*\{", content[pos:], flags=re.I)
        if not match:
            break

        start = pos + match.start()
        brace_start = pos + match.end() - 1
        depth = 0
        end = None

        for j in range(brace_start, len(content)):
            char = content[j]

            if char == "{":
                depth += 1
            elif char == "}":
                depth -= 1

                if depth == 0:
                    end = j + 1
                    break

        if end is None:
            break

        block = content[start:end]
        entries.append(block)
        pos = end

    records = []

    for n, block in enumerate(entries, 1):
        head = re.match(
            r"@\w+\s*\{\s*([^,\n]+)\s*,",
            block,
            flags=re.I
        )

        entry_id = head.group(1).strip() if head else str(n)

        def field(name):
            # Valor entre chaves, incluindo chaves internas simples.
            pattern = rf"(?is)\b{name}\s*=\s*\{{(.*?)\}}\s*,?"
            match = re.search(pattern, block)

            if match:
                return strip_bib_braces(match.group(1))

            # Valor entre aspas.
            pattern2 = rf'(?is)\b{name}\s*=\s*"(.*?)"\s*,?'
            match2 = re.search(pattern2, block)

            return clean(match2.group(1)) if match2 else ""

        authors = field("author")

        # BibTeX padrão usa "and". Alguns exports SciELO usam vírgulas;
        # nesses casos mantemos a sequência original, sem inventar autores.
        authors = re.sub(
            r"\s+and\s+",
            "; ",
            authors,
            flags=re.I
        )

        unique_id = field("unique-id") or field("unique_id")
        url = field("url")
        scopus_eid = _scopus_eid_from_text(url)

        records.append({
            "base": "",
            "arquivo": Path(path).name,
            "id_origem": unique_id or scopus_eid or entry_id,
            "autores": authors,
            "titulo": field("title"),
            "resumo": field("abstract"),
            "revista": field("journal"),
            "ano": field("year"),
            "doi": field("doi"),
        })

    if not records:
        raise RuntimeError(
            "Nenhum registro BibTeX foi reconhecido."
        )

    return records



def _scopus_eid_from_text(*values):
    """Extrai EID do Scopus a partir de metadados ou URL."""
    for value in values:
        value = clean(value)
        if not value:
            continue
        m = re.search(r"\b(2-s2\.0-[A-Za-z0-9._-]+)\b", value, flags=re.I)
        if m:
            return m.group(1)
        m = re.search(r"/publications/([^?/\s]+)", value, flags=re.I)
        if m:
            raw = m.group(1)
            return raw if raw.lower().startswith("2-s2.0-") else f"2-s2.0-{raw}"
    return ""


def parse_scopus_text_export(path):
    """
    Reconhece o Plain Text exportado diretamente pelo Scopus.

    Aceita o layout com cabeçalho "Scopus / EXPORT DATE", autores,
    AUTHOR FULL NAMES, título, linha (ANO) Fonte, DOI/URL e metadados finais.
    """
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")

    if not (
        content.lstrip().startswith("Scopus")
        and "EXPORT DATE:" in content[:300]
        and "scopus.com/pages/publications/" in content
    ):
        raise RuntimeError("TXT não reconhecido como exportação textual do Scopus.")

    # Remove apenas o cabeçalho geral.
    body = re.sub(r"(?is)^\s*Scopus\s*\n\s*EXPORT DATE:.*?\n", "", content, count=1)

    # Exportações atuais terminam cada registro em EID. Quando EID não existe,
    # usamos blocos separados por linhas em branco como fallback.
    eid_markers = list(re.finditer(r"(?m)^EID:\s*.+$", body))
    blocks = []
    if eid_markers:
        cursor = 0
        for marker in eid_markers:
            block = body[cursor:marker.end()].strip()
            if block:
                blocks.append(block)
            cursor = marker.end()
        tail = body[cursor:].strip()
        if tail:
            blocks.append(tail)
    else:
        blocks = [
            block.strip()
            for block in re.split(r"\n\s*\n(?=\S)", body)
            if block.strip()
        ]

    records = []

    for block_index, block in enumerate(blocks, 1):
        lines = [clean(line) for line in block.splitlines() if clean(line)]
        if not lines:
            continue

        url_line = next(
            (line for line in lines if "scopus.com/pages/publications/" in line),
            "",
        )
        eid_line = next((line for line in lines if line.upper().startswith("EID:")), "")
        record_id = _scopus_eid_from_text(eid_line, url_line) or str(block_index)

        source_index = None
        source_line = ""
        for i, line in enumerate(lines):
            if re.match(r"^\((?:19|20)\d{2}\)\s+", line):
                source_index = i
                source_line = line
                break
        if source_index is None:
            continue

        year_match = re.match(r"^\(((?:19|20)\d{2})\)\s+(.*)$", source_line)
        year = year_match.group(1) if year_match else ""
        source_rest = year_match.group(2) if year_match else source_line
        journal = source_rest.split(",", 1)[0].strip() if source_rest else ""
        title = lines[source_index - 1] if source_index >= 1 else ""

        full_author_line = next(
            (line for line in lines if line.upper().startswith("AUTHOR FULL NAMES:")),
            "",
        )
        if full_author_line:
            author_text = full_author_line.split(":", 1)[1].strip()
            author_parts = [x.strip() for x in author_text.split(";") if x.strip()]
            authors = "; ".join(
                re.sub(r"\s*\(\d+\)\s*$", "", author).strip()
                for author in author_parts
            )
        else:
            authors = ""
            if source_index >= 2:
                potential = lines[0]
                if not re.fullmatch(r"[\d;\s]+", potential):
                    authors = potential

        doi = ""
        doi_line = next((line for line in lines if line.upper().startswith("DOI:")), "")
        if doi_line:
            m = re.search(r"(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)", doi_line, flags=re.I)
            if m:
                doi = m.group(1)

        abstract = ""
        abstract_line = next((line for line in lines if line.upper().startswith("ABSTRACT:")), "")
        if abstract_line:
            abstract = abstract_line.split(":", 1)[1].strip()

        records.append({
            "base": "Scopus",
            "arquivo": Path(path).name,
            "id_origem": record_id,
            "autores": authors,
            "titulo": title,
            "resumo": abstract,
            "revista": journal,
            "ano": year,
            "doi": doi,
        })

    if not records:
        raise RuntimeError("Nenhum registro foi reconhecido no TXT do Scopus.")
    return records


def parse_scopus_csv(path):
    """Lê o CSV exportado diretamente pelo Scopus."""
    df = pd.read_csv(
        path,
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
        encoding_errors="replace",
    )
    columns = {str(c).strip().casefold(): c for c in df.columns}

    if "title" not in columns or "source title" not in columns:
        raise RuntimeError("CSV não reconhecido como exportação do Scopus.")

    def value_from(row, *names):
        for name in names:
            col = columns.get(name.casefold())
            if col is not None:
                value = clean(row[col])
                if value:
                    return value
        return ""

    records = []
    for n, row in df.iterrows():
        authors = value_from(row, "Author full names", "Authors")
        if authors:
            authors = "; ".join(
                re.sub(r"\s*\(\d+\)\s*$", "", part.strip()).strip()
                for part in authors.split(";")
                if part.strip()
            )

        link = value_from(row, "Link", "URL")
        eid = value_from(row, "EID") or _scopus_eid_from_text(link)

        records.append({
            "base": "Scopus",
            "arquivo": Path(path).name,
            "id_origem": eid or str(n + 1),
            "autores": authors,
            "titulo": value_from(row, "Title"),
            "resumo": value_from(row, "Abstract"),
            "revista": value_from(row, "Source title", "Journal"),
            "ano": value_from(row, "Year", "Publication year"),
            "doi": value_from(row, "DOI"),
        })

    if not records:
        raise RuntimeError("Nenhum registro Scopus foi encontrado no CSV.")
    return records

def parse_lilacs_reference_lines(path):
    """
    Reconhece LILACS/BVS exportado como referências completas,
    uma referência por linha, sem numeração:

    Autores. - Título. - Fonte;volume: páginas, ANO.

    Também tolera títulos bilíngues com um ou mais ' - ' internos.
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    nonblank = [
        html.unescape(
            line.strip()
        )
        for line in content.splitlines()
        if line.strip()
    ]

    records = []

    for record_number, line in enumerate(
        nonblank,
        1
    ):
        # O delimitador usado pelo export BVS é " - ".
        parts = [
            clean(part)
            for part in line.split(" - ")
        ]

        if len(parts) < 2:
            continue

        author_part = parts[0].strip()

        # Remove ponto final que separa autores do restante.
        author_part = re.sub(
            r"\.\s*$",
            "",
            author_part
        ).strip()

        if len(parts) >= 3:
            source_part = parts[-1].strip()
            title_part = " - ".join(
                parts[1:-1]
            ).strip()
        else:
            # Fallback: autores + restante.
            title_part = parts[1].strip()
            source_part = ""

        # Título costuma terminar em ponto antes do segundo " - ".
        title_part = re.sub(
            r"\.\s*$",
            "",
            title_part
        ).strip()

        year_candidates = re.findall(
            r"\b((?:19|20)\d{2})\b",
            source_part
        )

        year = (
            year_candidates[-1]
            if year_candidates
            else ""
        )

        # Fonte/periódico = texto antes do primeiro ";".
        # Para teses pode resultar em cidade/instituição, que é preferível
        # a perder o campo completamente.
        journal = (
            source_part.split(";", 1)[0]
            .strip()
            .rstrip(".")
            if source_part
            else ""
        )

        doi_match = re.search(
            r"\b(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
            line,
            flags=re.I
        )

        doi = (
            doi_match.group(1)
            if doi_match
            else ""
        )

        records.append({
            "base": "LILACS/BVS",
            "arquivo": Path(path).name,
            "id_origem": str(
                record_number
            ),
            "autores": author_part,
            "titulo": title_part,
            "revista": journal,
            "ano": year,
            "doi": doi,
        })

    if not records:
        raise RuntimeError(
            "Nenhuma referência em linhas do LILACS/BVS foi reconhecida."
        )

    return records


def detect_database_and_format(path, selected_database="Automático"):
    """Detecta a base e o estilo de exportação. A escolha manual funciona como fallback."""
    path = Path(path)
    ext = path.suffix.lower()
    requested = clean(selected_database) or "Automático"

    # Quando o usuário escolhe uma base explicitamente, preservamos essa escolha
    # e detectamos apenas o estilo do arquivo.
    if requested != "Automático":
        if ext in {".txt", ".ciw"}:
            try:
                return requested, detect_text_format(path, requested)
            except Exception:
                return requested, "Texto"
        if ext == ".nbib":
            return requested, "PubMed PubMed/MEDLINE"
        if ext == ".ris":
            labels = {
                "Embase": "Embase RIS",
                "LILACS/BVS": "LILACS/BVS RIS",
                "Scopus": "Scopus RIS",
                "Web of Science": "Web of Science RIS",
            }
            return requested, labels.get(requested, "RIS")
        if ext == ".csv":
            labels = {
                "Embase": "Embase CSV",
                "PubMed": "PubMed CSV",
                "LILACS/BVS": "LILACS/BVS CSV",
                "Scopus": "Scopus CSV",
                "Web of Science": "Web of Science CSV",
            }
            return requested, labels.get(requested, "CSV/Tabular")
        if ext == ".xml":
            return requested, "Embase XML" if requested == "Embase" else "XML"
        if ext == ".docx":
            return requested, "Embase MS Word" if requested == "Embase" else "MS Word"
        if ext == ".xlsx":
            if requested == "Embase":
                return requested, "Embase MS Excel"
            if requested == "Web of Science":
                return requested, "Web of Science Excel"
            return requested, "MS Excel"
        if ext == ".xls":
            return requested, "Web of Science Excel (.xls)" if requested == "Web of Science" else "MS Excel (.xls)"
        if ext == ".pdf":
            return requested, "Embase PDF" if requested == "Embase" else "PDF"
        if ext in {".html", ".htm"}:
            return requested, "Web of Science Printable HTML" if requested == "Web of Science" else "HTML"
        if ext == ".bib":
            labels = {
                "Web of Science": "Web of Science BibTeX",
                "Scopus": "Scopus BibTeX",
                "SciELO": "SciELO BibTeX",
            }
            return requested, labels.get(requested, "BibTeX")
        return requested, ext.lstrip(".").upper() or "Outro"

    if ext == ".nbib":
        return "PubMed", "PubMed PubMed/MEDLINE"
    if ext == ".ciw":
        return "Web of Science", "Web of Science Plain Text"

    if ext in {".txt", ".ciw"}:
        content = path.read_text(encoding="utf-8-sig", errors="replace")
        stripped = content.lstrip()
        first_line = content.splitlines()[0] if content.splitlines() else ""
        header_cells = {clean(x).casefold() for x in first_line.split("\t") if clean(x)} if "\t" in first_line else set()

        # WoS tabulado precisa ser reconhecido antes do Plain Text, pois o cabeçalho
        # começa por "PT\tAU..." e \s também reconheceria o TAB.
        if header_cells:
            if ({"pt", "ti", "so", "ut"}.issubset(header_cells)
                    or {"article title", "source title", "ut (unique wos id)"}.issubset(header_cells)):
                return "Web of Science", "Web of Science Tab-delimited"

        if re.search(r"(?m)^PMID-\s*\d+", content):
            return "PubMed", "PubMed PubMed/MEDLINE"
        if (
            re.search(r"(?m)^\s*\d+\.\s+.+", content)
            and re.search(r"(?im)^PMID:\s*\d+", content)
            and "Author information:" in content
        ):
            return "PubMed", "PubMed Abstract (text)"
        if re.search(r"(?m)^\s*\d+\s*:\s+", content) and re.search(r"\bPMID:\s*\d+", content, flags=re.I):
            return "PubMed", "PubMed Summary (text)"
        nonblank = [line.strip() for line in content.splitlines() if line.strip()]
        if nonblank and all(re.fullmatch(r"\d{5,10}", line) for line in nonblank):
            return "PubMed", "PubMed PMID"
        if stripped.startswith("Scopus") and "EXPORT DATE:" in content[:300] and "scopus.com/pages/publications/" in content:
            return "Scopus", "Scopus Plain Text"
        if "FN Clarivate" in content or re.search(r"(?m)^VR +", content) or re.search(r"(?m)^PT +", content):
            return "Web of Science", "Web of Science Plain Text"
        if re.search(r"(?m)^\s*RECORD\s+\d+\s*$", content) and re.search(r"(?m)^\s*TITLE\s*$", content):
            return "Embase", "Embase Plain Text"
        if re.search(r"(?m)^Record #\d+ of \d+\s*$", content):
            return "Cochrane Library", "Cochrane CENTRAL TXT"
        if re.search(r"(?m)^TY\s{0,2}-\s", content):
            if re.search(r"(?m)^DB\s{0,2}-\s*Embase\s*$", content, flags=re.I) or "embase.com/" in content.casefold():
                return "Embase", "Embase RIS salvo como TXT"
            if "bvsalud.org" in content.casefold() or re.search(r"(?m)^DB\s{0,2}-\s*LILACS\s*$", content, flags=re.I):
                return "LILACS/BVS", "LILACS/BVS RIS salvo como TXT"
            if re.search(r"(?m)^DB\s{0,2}-\s*Scopus\s*$", content, flags=re.I) or "scopus.com/pages/publications/" in content.casefold():
                return "Scopus", "Scopus RIS salvo como TXT"
            if re.search(r"(?m)^AN\s{0,2}-\s*WOS:", content, flags=re.I) or "web of science core collection" in content.casefold():
                return "Web of Science", "Web of Science RIS salvo como TXT"
            return "Outra", "RIS salvo como TXT"
        if " - " in content and re.search(r"(?:19|20)\d{2}", content):
            return "LILACS/BVS", "LILACS/BVS Citação"
        if "\t" in first_line:
            return "Outra", "TXT tabulado"
        return "Outra", "TXT não identificado"

    if ext == ".ris":
        content = path.read_text(encoding="utf-8-sig", errors="replace")
        if re.search(r"(?m)^DB\s{0,2}-\s*Embase\s*$", content, flags=re.I) or "embase.com/" in content.casefold():
            return "Embase", "Embase RIS"
        if "bvsalud.org" in content.casefold() or re.search(r"(?m)^DB\s{0,2}-\s*LILACS\s*$", content, flags=re.I):
            return "LILACS/BVS", "LILACS/BVS RIS"
        if re.search(r"(?m)^DB\s{0,2}-\s*Scopus\s*$", content, flags=re.I) or "scopus.com/pages/publications/" in content.casefold() or re.search(r"(?m)^N1\s{0,2}-\s*EID:", content, flags=re.I):
            return "Scopus", "Scopus RIS"
        if re.search(r"(?m)^AN\s{0,2}-\s*WOS:", content, flags=re.I) or "web of science core collection" in content.casefold():
            return "Web of Science", "Web of Science RIS"
        return "Outra", "RIS"

    if ext == ".csv":
        # O CSV do Embase é vertical: primeira coluna = nome do campo.
        try:
            with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
                preview_rows = []
                for n, row in enumerate(csv.reader(f)):
                    preview_rows.append(row)
                    if n >= 30:
                        break
            first_col = {clean(row[0]) for row in preview_rows if row and clean(row[0])}
        except Exception:
            first_col = set()
        if {"TITLE", "AUTHOR NAMES", "SOURCE"}.issubset(first_col) and ("PUI" in first_col or "MEDLINE PMID" in first_col):
            return "Embase", "Embase CSV"

        try:
            df = pd.read_csv(path, nrows=0, dtype=str, encoding="utf-8-sig", encoding_errors="replace")
            cols = {str(c).strip().casefold() for c in df.columns}
        except Exception:
            cols = set()
        if {"pmid", "title", "authors"}.issubset(cols):
            return "PubMed", "PubMed CSV"
        if "descriptor(s)" in cols or "fulltext url" in cols or ({"id", "title", "authors", "database"}.issubset(cols)):
            return "LILACS/BVS", "LILACS/BVS CSV"
        if "eid" in cols and "source title" in cols:
            return "Scopus", "Scopus CSV"
        if (
            "ut" in cols
            or "ut (unique wos id)" in cols
            or ("accession number" in cols and ("article title" in cols or "title" in cols))
        ):
            return "Web of Science", "Web of Science CSV"
        return "Outra", "CSV/Tabular"

    if ext == ".xml":
        try:
            head = path.read_text(encoding="utf-8-sig", errors="replace")[:12000].casefold()
        except Exception:
            head = ""
        if "<bibdataset" in head and ("embase" in head or "embaselink=" in head):
            return "Embase", "Embase XML"
        return "Outra", "XML"

    if ext == ".docx":
        try:
            text = _extract_docx_text_basic(path)[:12000]
        except Exception:
            text = ""
        if re.search(r"(?m)^\s*RECORD\s+\d+\s*$", text) and re.search(r"(?m)^\s*TITLE\s*$", text) and "AUTHOR NAMES" in text:
            return "Embase", "Embase MS Word"
        return "Outra", "MS Word"

    if ext == ".xlsx":
        try:
            rows = _read_xlsx_rows_basic(path)[:30]
            first_col = {clean(row[0]) for row in rows if row and clean(row[0])}
            header = {clean(x).casefold() for x in (rows[0] if rows else []) if clean(x)}
        except Exception:
            first_col = set()
            header = set()
        if {"TITLE", "AUTHOR NAMES", "SOURCE"}.issubset(first_col) and ("PUI" in first_col or "MEDLINE PMID" in first_col):
            return "Embase", "Embase MS Excel"
        if (
            {"article title", "source title", "ut (unique wos id)"}.issubset(header)
            or {"ti", "so", "ut"}.issubset(header)
        ):
            return "Web of Science", "Web of Science Excel"
        return "Outra", "MS Excel"

    if ext == ".xls":
        try:
            raw = path.read_bytes()[:200000]
            marker_text = raw.decode("latin1", errors="ignore").casefold()
        except Exception:
            marker_text = ""
        if (
            "web of science" in marker_text
            or "ut (unique wos id)" in marker_text
            or ("article title" in marker_text and "source title" in marker_text)
        ):
            return "Web of Science", "Web of Science Excel (.xls)"
        return "Outra", "MS Excel (.xls)"

    if ext == ".pdf":
        try:
            text = _extract_pdf_text_basic(path, max_pages=2)
        except Exception:
            text = ""
        if re.search(r"(?m)^\s*RECORD\s+\d+\s*$", text) and re.search(r"(?m)^\s*TITLE\s*$", text) and "AUTHOR NAMES" in text:
            return "Embase", "Embase PDF"
        return "Outra", "PDF"

    if ext in {".html", ".htm"}:
        try:
            head = path.read_text(encoding="utf-8-sig", errors="replace")[:150000]
        except Exception:
            head = ""
        if "Clarivate" in head and "Web of Science" in head and re.search(r"(?is)<b>\s*Record\s+\d+\s+of\s+\d+", head):
            return "Web of Science", "Web of Science Printable HTML"
        return "Outra", "HTML"

    if ext == ".bib":
        try:
            head = path.read_text(encoding="utf-8-sig", errors="replace")[:20000]
            folded = head.casefold()
        except Exception:
            head = ""
            folded = ""
        if (
            re.search(r"(?i)@\w+\s*\{\s*WOS:", head)
            or "unique-id = {wos:" in folded
            or "web-of-science" in folded
        ):
            return "Web of Science", "Web of Science BibTeX"
        if (
            folded.lstrip().startswith("scopus")
            or "source = {scopus}" in folded
            or "scopus.com/pages/publications/" in folded
        ):
            return "Scopus", "Scopus BibTeX"
        if "scielo" in folded:
            return "SciELO", "SciELO BibTeX"
        return "Outra", "BibTeX"

    return "Outra", ext.lstrip(".").upper() or "Outro"


def detect_text_format(path, selected_database):
    """Detecta automaticamente variantes conhecidas de TXT."""
    content = Path(path).read_text(encoding="utf-8-sig", errors="replace")
    stripped = content.lstrip()
    first_line = content.splitlines()[0] if content.splitlines() else ""
    header_cells = {clean(x).casefold() for x in first_line.split("\t") if clean(x)} if "\t" in first_line else set()

    if selected_database == "Web of Science" and header_cells:
        if ({"pt", "ti", "so", "ut"}.issubset(header_cells)
                or {"article title", "source title", "ut (unique wos id)"}.issubset(header_cells)):
            return "Web of Science Tab-delimited"

    if (
        selected_database == "Scopus"
        and stripped.startswith("Scopus")
        and "EXPORT DATE:" in content[:300]
        and "scopus.com/pages/publications/" in content
    ):
        return "Scopus Plain Text"

    if selected_database == "LILACS/BVS" and re.search(r"(?m)^\s*\d+\.\s+", content):
        return "LILACS/BVS numerado"
    if selected_database == "LILACS/BVS" and " - " in content:
        return "LILACS/BVS Citação"
    if selected_database == "Cochrane Library" and re.search(r"(?m)^Record #\d+ of \d+\s*$", content):
        return "Cochrane CENTRAL TXT"
    if (
        selected_database == "Embase"
        and re.search(r"(?m)^RECORD\s+\d+\s*$", content)
        and "\nTITLE\n" in content
    ):
        return "Embase RECORD/TITLE TXT"

    if selected_database == "PubMed":
        if re.search(r"(?m)^PMID-\s*\d+", content):
            return "PubMed PubMed/MEDLINE"
        if re.search(r"(?m)^\s*\d+\.\s+.+", content) and re.search(r"(?im)^PMID:\s*\d+", content) and "Author information:" in content:
            return "PubMed Abstract (text)"
        if re.search(r"(?m)^\s*\d+\s*:\s+", content) and re.search(r"\bPMID:\s*\d+", content, flags=re.I):
            return "PubMed Summary (text)"
        nonblank = [line.strip() for line in content.splitlines() if line.strip()]
        if nonblank and all(re.fullmatch(r"\d{5,10}", line) for line in nonblank):
            return "PubMed PMID"

    if selected_database == "Web of Science" and (
        "FN Clarivate" in content
        or re.search(r"(?m)^PT +", content)
        or re.search(r"(?m)^VR +", content)
    ):
        return "Web of Science Plain Text"

    if re.search(r"(?m)^PMID-\s*\d+", content):
        return "PubMed PubMed/MEDLINE"
    if header_cells and ({"pt", "ti", "so", "ut"}.issubset(header_cells)
                         or {"article title", "source title", "ut (unique wos id)"}.issubset(header_cells)):
        return "Web of Science Tab-delimited"
    if "FN Clarivate" in content or re.search(r"(?m)^PT +", content) or re.search(r"(?m)^VR +", content):
        return "Web of Science Plain Text"
    if re.search(r"(?m)^TY\s{0,2}-\s", content):
        if selected_database == "Scopus" or re.search(r"(?m)^DB\s{0,2}-\s*Scopus\s*$", content, flags=re.I):
            return "Scopus RIS salvo como TXT"
        if selected_database == "Web of Science" or re.search(r"(?m)^AN\s{0,2}-\s*WOS:", content, flags=re.I):
            return "Web of Science RIS salvo como TXT"
        return "RIS salvo como TXT"
    if "\t" in first_line:
        return "TXT tabulado"
    return "TXT não identificado"

def parse_lilacs_numbered(path):
    """
    Reconhece o formato LILACS/BVS em referências numeradas,
    como o arquivo enviado pelo usuário:

    1. Autores. Título. Revista. 2023;27(2):154-159.
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    lines = [
        line.strip()
        for line in content.splitlines()
        if line.strip()
    ]

    records = []

    for line in lines:
        numbered = re.match(
            r"^\s*(\d+)\.\s+(.*)$",
            line
        )

        if not numbered:
            continue

        record_id = numbered.group(1)
        body = numbered.group(2).strip()

        # Localiza o bloco ". YEAR;" mais à direita.
        year_matches = list(
            re.finditer(
                r"\.\s+((?:19|20)\d{2});",
                body
            )
        )

        if not year_matches:
            # Ainda preserva o registro para não perder referência.
            records.append({
                "base": "LILACS/BVS",
                "arquivo": Path(path).name,
                "id_origem": record_id,
                "autores": "",
                "titulo": body,
                "revista": "",
                "ano": "",
                "doi": "",
            })
            continue

        year_match = year_matches[-1]
        year = year_match.group(1)

        before_year = body[:year_match.start()]

        segments = [
            item.strip()
            for item in re.split(
                r"\.\s+",
                before_year
            )
            if item.strip()
        ]

        authors = segments[0] if segments else ""
        journal = segments[-1] if len(segments) >= 2 else ""

        if len(segments) >= 3:
            title = ". ".join(segments[1:-1])
        elif len(segments) == 2:
            title = segments[1]
            journal = ""
        else:
            title = ""

        doi_match = re.search(
            r"\b(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
            body,
            flags=re.I
        )

        doi = doi_match.group(1) if doi_match else ""

        records.append({
            "base": "LILACS/BVS",
            "arquivo": Path(path).name,
            "id_origem": record_id,
            "autores": authors,
            "titulo": title,
            "revista": journal,
            "ano": year,
            "doi": doi,
        })

    if not records:
        raise RuntimeError(
            "Nenhuma referência numerada LILACS/BVS foi reconhecida."
        )

    return records


def parse_cochrane_custom(path):
    """
    Reconhece exportação textual da Cochrane CENTRAL:

    Record #1 of 409
    ID:
    AU:
    TI:
    SO:
    YR:
    DOI:
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    parts = re.split(
        r"(?m)^Record #(\d+) of \d+\s*$",
        content
    )

    if len(parts) < 3:
        raise RuntimeError(
            "Formato Cochrane 'Record #N of TOTAL' não reconhecido."
        )

    records = []

    for i in range(1, len(parts), 2):
        sequence_number = parts[i]
        block = parts[i + 1]

        fields = defaultdict(list)
        current = None

        for raw_line in block.splitlines():
            line = raw_line.rstrip()

            field_match = re.match(
                r"^([A-Z]{2,4}):\s?(.*)$",
                line
            )

            if field_match:
                current = field_match.group(1)
                fields[current].append(
                    field_match.group(2).strip()
                )
            elif current and line.strip():
                # Continuação do campo anterior.
                fields[current][-1] += (
                    " " + line.strip()
                )

        record_id = (
            fields.get("ID")
            or [sequence_number]
        )[0]

        records.append({
            "base": "Cochrane Library",
            "arquivo": Path(path).name,
            "id_origem": record_id,
            "autores": "; ".join(
                fields.get("AU", [])
            ),
            "titulo": " ".join(
                fields.get("TI", [])
            ),
            "revista": " ".join(
                fields.get("SO", [])
            ),
            "ano": (
                fields.get("YR")
                or [""]
            )[0],
            "doi": (
                fields.get("DOI")
                or [""]
            )[0],
        })

    return records


def _bibtex_expected_record_count(path):
    """Conta somente entradas bibliográficas reais do arquivo BibTeX.

    Entradas auxiliares como @STRING, @COMMENT e @PREAMBLE não representam
    referências e, portanto, não entram na contagem esperada.
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )
    entry_types = re.findall(
        r"(?im)^\s*@([A-Za-z][A-Za-z0-9_-]*)\s*[\{(]",
        content
    )
    ignored = {"string", "comment", "preamble"}
    return sum(1 for entry_type in entry_types if entry_type.casefold() not in ignored)


def _bibtex_merge_key(record):
    """Cria uma chave estável para reconciliar leituras BibTeX parciais."""
    record_id = clean(record.get("id_origem", "")).casefold()
    if record_id:
        return ("id", record_id)

    doi = normalize_doi(record.get("doi", ""))
    if doi:
        return ("doi", doi)

    title = normalize_title(record.get("titulo", ""))
    year = normalize_year(record.get("ano", ""))
    if title:
        return ("title", title, year)

    return None


def parse_bibtex(path):
    """Lê BibTeX sem aceitar silenciosamente uma importação parcial.

    O leitor externo ``bibtexparser`` continua sendo priorizado por preservar
    bem a estrutura dos campos. Entretanto, alguns exports do Scopus podem
    conter tipos de entrada que resultam em leitura parcial em determinadas
    versões/configurações da biblioteca. Por isso, o leitor interno também é
    executado e as duas leituras são reconciliadas. Quando o arquivo contém
    mais entradas bibliográficas do que as reconhecidas pelo leitor externo,
    os registros ausentes são recuperados pelo parser interno.
    """
    expected_count = _bibtex_expected_record_count(path)
    external_records = []

    try:
        import bibtexparser

        with open(
            path,
            "r",
            encoding="utf-8-sig",
            errors="replace"
        ) as file:
            database = bibtexparser.load(file)

        for n, entry in enumerate(
            database.entries,
            1
        ):
            authors = strip_bib_braces(
                entry.get("author", "")
            )

            authors = re.sub(
                r"\s+and\s+",
                "; ",
                authors,
                flags=re.I
            )

            entry_url = entry.get("url", "") or entry.get("URL", "")
            unique_id = (
                entry.get("unique-id", "")
                or entry.get("unique_id", "")
                or entry.get("Unique-ID", "")
            )
            scopus_eid = _scopus_eid_from_text(entry_url)

            external_records.append({
                "base": "",
                "arquivo": Path(path).name,
                "id_origem": clean(unique_id) or scopus_eid or entry.get(
                    "ID",
                    str(n)
                ),
                "autores": authors,
                "titulo": strip_bib_braces(
                    entry.get("title", "")
                ),
                "resumo": strip_bib_braces(
                    entry.get("abstract", "")
                ),
                "revista": strip_bib_braces(
                    entry.get("journal", "")
                    or entry.get(
                        "booktitle",
                        ""
                    )
                ),
                "ano": strip_bib_braces(
                    entry.get("year", "")
                    or entry.get("date", "")
                ),
                "doi": strip_bib_braces(
                    entry.get("doi", "")
                ),
            })

    except Exception:
        # A ausência/falha da biblioteca externa não impede a leitura:
        # o parser interno abaixo permanece autossuficiente.
        external_records = []

    try:
        internal_records = parse_bibtex_fallback(path)
    except Exception:
        if external_records:
            return external_records
        raise

    if not external_records:
        return internal_records

    # Se o leitor externo já reconheceu todas as entradas bibliográficas,
    # preservamos sua leitura integralmente.
    if expected_count and len(external_records) == expected_count:
        return external_records

    # Mantém os metadados lidos pelo parser externo e acrescenta apenas os
    # registros que ficaram ausentes, identificados por ID/EID, DOI ou título.
    merged = list(external_records)
    seen = set()
    for record in merged:
        key = _bibtex_merge_key(record)
        if key is not None:
            seen.add(key)

    for record in internal_records:
        key = _bibtex_merge_key(record)
        if key is None or key not in seen:
            merged.append(record)
            if key is not None:
                seen.add(key)

    # Preferência 1: reconciliação atingiu exatamente a contagem física do BIB.
    if expected_count and len(merged) == expected_count:
        return merged

    # Preferência 2: o parser interno atingiu exatamente a contagem física.
    # Isso evita aceitar silenciosamente uma leitura externa parcial.
    if expected_count and len(internal_records) == expected_count:
        return internal_records

    # Fallback conservador: nunca reduz a quantidade reconhecida.
    return merged if len(merged) > len(external_records) else external_records


def parse_ris_internal(path):
    """
    Leitor RIS interno, sem dependência externa.

    Ele é usado como fallback geral e também diretamente pelo LILACS/BVS,
    pois o RIS da BVS traz campos como ID, AU, TI, JO, PY e DO.
    """
    content = Path(path).read_text(
        encoding="utf-8-sig",
        errors="replace"
    )

    raw_records = []
    fields = defaultdict(list)
    current_tag = None

    def flush_record():
        nonlocal fields, current_tag
        if fields:
            raw_records.append(fields)
        fields = defaultdict(list)
        current_tag = None

    for raw_line in content.splitlines():
        line = raw_line.rstrip("\r\n")
        match = re.match(
            r"^([A-Z0-9]{2})\s{0,2}-\s?(.*)$",
            line
        )

        if match:
            tag = match.group(1)
            value = match.group(2).strip()

            if tag == "ER":
                flush_record()
                continue

            current_tag = tag
            fields[tag].append(value)

        elif current_tag and line.strip():
            # Continuação de campo RIS quebrado em mais de uma linha.
            fields[current_tag][-1] += " " + line.strip()

    flush_record()

    if not raw_records:
        raise RuntimeError("Nenhum registro RIS foi reconhecido.")

    records = []

    for n, entry in enumerate(raw_records, 1):
        authors = entry.get("AU", []) or entry.get("A1", [])
        title = (entry.get("TI") or entry.get("T1") or [""])[0]
        journal = (
            entry.get("JO")
            or entry.get("JF")
            or entry.get("T2")
            or [""]
        )[0]
        year = (
            entry.get("PY")
            or entry.get("Y1")
            or entry.get("DA")
            or [""]
        )[0]
        record_id_values = (
            entry.get("ID")
            or entry.get("AN")
            or entry.get("U2")
            or entry.get("C8")
            or []
        )
        record_id = record_id_values[0] if record_id_values else ""

        if not clean(record_id):
            scopus_candidates = entry.get("N1", []) + entry.get("UR", [])
            record_id = _scopus_eid_from_text(*scopus_candidates)

        if not clean(record_id):
            record_id = str(n)

        doi = (entry.get("DO") or entry.get("DI") or [""])[0]

        if not clean(doi):
            # Alguns RIS trazem o DOI apenas como URL.
            for candidate in (
                entry.get("UR", [])
                + entry.get("L1", [])
                + entry.get("L2", [])
            ):
                doi_match = re.search(
                    r"(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
                    candidate,
                    flags=re.I
                )
                if doi_match:
                    doi = doi_match.group(1)
                    break

        records.append({
            "base": "",
            "arquivo": Path(path).name,
            "id_origem": clean(record_id) or str(n),
            "autores": "; ".join(clean(a) for a in authors if clean(a)),
            "titulo": clean(title),
            "resumo": clean((entry.get("AB") or entry.get("N2") or [""])[0]),
            "revista": clean(journal),
            "ano": clean(year),
            "doi": clean(doi),
        })

    return records


def parse_ris(path):
    """Lê RIS com rispy quando disponível e usa leitor interno como fallback."""
    try:
        import rispy

        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            entries = rispy.load(f)

        records = []

        for n, entry in enumerate(entries, 1):
            authors = (
                entry.get("authors")
                or entry.get("first_authors")
                or entry.get("primary_authors")
                or []
            )

            records.append({
                "base": "",
                "arquivo": Path(path).name,
                "id_origem": clean(entry.get("id", "")) or str(n),
                "autores": "; ".join(authors) if isinstance(authors, list) else clean(authors),
                "titulo": entry.get("title", "") or entry.get("primary_title", ""),
                "resumo": entry.get("abstract", "") or entry.get("notes_abstract", ""),
                "revista": (
                    entry.get("journal_name", "")
                    or entry.get("secondary_title", "")
                    or entry.get("alternate_title3", "")
                ),
                "ano": (
                    entry.get("year", "")
                    or entry.get("publication_year", "")
                    or entry.get("date", "")
                ),
                "doi": entry.get("doi", ""),
            })

        if records:
            return records

    except Exception:
        pass

    return parse_ris_internal(path)


def parse_lilacs_ris(path):
    """Reconhece o RIS exportado diretamente do LILACS/BVS."""
    records = parse_ris_internal(path)

    if not any(
        clean(record.get("titulo", ""))
        or clean(record.get("id_origem", ""))
        for record in records
    ):
        raise RuntimeError("RIS não reconhecido como exportação LILACS/BVS.")

    return records


def parse_lilacs_csv(path):
    """Reconhece o CSV exportado diretamente do LILACS/BVS."""
    df = pd.read_csv(
        path,
        sep=None,
        engine="python",
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
        encoding_errors="replace",
    )

    columns = {str(c).strip().lower(): c for c in df.columns}

    if not {"title", "authors"}.issubset(columns):
        raise RuntimeError("CSV não reconhecido como exportação LILACS/BVS.")

    def value_from(row, *names):
        for name in names:
            column = columns.get(name.lower())
            if column is not None:
                value = clean(row[column])
                if value:
                    return value
        return ""

    records = []

    for n, row in df.iterrows():
        record_id = value_from(
            row,
            "ID",
            "Accession number",
            "PMCID"
        )

        doi = value_from(row, "DOI")
        if not doi:
            doi_url = value_from(row, "Fulltext URL")
            doi_match = re.search(
                r"(10\.\d{4,9}/[-._;()/:A-Za-z0-9]+)",
                doi_url,
                flags=re.I
            )
            if doi_match:
                doi = doi_match.group(1)

        records.append({
            "base": "LILACS/BVS",
            "arquivo": Path(path).name,
            "id_origem": record_id or str(n + 1),
            "autores": value_from(row, "Authors", "Author"),
            "titulo": value_from(row, "Title"),
            "resumo": value_from(row, "Abstract"),
            "revista": value_from(row, "Journal", "Source"),
            "ano": value_from(
                row,
                "Publication year",
                "Year",
                "Entry Date"
            ),
            "doi": doi,
        })

    if not records:
        raise RuntimeError("Nenhum registro LILACS/BVS foi encontrado no CSV.")

    return records


def parse_csv_or_tabular(path, separator=None):
    df = pd.read_csv(
        path,
        sep=separator,
        engine="python" if separator is None else "c",
        dtype=str,
        keep_default_na=False,
        encoding_errors="replace",
    )

    columns = {str(c).strip().lower(): c for c in df.columns}

    def pick(*names):
        for name in names:
            if name.lower() in columns:
                return columns[name.lower()]
        return None

    mapping = {
        "id_origem": pick(
            "eid", "ut", "accession number", "id", "record id"
        ),
        "autores": pick(
            "author full names",
            "authors full names",
            "full author names",
            "authors",
            "author(s)",
            "author",
            "af",
            "au",
        ),
        "titulo": pick(
            "title", "document title", "article title", "ti"
        ),
        "resumo": pick(
            "abstract", "abstracts", "ab", "summary"
        ),
        "revista": pick(
            "source title",
            "journal",
            "publication name",
            "source",
            "so",
        ),
        "ano": pick(
            "year", "publication year", "py"
        ),
        "doi": pick(
            "doi", "digital object identifier", "di"
        ),
    }

    records = []

    for n, row in df.iterrows():
        record = {
            "base": "",
            "arquivo": Path(path).name,
        }

        for field, col in mapping.items():
            record[field] = row[col] if col else ""

        if not clean(record.get("id_origem", "")):
            record["id_origem"] = str(n + 1)

        records.append(record)

    return records


def parse_txt_generic(path, selected_database):
    detected = detect_text_format(path, selected_database)

    if detected in {"Scopus Plain Text", "Scopus TXT (EXPORT DATE)"}:
        return parse_scopus_text_export(path)

    if detected == "Web of Science Tab-delimited":
        return parse_wos_tabular(path)

    if detected == "Embase RECORD/TITLE TXT":
        return parse_embase_custom(path)

    if detected == "Cochrane CENTRAL TXT":
        return parse_cochrane_custom(path)

    if detected == "LILACS/BVS numerado":
        return parse_lilacs_numbered(path)

    if detected in {"LILACS/BVS Citação", "LILACS/BVS referências por linha"}:
        return parse_lilacs_reference_lines(path)

    if detected == "PubMed PubMed/MEDLINE":
        return parse_pubmed(path)
    if detected == "PubMed Summary (text)":
        return parse_pubmed_summary(path)
    if detected == "PubMed Abstract (text)":
        return parse_pubmed_abstract_text(path)
    if detected == "PubMed PMID":
        return parse_pubmed_pmid_only(path)

    if detected == "Web of Science Plain Text":
        return parse_wos_plain(path)

    if detected in {
        "RIS salvo como TXT",
        "Scopus RIS salvo como TXT",
        "Web of Science RIS salvo como TXT",
    }:
        return parse_ris_internal(path)

    if detected == "TXT tabulado":
        return parse_csv_or_tabular(path, separator="\t")

    # Última tentativa: delimitador automático.
    try:
        return parse_csv_or_tabular(path, separator=None)
    except Exception:
        pass

    raise RuntimeError(
        f"TXT não reconhecido para {selected_database}. Formato detectado: {detected}."
    )


def parse_file(path, selected_database):
    if clean(selected_database) == "Automático":
        selected_database, _ = detect_database_and_format(path, "Automático")

    ext = Path(path).suffix.lower()

    if ext == ".bib":
        records = parse_bibtex(path)

    elif ext == ".ris":
        if selected_database == "LILACS/BVS":
            records = parse_lilacs_ris(path)
        elif selected_database in {"Embase", "Scopus", "Web of Science"}:
            # Leitor interno preserva IDs específicos (PUI, WOS AN e Scopus EID).
            records = parse_ris_internal(path)
        else:
            records = parse_ris(path)

    elif ext == ".nbib":
        records = parse_pubmed(path)

    elif ext == ".csv":
        if selected_database == "Embase":
            try:
                records = parse_embase_csv(path)
            except Exception:
                records = parse_csv_or_tabular(path, separator=None)
        elif selected_database == "PubMed":
            try:
                records = parse_pubmed_csv(path)
            except Exception:
                records = parse_csv_or_tabular(path, separator=None)
        elif selected_database == "LILACS/BVS":
            try:
                records = parse_lilacs_csv(path)
            except Exception:
                records = parse_csv_or_tabular(path, separator=None)
        elif selected_database == "Scopus":
            try:
                records = parse_scopus_csv(path)
            except Exception:
                records = parse_csv_or_tabular(path, separator=None)
        elif selected_database == "Web of Science":
            try:
                records = parse_wos_csv(path)
            except Exception:
                records = parse_csv_or_tabular(path, separator=None)
        else:
            records = parse_csv_or_tabular(path, separator=None)

    elif ext in {".txt", ".ciw"}:
        records = parse_txt_generic(path, selected_database)

    elif ext == ".xml":
        if selected_database == "Embase":
            records = parse_embase_xml(path)
        else:
            raise RuntimeError("XML atualmente suportado para exportações do Embase.")

    elif ext == ".docx":
        if selected_database == "Embase":
            records = parse_embase_docx(path)
        else:
            raise RuntimeError("MS Word atualmente suportado para exportações do Embase.")

    elif ext == ".xlsx":
        if selected_database == "Embase":
            records = parse_embase_xlsx(path)
        elif selected_database == "Web of Science":
            records = parse_wos_excel(path)
        else:
            raise RuntimeError("MS Excel .xlsx atualmente suportado para Embase e Web of Science.")

    elif ext == ".xls":
        if selected_database == "Web of Science":
            records = parse_wos_excel(path)
        else:
            raise RuntimeError("Excel .xls atualmente suportado para exportações do Web of Science.")

    elif ext in {".html", ".htm"}:
        if selected_database == "Web of Science":
            records = parse_wos_html(path)
        else:
            raise RuntimeError("HTML atualmente suportado para exportações do Web of Science.")

    elif ext == ".pdf":
        if selected_database == "Embase":
            records = parse_embase_pdf(path)
        else:
            raise RuntimeError("PDF atualmente suportado para exportações do Embase.")

    else:
        raise RuntimeError(f"Formato não suportado: {ext}")

    # A base escolhida/detectada sempre prevalece.
    for record in records:
        record["base"] = selected_database

    return records


# ============================================================
# LIMPEZA DO BANCO
# ============================================================

def standardize(records):
    output = []

    for index, record in enumerate(records, 1):
        x = {field: "" for field in FIELDS}

        for field in FIELDS:
            if field in record:
                x[field] = record[field]

        for field in [
            "base",
            "arquivo",
            "id_origem",
            "autores",
            "titulo",
            "resumo",
            "revista",
        ]:
            x[field] = clean(x[field])

        x["ano"] = normalize_year(x["ano"])
        x["doi"] = normalize_doi(x["doi"])

        # Identificador interno: evita remover o registro errado quando
        # diferentes arquivos usam o mesmo ID local.
        x["uid"] = (
            f"{x['base']}|{x['arquivo']}|{x['id_origem']}|{index}"
        )

        output.append(x)

    return pd.DataFrame(output, columns=FIELDS)



def _review_meta_terms(text):
    """Retorna termos explícitos de revisão sistemática/meta-análise encontrados no texto."""
    normalized = normalize_title(text)
    if not normalized:
        return []

    patterns = [
        (r"\bsystematic review\b", "systematic review"),
        (r"\bsystematic literature review\b", "systematic literature review"),
        (r"\brevisao sistematica\b", "revisão sistemática"),
        (r"\brevision sistematica\b", "revisión sistemática"),
        (r"\bmeta analysis\b", "meta-analysis"),
        (r"\bmetaanalysis\b", "meta-analysis"),
        (r"\bmeta analytic\b", "meta-analytic"),
        (r"\bmeta analise\b", "meta-análise"),
        (r"\bmetaanalise\b", "meta-análise"),
        (r"\bmetanalise\b", "meta-análise"),
        (r"\bmeta analisis\b", "meta-análisis"),
        (r"\bmetaanalisis\b", "meta-análisis"),
        (r"\bmetanalisis\b", "meta-análisis"),
        (r"\boverview of systematic reviews\b", "overview of systematic reviews"),
        (r"\bumbrella review\b", "umbrella review"),
    ]

    found = []
    for pattern, label in patterns:
        if re.search(pattern, normalized):
            if label not in found:
                found.append(label)
    return found


def classify_review_meta_record(record):
    """
    Classifica revisões sistemáticas/meta-análises para a 2ª triagem.

    A 2ª triagem é CLASSIFICATÓRIA: nenhum registro é excluído.
    - termo explícito no TÍTULO => classificação sugerida como CONFIRMADA;
    - termo explícito somente no RESUMO => sinaliza para REVISAR.

    O usuário pode confirmar/corrigir a classificação ou marcar falso positivo.
    """
    title = clean(record.get("titulo", ""))
    abstract = clean(record.get("resumo", ""))

    title_terms = _review_meta_terms(title)
    abstract_terms = _review_meta_terms(abstract)

    if title_terms:
        joined = "; ".join(title_terms)
        has_systematic = any("systematic" in x or "sistem" in x for x in title_terms)
        has_meta = any("meta" in x for x in title_terms)
        if has_systematic and has_meta:
            kind = "REVISÃO SISTEMÁTICA + META-ANÁLISE"
        elif has_meta:
            kind = "META-ANÁLISE"
        elif any("umbrella" in x or "overview" in x for x in title_terms):
            kind = "REVISÃO DE REVISÕES"
        else:
            kind = "REVISÃO SISTEMÁTICA"
        return {
            "tipo_detectado": kind,
            "fonte_deteccao": "Título",
            "evidencia_2triagem": joined,
            "decisao_automatica": "CONFIRMADO",
        }

    if abstract_terms:
        return {
            "tipo_detectado": "POSSÍVEL REVISÃO SISTEMÁTICA/META-ANÁLISE",
            "fonte_deteccao": "Resumo",
            "evidencia_2triagem": "; ".join(abstract_terms),
            "decisao_automatica": "REVISAR",
        }

    return None


def normalize_second_screen_decision(value, detected_type=""):
    """Normaliza estados da 2ª triagem e migra projetos antigos sem excluir registros."""
    decision = clean(value).upper()
    detected_type = clean(detected_type)

    # Compatibilidade com projetos v0.14 e anteriores.
    if decision == "EXCLUIR":
        return "CONFIRMADO"
    if decision == "MANTER":
        if detected_type.startswith("POSSÍVEL"):
            return "REVISAR"
        return "NÃO É REVISÃO/META"

    if decision in {"CONFIRMADO", "REVISAR", "NÃO É REVISÃO/META"}:
        return decision
    return decision


# ============================================================
# COMPARAÇÃO
# ============================================================

# Colunas auxiliares usadas somente durante a deduplicação. Elas evitam
# recalcular normalizações milhares de vezes e nunca são exportadas.
_COMPARISON_HELPER_FIELDS = [
    "_norm_doi",
    "_norm_title",
    "_norm_year",
    "_title_len",
    "_base_clean",
]


def _prepared_value(record, helper_name, source_name, normalizer):
    """Usa o valor pré-calculado quando disponível; caso contrário normaliza."""
    value = record.get(helper_name, "")
    if value != "":
        return value
    return normalizer(record.get(source_name, ""))


def exact_match(reference, candidate):
    """Correspondência exata preservando a prioridade DOI -> título por registro."""
    doi_ref = _prepared_value(reference, "_norm_doi", "doi", normalize_doi)
    doi_cand = _prepared_value(candidate, "_norm_doi", "doi", normalize_doi)

    if doi_ref and doi_cand and doi_ref == doi_cand:
        return "DOI idêntico", 100.0

    title_ref = _prepared_value(reference, "_norm_title", "titulo", normalize_title)
    title_cand = _prepared_value(candidate, "_norm_title", "titulo", normalize_title)

    if title_ref and title_cand and title_ref == title_cand:
        return "Título idêntico", 100.0

    return None, None


def best_title_candidate(candidate, references, threshold):
    """
    Encontra o melhor título semelhante usando valores pré-normalizados.

    `references` pode ser uma lista/iterador de dicionários preparados ou um
    DataFrame (compatibilidade com chamadas antigas). O RapidFuzz faz a parte
    custosa em código otimizado, enquanto Python apenas aplica os mesmos filtros
    de ano e comprimento existentes na v0.13.
    """
    title_candidate = _prepared_value(
        candidate, "_norm_title", "titulo", normalize_title
    )

    if len(title_candidate) < 15:
        return None

    cand_year = _prepared_value(candidate, "_norm_year", "ano", clean)
    cand_base = clean(candidate.get("base", ""))
    cand_len = len(title_candidate)

    if isinstance(references, pd.DataFrame):
        iterator = (row for _, row in references.iterrows())
    else:
        iterator = iter(references)

    eligible_refs = []
    eligible_titles = []

    for reference in iterator:
        if clean(reference.get("base", "")) == cand_base:
            continue

        title_ref = _prepared_value(
            reference, "_norm_title", "titulo", normalize_title
        )
        ref_len = len(title_ref)

        if ref_len < 15:
            continue

        ref_year = _prepared_value(reference, "_norm_year", "ano", clean)
        if cand_year.isdigit() and ref_year.isdigit():
            if abs(int(cand_year) - int(ref_year)) > 1:
                continue

        ratio = min(cand_len, ref_len) / max(cand_len, ref_len)
        if ratio < 0.70:
            continue

        eligible_refs.append(reference)
        eligible_titles.append(title_ref)

    if not eligible_titles:
        return None

    match = process.extractOne(
        title_candidate,
        eligible_titles,
        scorer=fuzz.token_set_ratio,
        score_cutoff=threshold,
    )

    if match is None:
        return None

    _matched_title, best_score, choice_index = match
    reference = eligible_refs[choice_index]
    ref_index = reference.get("_history_pos", choice_index)
    return ref_index, reference, round(float(best_score), 1)


def _prepare_comparison_frame(all_df):
    """Pré-calcula uma única vez os campos usados repetidamente na comparação."""
    work = all_df.copy()
    work["_norm_doi"] = work["doi"].map(normalize_doi)
    work["_norm_title"] = work["titulo"].map(normalize_title)
    work["_norm_year"] = work["ano"].map(clean)
    work["_title_len"] = work["_norm_title"].str.len()
    work["_base_clean"] = work["base"].map(clean)
    return work


_TITLE_LENGTH_BIN = 8


def _best_title_candidate_indexed(
    candidate,
    threshold,
    history_by_year_len_records,
    history_by_year_len_titles,
    history_by_len_records,
    history_by_len_titles,
):
    """Fuzzy rápido preservando exatamente os filtros de ano e comprimento."""
    title_candidate = candidate.get("_norm_title", "")
    cand_len = len(title_candidate)
    if cand_len < 15:
        return None

    # Mesma condição ratio >= 0.70 da implementação original, convertida em
    # intervalo matematicamente equivalente de comprimentos permitidos.
    min_len = max(15, math.ceil(0.70 * cand_len))
    max_len = math.floor(cand_len / 0.70)
    first_bin = min_len // _TITLE_LENGTH_BIN
    last_bin = max_len // _TITLE_LENGTH_BIN

    cand_year = candidate.get("_norm_year", "")
    if cand_year.isdigit():
        year = int(cand_year)
        year_keys = (str(year - 1), str(year), str(year + 1), "")
        sources = [
            (
                history_by_year_len_records,
                history_by_year_len_titles,
                year_key,
            )
            for year_key in year_keys
        ]
    else:
        sources = [(history_by_len_records, history_by_len_titles, None)]

    best_reference = None
    best_score = -1.0
    best_pos = 10**18

    for records_index, titles_index, year_key in sources:
        for length_bin in range(first_bin, last_bin + 1):
            key = (
                (year_key, length_bin)
                if year_key is not None
                else length_bin
            )
            records = records_index.get(key)
            if not records:
                continue
            titles = titles_index[key]

            bin_low = length_bin * _TITLE_LENGTH_BIN
            bin_high = bin_low + _TITLE_LENGTH_BIN - 1
            fully_inside = bin_low >= min_len and bin_high <= max_len

            if fully_inside:
                match = process.extractOne(
                    title_candidate,
                    titles,
                    scorer=fuzz.token_set_ratio,
                    score_cutoff=threshold,
                )
                if match is None:
                    continue
                _choice, score, local_index = match
                reference = records[local_index]
            else:
                # Apenas bins de borda precisam deste filtro curto.
                filtered_records = []
                filtered_titles = []
                for reference, title in zip(records, titles):
                    ref_len = reference.get("_title_len", len(title))
                    if min_len <= ref_len <= max_len:
                        filtered_records.append(reference)
                        filtered_titles.append(title)
                if not filtered_titles:
                    continue
                match = process.extractOne(
                    title_candidate,
                    filtered_titles,
                    scorer=fuzz.token_set_ratio,
                    score_cutoff=threshold,
                )
                if match is None:
                    continue
                _choice, score, local_index = match
                reference = filtered_records[local_index]

            score = float(score)
            history_pos = reference.get("_history_pos", 10**18)
            if (
                score > best_score
                or (score == best_score and history_pos < best_pos)
            ):
                best_score = score
                best_pos = history_pos
                best_reference = reference

    if best_reference is None:
        return None

    return best_pos, best_reference, round(best_score, 1)


def sequential_compare(all_df, reference_database=None, threshold=95, progress_callback=None):
    """
    Comparação sequencial otimizada entre bases.

    Mantém a lógica da v0.13:
    - maior base primeiro;
    - nunca compara registros da mesma base entre si;
    - duplicatas exatas não entram no banco final, mas continuam no histórico;
    - correspondências fuzzy permanecem no banco até decisão manual.

    Otimizações:
    - DOI/título/ano são normalizados uma única vez;
    - DOI e título exatos usam índices O(1), sem varrer o histórico inteiro;
    - fuzzy recebe somente anos plausíveis (±1 e registros sem ano);
    - RapidFuzz seleciona o melhor candidato em código otimizado;
    - progresso é emitido apenas quando o percentual inteiro muda.
    """
    if all_df.empty:
        return (
            pd.DataFrame(columns=FIELDS),
            pd.DataFrame(columns=COMPARISON_FIELDS),
            [],
        )

    work = _prepare_comparison_frame(all_df)

    loaded_order = list(dict.fromkeys(work["base"].tolist()))
    load_position = {db: i for i, db in enumerate(loaded_order)}
    counts_dict = work.groupby("base").size().to_dict()

    databases = sorted(
        counts_dict.keys(),
        key=lambda db: (-counts_dict[db], load_position.get(db, 999999)),
    )
    main_database = databases[0]
    comparison_order = databases

    first_base = work[work["base"] == main_database].copy().reset_index(drop=True)

    # Banco final em lista de dicionários: evita pd.concat repetido.
    consolidated_records = first_base[FIELDS].to_dict("records")

    # Histórico e índices. Os dicionários mantêm a primeira ocorrência, igual à
    # varredura sequencial antiga (primeiro registro encontrado vence).
    history_records = []
    doi_index = {}
    title_index = {}

    # Índices para fuzzy: listas de registros e títulos ficam prontas por ano
    # e faixa de comprimento, evitando recriar grandes listas a cada artigo.
    history_by_year_len_records = defaultdict(list)
    history_by_year_len_titles = defaultdict(list)
    history_by_len_records = defaultdict(list)
    history_by_len_titles = defaultdict(list)

    def add_to_history(records):
        for raw in records:
            record = dict(raw)
            record["_history_pos"] = len(history_records)
            history_records.append(record)

            doi = record.get("_norm_doi", "")
            title = record.get("_norm_title", "")
            if doi:
                doi_index.setdefault(doi, record)
            if title:
                title_index.setdefault(title, record)

            title_len = int(record.get("_title_len", 0) or 0)
            if title and title_len >= 15:
                length_bin = title_len // _TITLE_LENGTH_BIN
                year_key = record.get("_norm_year", "")
                yl_key = (year_key, length_bin)

                history_by_year_len_records[yl_key].append(record)
                history_by_year_len_titles[yl_key].append(title)
                history_by_len_records[length_bin].append(record)
                history_by_len_titles[length_bin].append(title)

    add_to_history(first_base.to_dict("records"))

    comparisons = []
    comparison_number = 1
    total_candidates = max(1, len(work) - len(first_base))
    processed_candidates = 0
    last_progress_int = -1

    def report_progress(database, force=False):
        nonlocal last_progress_int
        if progress_callback is None:
            return
        percent = (processed_candidates / total_candidates) * 100
        percent_int = int(percent)
        if force or percent_int != last_progress_int:
            last_progress_int = percent_int
            progress_callback(
                percent,
                (
                    f"{database}: {processed_candidates}/{total_candidates}"
                    f" • {percent:.0f}%"
                ),
            )

    if progress_callback is not None:
        progress_callback(0, "Iniciando comparação • 0%")

    for database in comparison_order[1:]:
        candidates_df = work[work["base"] == database].copy().reset_index(drop=True)
        candidate_records = candidates_df.to_dict("records")

        for candidate in candidate_records:
            confirmed_match = None

            # --------------------------------------------------
            # 1) Correspondência exata por índices
            # --------------------------------------------------
            exact_candidates = []
            cand_doi = candidate.get("_norm_doi", "")
            cand_title = candidate.get("_norm_title", "")

            if cand_doi and cand_doi in doi_index:
                exact_candidates.append(doi_index[cand_doi])
            if cand_title and cand_title in title_index:
                title_ref = title_index[cand_title]
                if not exact_candidates or title_ref is not exact_candidates[0]:
                    exact_candidates.append(title_ref)

            if exact_candidates:
                # Preserva a semântica da v0.13: vence o registro que apareceria
                # primeiro na varredura histórica; dentro dele DOI tem prioridade.
                reference = min(
                    exact_candidates,
                    key=lambda r: r.get("_history_pos", 10**18),
                )
                criterion, similarity = exact_match(reference, candidate)
                if criterion:
                    confirmed_match = (reference, criterion, similarity)

            if confirmed_match is not None:
                reference, criterion, similarity = confirmed_match
                evidence = build_evidence(reference, candidate, similarity)

                comparisons.append({
                    "comparacao_id": f"C{comparison_number:05d}",
                    "status": "DUPLICATA CONFIRMADA",
                    "criterio": criterion,
                    "similaridade": similarity,
                    "evidencia": evidence["compact"],
                    "uid_referencia": reference["uid"],
                    "base_referencia": reference["base"],
                    "id_referencia": reference["id_origem"],
                    "autores_referencia": reference["autores"],
                    "ano_referencia": reference["ano"],
                    "titulo_referencia": reference["titulo"],
                    "doi_referencia": reference["doi"],
                    "uid_comparado": candidate["uid"],
                    "base_comparada": candidate["base"],
                    "id_comparado": candidate["id_origem"],
                    "autores_comparado": candidate["autores"],
                    "ano_comparado": candidate["ano"],
                    "titulo_comparado": candidate["titulo"],
                    "doi_comparado": candidate["doi"],
                    "decisao_manual": "DUPLICATA",
                })
                comparison_number += 1
                processed_candidates += 1
                report_progress(database)
                continue

            # --------------------------------------------------
            # 2) Correspondência por similaridade
            # --------------------------------------------------
            possible = _best_title_candidate_indexed(
                candidate,
                threshold,
                history_by_year_len_records,
                history_by_year_len_titles,
                history_by_len_records,
                history_by_len_titles,
            )

            if possible is not None:
                _, reference, similarity = possible
                evidence = build_evidence(reference, candidate, similarity)

                comparisons.append({
                    "comparacao_id": f"C{comparison_number:05d}",
                    "status": "REVISAR",
                    "criterio": "Título semelhante",
                    "similaridade": similarity,
                    "evidencia": evidence["compact"],
                    "uid_referencia": reference["uid"],
                    "base_referencia": reference["base"],
                    "id_referencia": reference["id_origem"],
                    "autores_referencia": reference["autores"],
                    "ano_referencia": reference["ano"],
                    "titulo_referencia": reference["titulo"],
                    "doi_referencia": reference["doi"],
                    "uid_comparado": candidate["uid"],
                    "base_comparada": candidate["base"],
                    "id_comparado": candidate["id_origem"],
                    "autores_comparado": candidate["autores"],
                    "ano_comparado": candidate["ano"],
                    "titulo_comparado": candidate["titulo"],
                    "doi_comparado": candidate["doi"],
                    "decisao_manual": "",
                })
                comparison_number += 1

            # Fuzzy pendente não é excluído até decisão humana.
            consolidated_records.append({
                field: candidate.get(field, "") for field in FIELDS
            })
            processed_candidates += 1
            report_progress(database)

        # A base inteira só entra no histórico depois de ser processada,
        # preservando a regra de nunca comparar uma base com ela mesma.
        add_to_history(candidate_records)

    if progress_callback is not None:
        progress_callback(100, "Concluído • 100%")

    return (
        pd.DataFrame(consolidated_records, columns=FIELDS).reset_index(drop=True),
        pd.DataFrame(comparisons, columns=COMPARISON_FIELDS),
        comparison_order,
    )


class SimpleReviewApp(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title(f"{APP_TITLE} v{VERSION}")
        self.geometry("1440x900")
        self.minsize(1120, 700)

        self.file_entries = []
        self.base_data = {}
        self.all_df = pd.DataFrame(columns=FIELDS)
        self.base_final_df = pd.DataFrame(columns=FIELDS)
        self.final_df = pd.DataFrame(columns=FIELDS)
        self.comparison_df = pd.DataFrame(columns=COMPARISON_FIELDS)
        self.comparison_order = []
        self.triage_df = pd.DataFrame(columns=TRIAGE_FIELDS)  # legado v0.9; sem interface nesta versão
        self.review_filter_df = pd.DataFrame(columns=SECOND_SCREEN_FIELDS)
        self.second_screen_applied = False
        self.current_project_path = None
        self.audit_events = []
        self.last_biblio_merge_summary = {}

        # Idioma da interface: pode ser alterado a qualquer momento sem modificar
        # os dados internos do projeto. Português é o padrão para compatibilidade.
        self.language_code = "pt"
        self.language_var = tk.StringVar(value="Português")

        # Estado do processamento assíncrono. O worker nunca toca diretamente
        # no Tkinter; ele envia eventos por fila e a thread principal os aplica.
        self._compare_queue = queue.Queue()
        self._compare_running = False
        self._compare_generation = 0
        self._tree_fill_generation = 0

        self.prisma_extra = {
            "records_excluded_title_abstract": 0,
            "reports_not_retrieved": 0,
            "full_text_excluded": 0,
            "studies_included": 0,
            "full_text_reasons": "",
        }

        self._configure_styles()
        self._build_ui()

    def _configure_styles(self):
        """Tema azul leve e legível."""
        self.configure(bg="#EAF2F8")

        style = ttk.Style(self)

        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure(
            "TFrame",
            background="#EAF2F8"
        )
        style.configure(
            "TLabel",
            background="#EAF2F8",
            foreground="#163A5F",
            font=("Segoe UI", 9)
        )
        style.configure(
            "TLabelframe",
            background="#EAF2F8",
            bordercolor="#6FA8DC",
            relief="solid"
        )
        style.configure(
            "TLabelframe.Label",
            background="#EAF2F8",
            foreground="#0B4F8A",
            font=("Segoe UI", 10, "bold")
        )
        style.configure(
            "TButton",
            background="#1769AA",
            foreground="white",
            padding=(9, 5),
            font=("Segoe UI", 9, "bold")
        )
        style.map(
            "TButton",
            background=[
                ("active", "#0B4F8A"),
                ("pressed", "#083B66")
            ],
            foreground=[
                ("disabled", "#D9E2EC"),
                ("!disabled", "white")
            ]
        )
        style.configure(
            "TCombobox",
            fieldbackground="white",
            background="white"
        )
        style.configure(
            "TSpinbox",
            fieldbackground="white"
        )
        style.configure(
            "Treeview",
            background="white",
            fieldbackground="white",
            foreground="#1F2D3D",
            rowheight=25
        )
        style.configure(
            "Treeview.Heading",
            background="#0B4F8A",
            foreground="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat"
        )

        style.configure(
            "Blue.Horizontal.TProgressbar",
            troughcolor="#D6E9F8",
            background="#1769AA",
            bordercolor="#A9CCE3",
            lightcolor="#1769AA",
            darkcolor="#1769AA",
            thickness=16
        )
        style.map(
            "Treeview.Heading",
            background=[
                ("active", "#1769AA")
            ]
        )

    def t(self, key, **kwargs):
        """Retorna um texto da interface no idioma atualmente selecionado."""
        item = UI_TEXT.get(key, {})
        text = item.get(self.language_code) or item.get("pt") or key
        if kwargs:
            try:
                text = text.format(**kwargs)
            except Exception:
                pass
        return text

    def display_database_name(self, internal_name):
        item = DATABASE_UI_NAMES.get(internal_name, {})
        return item.get(self.language_code) or item.get("pt") or internal_name

    def internal_database_name(self, display_name):
        for internal_name, names in DATABASE_UI_NAMES.items():
            if display_name in set(names.values()):
                return internal_name
        return display_name

    def display_value(self, value):
        value = clean(value)
        item = DISPLAY_VALUE_TRANSLATIONS.get(value, {})
        return item.get(self.language_code) or item.get("pt") or value

    def _language_name_from_code(self, code):
        for name, lang_code in LANGUAGE_OPTIONS.items():
            if lang_code == code:
                return name
        return "Português"

    def _selected_database_internal(self):
        if not hasattr(self, "database_var"):
            return "Automático"
        return self.internal_database_name(self.database_var.get())

    def _populate_files_tree_from_state(self):
        if not hasattr(self, "files_tree"):
            return
        self.files_tree.delete(*self.files_tree.get_children())
        for entry in self.file_entries:
            file_path = entry.get("path", "")
            database = entry.get("database", "Outra")
            style = entry.get("style", "")
            self.files_tree.insert(
                "", "end",
                values=(self.display_database_name(database), file_path, Path(file_path).suffix.lower(), style)
            )

    def _refresh_reference_label(self):
        if not hasattr(self, "reference_var"):
            return
        counts = self.all_df.groupby("base").size().sort_values(ascending=False) if not self.all_df.empty else pd.Series(dtype=int)
        if not counts.empty:
            base = counts.index[0]
            self.reference_var.set(
                f"{self.display_database_name(base)} ({int(counts.iloc[0])} {self.t('records')})"
            )
        else:
            self.reference_var.set("-")

    def change_language(self, *_args):
        """Reconstrói a interface principal no idioma escolhido preservando o projeto."""
        selected_name = self.language_var.get()
        new_code = LANGUAGE_OPTIONS.get(selected_name, "pt")
        if new_code == self.language_code:
            return

        selected_database = self._selected_database_internal()
        threshold = int(self.threshold_var.get()) if hasattr(self, "threshold_var") else 95
        progress = float(self.progress_var.get()) if hasattr(self, "progress_var") else 0
        progress_text = self.progress_label_var.get() if hasattr(self, "progress_label_var") else ""

        self.language_code = new_code

        # Cancela apenas preenchimentos visuais em andamento; os dados permanecem.
        self._tree_fill_generation += 1
        for child in list(self.winfo_children()):
            child.destroy()
        self.config(menu=tk.Menu(self))
        self._build_ui()

        self.threshold_var.set(threshold)
        self.database_var.set(self.display_database_name(selected_database))
        self._populate_files_tree_from_state()
        self._refresh_reference_label()
        self.progress_var.set(progress)
        if progress <= 0:
            self.progress_label_var.set(self.t("waiting_processing"))
        elif progress >= 100:
            self.progress_label_var.set({"pt":"Concluído • 100%", "en":"Completed • 100%", "es":"Completado • 100%"}.get(self.language_code, progress_text))
        else:
            self.progress_label_var.set(self.translate_literal(progress_text))
        self.fill_comparison_tree()
        self.update_summary()

    def translate_literal(self, text):
        """Traduz texto estático ou dinâmico mostrado ao usuário."""
        if text is None or self.language_code == "pt":
            return text

        original = str(text)

        # 1) Correspondência exata das janelas secundárias já catalogadas.
        item = LITERAL_UI_TEXT.get(original, {})
        if self.language_code in item:
            return item[self.language_code]

        # 2) Correspondência exata/por trechos para mensagens dinâmicas,
        #    mantendo números, caminhos, nomes de arquivos e detalhes técnicos.
        translated = original
        for source in sorted(MESSAGE_UI_TEXT, key=len, reverse=True):
            target = MESSAGE_UI_TEXT[source].get(self.language_code)
            if target and source in translated:
                translated = translated.replace(source, target)

        # 3) Estados internos que aparecem embutidos em mensagens.
        for source, translations in DISPLAY_VALUE_TRANSLATIONS.items():
            target = translations.get(self.language_code)
            if target and source in translated:
                translated = translated.replace(source, target)

        return translated

    # Wrappers: qualquer aviso/caixa aberta por uma ação respeita o idioma atual.
    def _showinfo(self, title, message, **kwargs):
        return messagebox.showinfo(
            self.translate_literal(title), self.translate_literal(message), **kwargs
        )

    def _showwarning(self, title, message, **kwargs):
        return messagebox.showwarning(
            self.translate_literal(title), self.translate_literal(message), **kwargs
        )

    def _showerror(self, title, message, **kwargs):
        return messagebox.showerror(
            self.translate_literal(title), self.translate_literal(message), **kwargs
        )

    def _askyesno(self, title, message, **kwargs):
        return messagebox.askyesno(
            self.translate_literal(title), self.translate_literal(message), **kwargs
        )

    def _askstring(self, title, prompt, **kwargs):
        return simpledialog.askstring(
            self.translate_literal(title), self.translate_literal(prompt), **kwargs
        )

    def _localized_filetypes(self, filetypes):
        if not filetypes or self.language_code == "pt":
            return filetypes
        return [
            (self.translate_literal(label), pattern)
            for label, pattern in filetypes
        ]

    def _askopenfilenames(self, **kwargs):
        if "title" in kwargs:
            kwargs["title"] = self.translate_literal(kwargs["title"])
        if "filetypes" in kwargs:
            kwargs["filetypes"] = self._localized_filetypes(kwargs["filetypes"])
        return filedialog.askopenfilenames(**kwargs)

    def _askopenfilename(self, **kwargs):
        if "title" in kwargs:
            kwargs["title"] = self.translate_literal(kwargs["title"])
        if "filetypes" in kwargs:
            kwargs["filetypes"] = self._localized_filetypes(kwargs["filetypes"])
        return filedialog.askopenfilename(**kwargs)

    def _askdirectory(self, **kwargs):
        if "title" in kwargs:
            kwargs["title"] = self.translate_literal(kwargs["title"])
        return filedialog.askdirectory(**kwargs)

    def _asksaveasfilename(self, **kwargs):
        if "title" in kwargs:
            kwargs["title"] = self.translate_literal(kwargs["title"])
        if "filetypes" in kwargs:
            kwargs["filetypes"] = self._localized_filetypes(kwargs["filetypes"])
        return filedialog.asksaveasfilename(**kwargs)

    def _localize_widget_tree(self, widget):
        """Traduz textos estáticos de janelas secundárias após sua construção."""
        try:
            if isinstance(widget, tk.Toplevel):
                widget.title(self.translate_literal(widget.title()))
        except Exception:
            pass
        try:
            current = widget.cget("text")
            translated = self.translate_literal(current)
            if translated != current:
                widget.configure(text=translated)
        except Exception:
            pass
        try:
            if isinstance(widget, ttk.Treeview):
                for column in widget["columns"]:
                    heading_text = widget.heading(column, "text")
                    translated_heading = self.translate_literal(heading_text)
                    if translated_heading != heading_text:
                        widget.heading(column, text=translated_heading)
        except Exception:
            pass
        for child in widget.winfo_children():
            self._localize_widget_tree(child)

    def open_about_dialog(self):
        """Exibe autoria, licença, versão e orientação de citação do software."""
        if self.language_code == "en":
            message = (
                f"{APP_TITLE} v{VERSION}\n\n"
                f"Developed by {AUTHOR_NAME}\n"
                f"E-mail: {AUTHOR_EMAIL}\n"
                f"{COPYRIGHT_NOTICE}\n"
                f"License: {LICENSE_NAME}\n\n"
                "BioReviewPy is a semi-automated and auditable tool designed to support "
                "bibliographic data preparation and evidence-synthesis workflows.\n\n"
                f"Repository: {PROJECT_URL}\n\n"
                "If you use BioReviewPy in research, please cite the software according "
                "to the CITATION.cff file and the DOI of the corresponding release."
            )
            title = "About BioReviewPy"
        elif self.language_code == "es":
            message = (
                f"{APP_TITLE} v{VERSION}\n\n"
                f"Desarrollado por {AUTHOR_NAME}\n"
                f"Correo electrónico: {AUTHOR_EMAIL}\n"
                f"{COPYRIGHT_NOTICE}\n"
                f"Licencia: {LICENSE_NAME}\n\n"
                "BioReviewPy es una herramienta semiautomatizada y auditable para apoyar "
                "la preparación de datos bibliográficos y los flujos de síntesis de evidencia.\n\n"
                f"Repositorio: {PROJECT_URL}\n\n"
                "Si utiliza BioReviewPy en investigación, cite el software de acuerdo con "
                "el archivo CITATION.cff y el DOI de la versión correspondiente."
            )
            title = "Acerca de BioReviewPy"
        else:
            message = (
                f"{APP_TITLE} v{VERSION}\n\n"
                f"Desenvolvido por {AUTHOR_NAME}\n"
                f"E-mail: {AUTHOR_EMAIL}\n"
                f"{COPYRIGHT_NOTICE}\n"
                f"Licença: {LICENSE_NAME}\n\n"
                "BioReviewPy é uma ferramenta semiautomatizada e auditável desenvolvida para "
                "apoiar a preparação de dados bibliográficos e fluxos de síntese de evidências.\n\n"
                f"Repositório: {PROJECT_URL}\n\n"
                "Ao utilizar o BioReviewPy em pesquisas, cite o software conforme o arquivo "
                "CITATION.cff e o DOI da versão correspondente."
            )
            title = "Sobre o BioReviewPy"
        messagebox.showinfo(title, message, parent=self)

    def _build_ui(self):
        # Menu discreto: evita poluir a grade principal.
        menubar = tk.Menu(self)
        project_menu = tk.Menu(menubar, tearoff=0)
        project_menu.add_command(label=self.t("new_project"), command=self.new_project)
        project_menu.add_separator()
        project_menu.add_command(label=self.t("open_project"), command=self.open_project)
        project_menu.add_command(label=self.t("save_project"), command=self.save_project)
        project_menu.add_command(label=self.t("save_project_as"), command=lambda: self.save_project(save_as=True))
        menubar.add_cascade(label=self.t("project"), menu=project_menu)

        tools_menu = tk.Menu(menubar, tearoff=0)
        tools_menu.add_command(label=self.t("dashboard"), command=self.open_dashboard_window)
        tools_menu.add_separator()
        tools_menu.add_command(label=self.t("second_screen"), command=self.open_second_screen_window)
        tools_menu.add_command(label=self.t("search_article"), command=self.open_search_window)
        tools_menu.add_command(label=self.t("format_refs"), command=self.open_reference_formatter_window)
        tools_menu.add_separator()
        tools_menu.add_command(label=self.t("prisma_auto"), command=self.open_prisma_window)
        tools_menu.add_command(label=self.t("audit_report"), command=self.export_audit_dialog)
        menubar.add_cascade(label=self.t("tools"), menu=tools_menu)

        biblio_menu = tk.Menu(menubar, tearoff=0)
        biblio_menu.add_command(
            label=self.t("biblio_merge"),
            command=self.open_bibliometric_merge_window,
        )
        menubar.add_cascade(label=self.t("bibliometrics"), menu=biblio_menu)

        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label=self.t("about"), command=self.open_about_dialog)
        menubar.add_cascade(label=self.t("help"), menu=help_menu)
        self.config(menu=menubar)

        # Header azul
        header = tk.Frame(
            self,
            bg="#0B4F8A",
            padx=14,
            pady=10
        )
        header.pack(fill="x")

        language_box = tk.Frame(header, bg="#0B4F8A")
        language_box.pack(side="right", padx=(10, 0), anchor="n")
        tk.Label(
            language_box,
            text=self.t("language"),
            bg="#0B4F8A",
            fg="white",
            font=("Segoe UI", 9, "bold")
        ).pack(side="left", padx=(0, 6))
        language_combo = ttk.Combobox(
            language_box,
            textvariable=self.language_var,
            values=list(LANGUAGE_OPTIONS.keys()),
            state="readonly",
            width=11,
        )
        language_combo.pack(side="left")
        language_combo.bind("<<ComboboxSelected>>", self.change_language)

        tk.Label(
            header,
            text=f"{APP_TITLE}   v{VERSION}",
            bg="#0B4F8A",
            fg="white",
            font=("Segoe UI", 18, "bold")
        ).pack()

        tk.Label(
            header,
            text=self.t("subtitle"),
            bg="#0B4F8A",
            fg="#DDEEFF",
            font=("Segoe UI", 9)
        ).pack(pady=(2, 0))

        # Créditos centralizados
        credit_banner = tk.Frame(
            self,
            bg="#D6E9F8",
            pady=5
        )
        credit_banner.pack(
            fill="x",
            padx=10,
            pady=(6, 8)
        )

        tk.Label(
            credit_banner,
            text=self.t("credits"),
            bg="#D6E9F8",
            fg="#0B4F8A",
            font=("Segoe UI", 9, "bold")
        ).pack()

        # Arquivos
        files_box = ttk.LabelFrame(
            self,
            text=self.t("load_bases"),
            padding=8
        )
        files_box.pack(fill="x", padx=10, pady=(0, 8))

        ttk.Label(
            files_box,
            text=self.t("database")
        ).grid(row=0, column=0, padx=(0, 4), sticky="w")

        self.database_var = tk.StringVar(value=self.display_database_name("Automático"))

        self.database_combo = ttk.Combobox(
            files_box,
            textvariable=self.database_var,
            values=[self.display_database_name(db) for db in DATABASES],
            state="readonly",
            width=23,
        )
        self.database_combo.grid(row=0, column=1, padx=4)

        ttk.Button(
            files_box,
            text=self.t("add_files"),
            command=self.add_files
        ).grid(row=0, column=2, padx=8)

        ttk.Button(
            files_box,
            text=self.t("remove_selected"),
            command=self.remove_selected_files
        ).grid(row=0, column=3, padx=4)

        ttk.Button(
            files_box,
            text=self.t("clear_all"),
            command=self.clear_files
        ).grid(row=0, column=4, padx=4)

        ttk.Button(
            files_box,
            text=self.t("process"),
            command=self.process_files
        ).grid(row=0, column=5, padx=(25, 4))

        ttk.Label(
            files_box,
            text=self.t("auto_detection_note"),
            wraplength=1320,
            justify="left",
            font=("Segoe UI", 8)
        ).grid(row=1, column=0, columnspan=6, sticky="w", pady=(7, 0))

        # Lista de arquivos
        list_frame = ttk.Frame(self)
        list_frame.pack(fill="x", padx=10, pady=(0, 8))

        self.files_tree = ttk.Treeview(
            list_frame,
            columns=["base", "arquivo", "formato", "estilo"],
            show="headings",
            height=6,
            selectmode="extended"
        )

        self.files_tree.heading("base", text=self.t("database").rstrip(":"))
        self.files_tree.heading("arquivo", text=self.t("file"))
        self.files_tree.heading("formato", text=self.t("extension"))
        self.files_tree.heading("estilo", text=self.t("detected_format"))

        self.files_tree.column("base", width=165)
        self.files_tree.column("arquivo", width=650)
        self.files_tree.column("formato", width=75)
        self.files_tree.column("estilo", width=260)

        y_files = ttk.Scrollbar(
            list_frame,
            orient="vertical",
            command=self.files_tree.yview
        )

        self.files_tree.configure(yscrollcommand=y_files.set)

        self.files_tree.pack(side="left", fill="x", expand=True)
        y_files.pack(side="right", fill="y")

        # Opções comparação
        compare_box = ttk.LabelFrame(
            self,
            text=self.t("comparison"),
            padding=(12, 9)
        )
        compare_box.pack(
            fill="x",
            padx=10,
            pady=(0, 8)
        )

        # Linha superior: controles bem espaçados.
        controls_row = ttk.Frame(
            compare_box
        )
        controls_row.pack(
            fill="x"
        )

        ttk.Label(
            controls_row,
            text=self.t("auto_reference")
        ).pack(
            side="left"
        )

        self.reference_var = tk.StringVar(
            value="-"
        )

        ttk.Label(
            controls_row,
            textvariable=self.reference_var,
            font=("Segoe UI", 10, "bold")
        ).pack(
            side="left",
            padx=(6, 36)
        )

        ttk.Label(
            controls_row,
            text=self.t("review_similarity")
        ).pack(
            side="left",
            padx=(0, 6)
        )

        self.threshold_var = tk.IntVar(
            value=95
        )

        ttk.Spinbox(
            controls_row,
            from_=85,
            to=99,
            textvariable=self.threshold_var,
            width=5
        ).pack(
            side="left"
        )

        ttk.Label(
            controls_row,
            text="%"
        ).pack(
            side="left",
            padx=(3, 34)
        )

        # Espaço elástico antes do botão para evitar aparência "estourada".
        ttk.Frame(
            controls_row
        ).pack(
            side="left",
            fill="x",
            expand=True
        )

        ttk.Button(
            controls_row,
            text=self.t("update_comparison"),
            command=self.recompare
        ).pack(
            side="right",
            padx=(30, 4)
        )

        # Linha intermediária: explicação.
        info_row = ttk.Frame(
            compare_box
        )
        info_row.pack(
            fill="x",
            pady=(7, 5)
        )

        ttk.Label(
            info_row,
            text=(
                self.t("comparison_note")
            )
        ).pack(
            side="left"
        )

        # Linha inferior: barra de progresso.
        progress_row = ttk.Frame(
            compare_box
        )
        progress_row.pack(
            fill="x",
            pady=(2, 0)
        )

        self.progress_var = tk.DoubleVar(
            value=0
        )

        self.progress_label_var = tk.StringVar(
            value=self.t("waiting_processing")
        )

        self.progress_bar = ttk.Progressbar(
            progress_row,
            variable=self.progress_var,
            maximum=100,
            mode="determinate",
            style="Blue.Horizontal.TProgressbar"
        )
        self.progress_bar.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 12)
        )

        ttk.Label(
            progress_row,
            textvariable=self.progress_label_var,
            width=27,
            anchor="e",
            font=("Segoe UI", 9, "bold")
        ).pack(
            side="right"
        )

        # Resumo
        summary_box = ttk.LabelFrame(
            self,
            text=self.t("summary"),
            padding=6
        )
        summary_box.pack(fill="x", padx=10, pady=(0, 8))

        self.summary_tree = ttk.Treeview(
            summary_box,
            columns=[
                "base",
                "bruto",
                "duplicatas",
                "pendentes",
                "final"
            ],
            show="headings",
            height=7
        )

        self.summary_tree.heading(
            "base",
            text=self.t("database").rstrip(":")
        )
        self.summary_tree.heading(
            "bruto",
            text=self.t("raw")
        )
        self.summary_tree.heading(
            "duplicatas",
            text=self.t("duplicates_removed")
        )
        self.summary_tree.heading(
            "pendentes",
            text=self.t("pending")
        )
        self.summary_tree.heading(
            "final",
            text=self.t("final")
        )

        self.summary_tree.column(
            "base",
            width=230
        )
        self.summary_tree.column(
            "bruto",
            width=90,
            anchor="center"
        )
        self.summary_tree.column(
            "duplicatas",
            width=155,
            anchor="center"
        )
        self.summary_tree.column(
            "pendentes",
            width=105,
            anchor="center"
        )
        self.summary_tree.column(
            "final",
            width=90,
            anchor="center"
        )

        self.summary_tree.pack(
            side="left",
            fill="x",
            expand=False
        )

        self.summary_text = tk.StringVar(value=self.t("no_files_processed"))

        ttk.Label(
            summary_box,
            textvariable=self.summary_text,
            justify="left",
            font=("Segoe UI", 10, "bold")
        ).pack(side="left", padx=30)

        # Comparações
        duplicates_box = ttk.LabelFrame(
            self,
            text=self.t("duplicate_section"),
            padding=6
        )
        duplicates_box.pack(fill="both", expand=True, padx=10, pady=(0, 8))

        buttons = ttk.Frame(duplicates_box)
        buttons.pack(fill="x", pady=(0, 5))

        ttk.Button(
            buttons,
            text=self.t("confirm_duplicate"),
            command=lambda: self.set_manual_decision("DUPLICATA")
        ).pack(side="left", padx=4)

        ttk.Button(
            buttons,
            text=self.t("not_duplicate"),
            command=lambda: self.set_manual_decision("DIFERENTE")
        ).pack(side="left", padx=4)

        ttk.Button(
            buttons,
            text=self.t("all_duplicate"),
            command=lambda: self.set_all_pending_decisions("DUPLICATA")
        ).pack(side="left", padx=4)

        ttk.Button(
            buttons,
            text=self.t("all_different"),
            command=lambda: self.set_all_pending_decisions("DIFERENTE")
        ).pack(side="left", padx=4)

        ttk.Button(
            buttons,
            text=self.t("evidence"),
            command=self.show_evidence_details
        ).pack(side="left", padx=4)

        ttk.Label(
            buttons,
            text=self.t("color_legend")
        ).pack(side="right")

        table_frame = ttk.Frame(duplicates_box)
        table_frame.pack(fill="both", expand=True)

        columns = [
            "status",
            "criterio",
            "similaridade",
            "base_referencia",
            "titulo_referencia",
            "base_comparada",
            "titulo_comparado",
            "decisao_manual",
        ]

        self.comparison_tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            selectmode="extended"
        )

        headings = {
            "status": self.t("status"),
            "criterio": self.t("criterion"),
            "similaridade": "%",
            "base_referencia": self.t("existing_base"),
            "titulo_referencia": self.t("existing_title"),
            "base_comparada": self.t("new_base"),
            "titulo_comparado": self.t("new_title"),
            "decisao_manual": self.t("decision"),
        }

        widths = {
            "status": 150,
            "criterio": 130,
            "similaridade": 60,
            "base_referencia": 130,
            "titulo_referencia": 420,
            "base_comparada": 130,
            "titulo_comparado": 420,
            "decisao_manual": 100,
        }

        for column in columns:
            self.comparison_tree.heading(
                column,
                text=headings[column]
            )
            self.comparison_tree.column(
                column,
                width=widths[column],
                minwidth=60
            )

        y_comp = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.comparison_tree.yview
        )

        x_comp = ttk.Scrollbar(
            table_frame,
            orient="horizontal",
            command=self.comparison_tree.xview
        )

        self.comparison_tree.configure(
            yscrollcommand=y_comp.set,
            xscrollcommand=x_comp.set
        )

        self.comparison_tree.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        y_comp.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        x_comp.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        # Exportação
        export_box = ttk.LabelFrame(
            self,
            text=self.t("export"),
            padding=8
        )
        export_box.pack(fill="x", padx=10, pady=(0, 10))

        ttk.Button(
            export_box,
            text=self.t("export_all"),
            command=self.export_all
        ).pack(side="left", padx=4)

        ttk.Label(
            export_box,
            text=(
                self.t("export_note")
            )
        ).pack(side="left", padx=15)

        self.status_var = tk.StringVar(value=self.t("ready"))

        tk.Label(
            self,
            textvariable=self.status_var,
            anchor="w",
            relief="sunken",
            bg="#D6E9F8",
            fg="#0B4F8A",
            font=("Segoe UI", 9, "bold")
        ).pack(side="bottom", fill="x")

    # --------------------------------------------------------
    # Arquivos
    # --------------------------------------------------------

    def add_files(self):
        requested_database = self._selected_database_internal()

        files = self._askopenfilenames(
            title=f"Adicionar arquivos - {requested_database}",
            filetypes=[
                (
                    "Arquivos bibliográficos",
                    "*.bib *.ris *.nbib *.txt *.ciw *.csv *.xml *.docx *.xlsx *.xls *.html *.htm *.pdf"
                ),
                ("BibTeX", "*.bib"),
                ("RIS", "*.ris"),
                ("PubMed NBIB", "*.nbib"),
                ("Texto", "*.txt *.ciw"),
                ("CSV", "*.csv"),
                ("XML", "*.xml"),
                ("MS Word", "*.docx"),
                ("MS Excel", "*.xlsx *.xls"),
                ("HTML", "*.html *.htm"),
                ("PDF", "*.pdf"),
                ("Todos", "*.*"),
            ]
        )

        existing_paths = {entry["path"] for entry in self.file_entries}

        for filename in files:
            if filename in existing_paths:
                continue

            path = Path(filename)
            detected_database, detected_style = detect_database_and_format(
                path,
                requested_database
            )

            self.file_entries.append({
                "path": filename,
                "database": detected_database,
                "style": detected_style,
                "requested_database": requested_database,
            })
            existing_paths.add(filename)

            self.files_tree.insert(
                "",
                "end",
                values=(
                    self.display_database_name(detected_database),
                    str(path),
                    path.suffix.lower(),
                    detected_style
                )
            )

        if files and requested_database == "Automático":
            self.status_var.set(
                self.translate_literal("Arquivos adicionados com detecção automática de base/formato.")
            )

    def remove_selected_files(self):
        selected = self.files_tree.selection()

        for item in selected:
            values = self.files_tree.item(item, "values")

            if values:
                database = self.internal_database_name(values[0])
                path = values[1]

                self.file_entries = [
                    entry
                    for entry in self.file_entries
                    if not (
                        entry["database"] == database
                        and entry["path"] == path
                    )
                ]

            self.files_tree.delete(item)

    def clear_files(self):
        self.file_entries = []
        self.files_tree.delete(*self.files_tree.get_children())

        self.base_data = {}
        self.all_df = pd.DataFrame(columns=FIELDS)
        self.base_final_df = pd.DataFrame(columns=FIELDS)
        self.final_df = pd.DataFrame(columns=FIELDS)
        self.comparison_df = pd.DataFrame(
            columns=COMPARISON_FIELDS
        )
        self.comparison_order = []
        self.triage_df = pd.DataFrame(columns=TRIAGE_FIELDS)  # legado v0.9; sem interface nesta versão
        self.review_filter_df = pd.DataFrame(columns=SECOND_SCREEN_FIELDS)
        self.second_screen_applied = False
        self.current_project_path = None
        self.audit_events = []
        self.last_biblio_merge_summary = {}

        # Estado do processamento assíncrono. O worker nunca toca diretamente
        # no Tkinter; ele envia eventos por fila e a thread principal os aplica.
        self._compare_queue = queue.Queue()
        self._compare_running = False
        self._compare_generation = 0
        self._tree_fill_generation = 0

        self.prisma_extra = {
            "records_excluded_title_abstract": 0,
            "reports_not_retrieved": 0,
            "full_text_excluded": 0,
            "studies_included": 0,
            "full_text_reasons": "",
        }

        self.summary_tree.delete(
            *self.summary_tree.get_children()
        )

        self.comparison_tree.delete(
            *self.comparison_tree.get_children()
        )

        self.reference_var.set("-")

        self.summary_text.set(self.t("no_files_processed"))
        self.status_var.set(self.t("ready"))

        if hasattr(
            self,
            "progress_var"
        ):
            self.reset_progress()

    # --------------------------------------------------------
    # Processamento
    # --------------------------------------------------------

    def set_progress(self, percent, message=None):
        """Atualiza barra e texto sem bloquear a interface."""
        try:
            percent = max(
                0.0,
                min(
                    100.0,
                    float(percent)
                )
            )
        except Exception:
            percent = 0.0

        self.progress_var.set(
            percent
        )

        if message is None:
            message = (
                f"{percent:.0f}%"
            )

        self.progress_label_var.set(
            self.translate_literal(message)
        )

        self.update_idletasks()

    def reset_progress(self):
        self.set_progress(
            0,
            self.t("waiting_processing")
        )

    def process_files(self):
        if self._compare_running:
            self._showinfo(
                "Processamento",
                "A comparação ainda está em andamento. Aguarde a conclusão antes de processar novamente."
            )
            return

        if not self.file_entries:
            self._showwarning(
                "Sem arquivos",
                "Adicione pelo menos um arquivo."
            )
            return

        self.set_progress(
            0,
            "Preparando arquivos • 0%"
        )

        records_by_base = defaultdict(list)
        errors = []

        total_files = max(
            1,
            len(self.file_entries)
        )

        for file_index, entry in enumerate(
            self.file_entries,
            1
        ):
            try:
                records = parse_file(
                    entry["path"],
                    entry["database"]
                )

                records_by_base[
                    entry["database"]
                ].extend(records)

            except Exception as exc:
                errors.append(
                    f"{entry['database']} | "
                    f"{Path(entry['path']).name}: {exc}"
                )

            # Importação ocupa os primeiros 20% da barra.
            import_percent = (
                file_index
                / total_files
            ) * 20

            self.set_progress(
                import_percent,
                (
                    f"Importando arquivos • "
                    f"{import_percent:.0f}%"
                )
            )

        if not records_by_base:
            self._showerror(
                "Falha",
                "Nenhum registro pôde ser importado.\n\n"
                + "\n".join(errors)
            )
            return

        self.base_data = {}

        frames = []

        for database, records in records_by_base.items():
            df = standardize(records)

            # IDs locais para auditoria
            df = df.reset_index(drop=True)
            df["id_origem"] = df["id_origem"].where(
                df["id_origem"].astype(str).str.strip() != "",
                pd.Series(
                    [str(i + 1) for i in range(len(df))],
                    index=df.index
                )
            )

            self.base_data[database] = df
            frames.append(df)

        self.all_df = pd.concat(
            frames,
            ignore_index=True
        )

        counts = {
            database: len(df)
            for database, df in self.base_data.items()
        }

        ordered = sorted(
            counts,
            key=counts.get,
            reverse=True
        )

        # A maior base é sempre a referência automática.
        self.reference_var.set(
            f"{self.display_database_name(ordered[0])} ({counts[ordered[0]]} {self.t('records')})"
        )

        self.update_summary()
        self.set_progress(
            20,
            "Arquivos importados • 20%"
        )
        def finish_processing_message():
            if errors:
                self._showwarning(
                    "Processado com avisos",
                    (
                        "Alguns arquivos não foram lidos:\n\n"
                        + "\n".join(errors)
                        + "\n\nA ferramenta aceita várias variantes por base. "
                        "Se um novo formato aparecer, envie o arquivo para adicionarmos "
                        "mais um detector sem alterar os anteriores."
                    )
                )
            else:
                self._showinfo(
                    "Processado",
                    f"{len(self.all_df)} registros foram importados "
                    f"de {len(self.base_data)} base(s)."
                )

        self.recompare(
            show_message=False,
            progress_start=20,
            progress_end=100,
            completion_callback=finish_processing_message,
        )

    def recompare(
        self,
        show_message=True,
        progress_start=0,
        progress_end=100,
        completion_callback=None,
    ):
        """Executa a comparação em worker thread sem bloquear a interface."""
        if self.all_df.empty:
            return

        if self._compare_running:
            if show_message:
                self._showinfo(
                    "Comparação",
                    "Já existe uma comparação em andamento."
                )
            return

        self.status_var.set(self.translate_literal("Comparando as bases..."))
        self.set_progress(
            progress_start,
            f"Iniciando comparação • {progress_start:.0f}%",
        )

        # Tudo que pertence ao Tkinter é lido antes de iniciar o worker.
        threshold = int(self.threshold_var.get())
        data_snapshot = self.all_df.copy()

        self._compare_running = True
        self._compare_generation += 1
        generation = self._compare_generation
        self._compare_queue = queue.Queue()

        def comparison_progress(internal_percent, internal_message):
            external_percent = (
                progress_start
                + (internal_percent / 100)
                * (progress_end - progress_start)
            )
            message_base = internal_message.split("•")[0].strip()
            self._compare_queue.put((
                "progress",
                external_percent,
                f"{message_base} • {external_percent:.0f}%",
            ))

        def worker():
            try:
                result = sequential_compare(
                    data_snapshot,
                    reference_database=None,
                    threshold=threshold,
                    progress_callback=comparison_progress,
                )
            except Exception as exc:
                self._compare_queue.put(("error", str(exc)))
                return
            self._compare_queue.put(("done", result))

        threading.Thread(
            target=worker,
            name="review-tool-dedup",
            daemon=True,
        ).start()

        self.after(
            35,
            lambda: self._poll_compare_queue(
                generation,
                show_message,
                progress_end,
                completion_callback,
            ),
        )

    def _poll_compare_queue(
        self,
        generation,
        show_message,
        progress_end,
        completion_callback,
    ):
        """Aplica na thread principal os eventos produzidos pelo worker."""
        if generation != self._compare_generation:
            return

        finished = False

        while True:
            try:
                event = self._compare_queue.get_nowait()
            except queue.Empty:
                break

            kind = event[0]

            if kind == "progress":
                _, percent, message = event
                self.set_progress(percent, message)

            elif kind == "error":
                self._compare_running = False
                finished = True
                self.status_var.set(self.translate_literal("Erro durante a comparação"))
                self._showerror(
                    "Erro de comparação",
                    event[1],
                )

            elif kind == "done":
                self._compare_running = False
                finished = True
                self._apply_compare_result(
                    event[1],
                    show_message=show_message,
                    progress_end=progress_end,
                    completion_callback=completion_callback,
                )

        if self._compare_running and not finished:
            self.after(
                35,
                lambda: self._poll_compare_queue(
                    generation,
                    show_message,
                    progress_end,
                    completion_callback,
                ),
            )

    def _apply_compare_result(
        self,
        result,
        show_message=True,
        progress_end=100,
        completion_callback=None,
    ):
        self.base_final_df, self.comparison_df, self.comparison_order = result

        # Sem decisões manuais, o final inicial é a base após duplicatas exatas.
        self.final_df = self.base_final_df.copy()
        self.second_screen_applied = False
        self.review_filter_df = pd.DataFrame(columns=SECOND_SCREEN_FIELDS)
        self.sync_triage_df()

        counts = (
            self.all_df
            .groupby("base")
            .size()
            .sort_values(ascending=False)
        )

        if not counts.empty:
            main = counts.index[0]
            self.reference_var.set(
                f"{self.display_database_name(main)} ({int(counts.iloc[0])} {self.t('records')})"
            )

        self.log_audit_event(
            "DEDUPLICAÇÃO CONCLUÍDA",
            f"Limiar fuzzy: {int(self.threshold_var.get())}%; ordem: {' → '.join(self.comparison_order)}",
        )
        self.update_summary()
        self.set_progress(
            progress_end,
            "Concluído • 100%"
            if progress_end >= 100
            else f"Concluído • {progress_end:.0f}%",
        )
        self.status_var.set(
            self.translate_literal(f"Comparação concluída • {len(self.final_df)} registros no banco atual")
        )

        def finish_ui():
            if show_message:
                self._showinfo(
                    "Comparação atualizada",
                    (
                        f"Base principal: "
                        f"{counts.index[0] if not counts.empty else '-'}\n"
                        f"Registros finais atuais: {len(self.final_df)}\n"
                        f"Comparações registradas: {len(self.comparison_df)}"
                    ),
                )
            if completion_callback is not None:
                completion_callback()

        # A tabela é populada aos poucos para a janela continuar responsiva.
        self.fill_comparison_tree(on_complete=finish_ui)

    def base_report_df(self):
        """
        Mini-relatório por base.

        Duplicatas removidas são atribuídas à base que entrou depois
        na comparação. Assim, a base principal preserva sua cópia.
        """
        rows = []

        if not self.base_data:
            return pd.DataFrame(
                columns=[
                    "Base",
                    "Bruto",
                    "Duplicatas removidas",
                    "Pendentes",
                    "Final"
                ]
            )

        for database, df in sorted(
            self.base_data.items(),
            key=lambda item: len(item[1]),
            reverse=True
        ):
            bruto = len(df)

            duplicatas = 0
            pendentes = 0

            if (
                self.comparison_df is not None
                and not self.comparison_df.empty
            ):
                duplicatas = int(
                    (
                        (
                            self.comparison_df[
                                "base_comparada"
                            ]
                            == database
                        )
                        & (
                            self.comparison_df[
                                "status"
                            ].isin([
                                "DUPLICATA CONFIRMADA",
                                "DUPLICATA CONFIRMADA MANUAL"
                            ])
                        )
                    ).sum()
                )

                pendentes = int(
                    (
                        (
                            self.comparison_df[
                                "base_comparada"
                            ]
                            == database
                        )
                        & (
                            self.comparison_df[
                                "status"
                            ]
                            == "REVISAR"
                        )
                    ).sum()
                )

            final_count = 0

            if (
                self.final_df is not None
                and not self.final_df.empty
            ):
                final_count = int(
                    (
                        self.final_df["base"]
                        == database
                    ).sum()
                )

            rows.append({
                "Base": database,
                "Bruto": bruto,
                "Duplicatas removidas": duplicatas,
                "Pendentes": pendentes,
                "Final": final_count,
            })

        return pd.DataFrame(rows)

    def update_summary(self):
        self.summary_tree.delete(
            *self.summary_tree.get_children()
        )

        mini = self.base_report_df()

        if mini.empty:
            self.summary_text.set(
                self.t("no_files_processed")
            )
            return

        for _, row in mini.iterrows():
            self.summary_tree.insert(
                "",
                "end",
                values=(
                    self.display_database_name(row["Base"]),
                    int(row["Bruto"]),
                    int(row["Duplicatas removidas"]),
                    int(row["Pendentes"]),
                    int(row["Final"]),
                )
            )

        total_bruto = int(
            mini["Bruto"].sum()
        )
        total_duplicates = int(
            mini["Duplicatas removidas"].sum()
        )
        total_pending = int(
            mini["Pendentes"].sum()
        )
        total_final = int(
            mini["Final"].sum()
        )

        self.summary_tree.insert(
            "",
            "end",
            values=(
                "TOTAL",
                total_bruto,
                total_duplicates,
                total_pending,
                total_final,
            ),
            tags=("total_row",)
        )

        self.summary_tree.tag_configure(
            "total_row",
            background="#D6E9F8",
            foreground="#0B4F8A",
            font=("Segoe UI", 9, "bold")
        )

        auto_duplicates = 0
        manual_duplicates = 0

        if not self.comparison_df.empty:
            auto_duplicates = int(
                (
                    self.comparison_df["status"]
                    == "DUPLICATA CONFIRMADA"
                ).sum()
            )

            manual_duplicates = int(
                (
                    self.comparison_df["status"]
                    == "DUPLICATA CONFIRMADA MANUAL"
                ).sum()
            )

        order_text = (
            " → ".join(
                self.comparison_order
            )
            if self.comparison_order
            else "-"
        )

        if getattr(self, "second_screen_applied", False):
            stats = self.second_screen_stats()
            second_text = self.t(
                "second_summary",
                confirmed=stats["confirmed"],
                pending=stats["pending"],
            )
        else:
            second_text = self.t("second_not_run")

        display_order = " → ".join(
            self.display_database_name(db) for db in self.comparison_order
        ) if self.comparison_order else "-"

        self.summary_text.set(
            f"{self.t('main_database')}: {self.reference_var.get()}\n"
            f"{self.t('order')}: {display_order}\n"
            f"{self.t('auto_duplicates')}: {auto_duplicates}\n"
            f"{self.t('manual_duplicates')}: {manual_duplicates}\n"
            f"{self.t('pending_check')}: {total_pending}\n"
            f"{self.t('total_after')}: {total_final}\n"
            f"{second_text}"
        )

    # --------------------------------------------------------
    # Comparação visual
    # --------------------------------------------------------

    def fill_comparison_tree(self, on_complete=None, chunk_size=250):
        """Preenche a grade em blocos para evitar congelamento com muitas linhas."""
        self._tree_fill_generation += 1
        generation = self._tree_fill_generation

        self.comparison_tree.delete(
            *self.comparison_tree.get_children()
        )

        self.comparison_tree.tag_configure(
            "pending",
            background="#FFF2A8"
        )
        self.comparison_tree.tag_configure(
            "confirmed",
            background="#D9EAD3"
        )
        self.comparison_tree.tag_configure(
            "different",
            background="#E7E6E6"
        )

        if self.comparison_df.empty:
            if on_complete is not None:
                on_complete()
            return

        columns = list(self.comparison_tree["columns"])
        rows = self.comparison_df.to_dict("records")

        def insert_chunk(start_index=0):
            if generation != self._tree_fill_generation:
                return

            stop_index = min(start_index + chunk_size, len(rows))

            for row in rows[start_index:stop_index]:
                status = clean(row.get("status", ""))

                if status == "REVISAR":
                    color_tag = "pending"
                elif status in {
                    "DUPLICATA CONFIRMADA",
                    "DUPLICATA CONFIRMADA MANUAL",
                }:
                    color_tag = "confirmed"
                else:
                    color_tag = "different"

                self.comparison_tree.insert(
                    "",
                    "end",
                    values=[
                        (
                            self.display_database_name(clean(row.get(column, "")))
                            if column in {"base_referencia", "base_comparada"}
                            else self.display_value(clean(row.get(column, "")))
                            if column in {"status", "criterio", "decisao_manual"}
                            else clean(row.get(column, ""))
                        )
                        for column in columns
                    ],
                    tags=(
                        row["comparacao_id"],
                        color_tag,
                    ),
                )

            if stop_index < len(rows):
                self.after(1, lambda: insert_chunk(stop_index))
            else:
                self.comparison_tree.update_idletasks()
                if on_complete is not None:
                    on_complete()

        insert_chunk(0)

    def set_manual_decision(self, decision):
        if self.comparison_df.empty:
            return

        selected = self.comparison_tree.selection()

        if not selected:
            self._showinfo(
                "Seleção",
                "Selecione pelo menos uma linha amarela (REVISAR)."
            )
            return

        ids = []

        for item in selected:
            tags = self.comparison_tree.item(
                item,
                "tags"
            )

            if tags:
                ids.append(tags[0])

        mask = (
            self.comparison_df["comparacao_id"].isin(ids)
            & (
                self.comparison_df["status"]
                .isin([
                    "REVISAR",
                    "DUPLICATA CONFIRMADA MANUAL",
                    "NÃO DUPLICATA"
                ])
            )
            & (
                self.comparison_df["criterio"]
                == "Título semelhante"
            )
        )

        changed = int(mask.sum())

        self.set_progress(
            0,
            "Atualizando decisão • 0%"
        )

        if changed == 0:
            self._showinfo(
                "Sem alteração",
                "A seleção não contém um caso de similaridade que possa ser revisado."
            )
            return

        if decision == "DUPLICATA":
            self.comparison_df.loc[mask, "decisao_manual"] = "DUPLICATA"
            self.comparison_df.loc[mask, "status"] = "DUPLICATA CONFIRMADA MANUAL"
        else:
            self.comparison_df.loc[mask, "decisao_manual"] = "DIFERENTE"
            self.comparison_df.loc[mask, "status"] = "NÃO DUPLICATA"

        # Rápido: NÃO recalcula toda a comparação.
        self.rebuild_final_from_decisions()
        self.log_audit_event(
            "DECISÃO MANUAL DE DUPLICATA",
            f"{changed} comparação(ões): {decision}",
        )

        # Refresh visual imediato.
        self.fill_comparison_tree()
        self.update_summary()
        self.update_idletasks()

        self.set_progress(
            100,
            "Decisão atualizada • 100%"
        )

        self.status_var.set(
            self.translate_literal(f"OK • {changed} decisão(ões) atualizada(s)")
        )

        self._showinfo(
            "OK",
            (
                f"{changed} comparação(ões) confirmada(s) como duplicata."
                if decision == "DUPLICATA"
                else f"{changed} comparação(ões) mantida(s) como diferente."
            )
        )

    def set_all_pending_decisions(self, decision):
        """Aplica uma decisão em lote a todos os casos amarelos (REVISAR)."""
        if self.comparison_df.empty:
            self._showinfo("Duplicatas", "Não há comparações disponíveis.")
            return

        mask = (
            (self.comparison_df["status"] == "REVISAR")
            & (self.comparison_df["criterio"] == "Título semelhante")
        )
        count = int(mask.sum())
        if count == 0:
            self._showinfo("Duplicatas", "Não há duplicatas pendentes para decidir.")
            return

        if decision == "DUPLICATA":
            question = (
                f"Confirmar TODOS os {count} casos amarelos como duplicata?\n\n"
                "Esses registros serão removidos do banco final."
            )
        else:
            question = (
                f"Marcar TODOS os {count} casos amarelos como não duplicata?\n\n"
                "Esses registros permanecerão no banco final."
            )

        if not self._askyesno("Decisão em lote", question):
            return

        if decision == "DUPLICATA":
            self.comparison_df.loc[mask, "decisao_manual"] = "DUPLICATA"
            self.comparison_df.loc[mask, "status"] = "DUPLICATA CONFIRMADA MANUAL"
        else:
            self.comparison_df.loc[mask, "decisao_manual"] = "DIFERENTE"
            self.comparison_df.loc[mask, "status"] = "NÃO DUPLICATA"

        self.rebuild_final_from_decisions()
        self.log_audit_event(
            "DECISÃO EM LOTE DE DUPLICATAS",
            f"{count} comparação(ões): {decision}",
        )
        self.fill_comparison_tree()
        self.update_summary()
        self.update_idletasks()
        self.set_progress(100, f"Decisão em lote aplicada a {count} caso(s) • 100%")
        self.status_var.set(self.translate_literal(f"OK • {count} decisão(ões) em lote atualizada(s)"))

        self._showinfo(
            "Decisão em lote concluída",
            (
                f"{count} caso(s) confirmado(s) como duplicata."
                if decision == "DUPLICATA"
                else f"{count} caso(s) mantido(s) como não duplicata."
            ),
        )

    def rebuild_final_from_decisions(self):
        """
        Atualização rápida: parte do banco já deduplicado por DOI/título exato
        e apenas remove os UIDs que o usuário confirmou como duplicata fuzzy.
        Não refaz a comparação inteira.
        """
        if self.base_final_df.empty:
            self.final_df = self.base_final_df.copy()
            return

        final = self.base_final_df.copy()

        if not self.comparison_df.empty:
            manual_duplicates = self.comparison_df[
                self.comparison_df["status"]
                == "DUPLICATA CONFIRMADA MANUAL"
            ]

            uids_to_remove = set(
                clean(uid)
                for uid in manual_duplicates["uid_comparado"].tolist()
                if clean(uid)
            )

            if uids_to_remove:
                final = final[
                    ~final["uid"].isin(uids_to_remove)
                ].copy()

        self.final_df = final.reset_index(drop=True)
        self.sync_triage_df()
        if getattr(self, "second_screen_applied", False):
            self.sync_review_filter_df()

    # --------------------------------------------------------
    # Projeto, 2ª triagem, busca e PRISMA
    # --------------------------------------------------------

    def new_project(self):
        """Limpa o estado atual e inicia um novo projeto."""
        if not self.all_df.empty:
            answer = self._askyesno(
                "Novo projeto",
                "Iniciar um novo projeto apagará da tela os dados atuais.\n\n"
                "Salve o projeto antes, se quiser continuar depois. Deseja prosseguir?"
            )
            if not answer:
                return
        self.clear_files()
        self.log_audit_event("NOVO PROJETO", "Estado do projeto reiniciado.")

    def log_audit_event(self, action, details=""):
        """Registra eventos relevantes para reprodutibilidade e auditoria."""
        try:
            raw_n = int(len(self.all_df))
            final_n = int(len(self.final_df))
        except Exception:
            raw_n = final_n = 0
        pending = 0
        try:
            if self.comparison_df is not None and not self.comparison_df.empty:
                pending = int((self.comparison_df["status"] == "REVISAR").sum())
        except Exception:
            pass
        self.audit_events.append({
            "data_hora": datetime.now().isoformat(timespec="seconds"),
            "acao": clean(action),
            "detalhes": clean(details),
            "registros_brutos": raw_n,
            "registros_finais": final_n,
            "pendentes_duplicata": pending,
            "versao_programa": VERSION,
        })

    @staticmethod
    def _sha256_file(path):
        try:
            h = hashlib.sha256()
            with open(path, "rb") as handle:
                while True:
                    chunk = handle.read(1024 * 1024)
                    if not chunk:
                        break
                    h.update(chunk)
            return h.hexdigest()
        except Exception:
            return ""

    def audit_files_df(self):
        rows = []
        for entry in self.file_entries:
            path = Path(entry.get("path", ""))
            exists = path.exists() and path.is_file()
            try:
                stat = path.stat() if exists else None
            except Exception:
                stat = None
            rows.append({
                "Base": clean(entry.get("database", "")),
                "Arquivo": path.name or clean(entry.get("path", "")),
                "Caminho_original": str(path),
                "Formato_detectado": clean(entry.get("style", "")),
                "Base_solicitada": clean(entry.get("requested_database", "")),
                "Arquivo_localizado": "SIM" if exists else "NÃO",
                "Tamanho_bytes": int(stat.st_size) if stat else "",
                "Modificado_em": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds") if stat else "",
                "SHA256": self._sha256_file(path) if exists else "",
            })
        return pd.DataFrame(rows)

    def project_metrics(self):
        mini = self.base_report_df()
        auto_dup = manual_dup = pending = 0
        if self.comparison_df is not None and not self.comparison_df.empty:
            auto_dup = int((self.comparison_df["status"] == "DUPLICATA CONFIRMADA").sum())
            manual_dup = int((self.comparison_df["status"] == "DUPLICATA CONFIRMADA MANUAL").sum())
            pending = int((self.comparison_df["status"] == "REVISAR").sum())
        second = self.second_screen_stats() if self.second_screen_applied else {
            "confirmed": 0, "pending": 0, "detected": 0, "dismissed": 0,
            "type_counts": {}, "base_counts": {},
        }
        return {
            "bruto": int(len(self.all_df)),
            "duplicatas_auto": auto_dup,
            "duplicatas_manual": manual_dup,
            "duplicatas_total": max(0, int(len(self.all_df)) - int(len(self.final_df))),
            "pendentes": pending,
            "final": int(len(self.final_df)),
            "bases": int(len(self.base_data)),
            "segunda_confirmados": int(second.get("confirmed", 0)),
            "segunda_pendentes": int(second.get("pending", 0)),
            "mini": mini,
        }

    def open_dashboard_window(self):
        if self.all_df.empty:
            self._showwarning("Painel", "Processe pelo menos uma base para abrir o painel do projeto.")
            return
        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("Painel do projeto")
        win.geometry("1050x720")
        win.minsize(850, 620)

        header = tk.Frame(win, bg="#0B4F8A", pady=12)
        header.pack(fill="x")
        tk.Label(header, text="PAINEL DO PROJETO", bg="#0B4F8A", fg="white", font=("Segoe UI", 18, "bold")).pack()
        project_name = Path(self.current_project_path).name if self.current_project_path else "Projeto ainda não salvo"
        tk.Label(header, text=project_name, bg="#0B4F8A", fg="#DDEEFF", font=("Segoe UI", 9)).pack()

        cards = ttk.Frame(win, padding=12)
        cards.pack(fill="x")
        for col in range(4):
            cards.columnconfigure(col, weight=1)

        card_vars = {}
        specs = [
            ("bruto", "REGISTROS BRUTOS"),
            ("duplicatas_total", "DUPLICATAS REMOVIDAS"),
            ("final", "APÓS DEDUPLICAÇÃO"),
            ("pendentes", "DUPLICATAS PENDENTES"),
            ("bases", "BASES IMPORTADAS"),
            ("segunda_confirmados", "REVISÕES/META CONFIRMADAS"),
            ("segunda_pendentes", "2ª TRIAGEM PENDENTE"),
            ("salvo", "PROJETO"),
        ]
        for idx, (key, title) in enumerate(specs):
            box = ttk.LabelFrame(cards, text=title, padding=10)
            box.grid(row=idx // 4, column=idx % 4, padx=5, pady=5, sticky="nsew")
            var = tk.StringVar(value="-")
            card_vars[key] = var
            ttk.Label(box, textvariable=var, font=("Segoe UI", 18, "bold")).pack(pady=8)

        table_box = ttk.LabelFrame(win, text="Resumo por base", padding=8)
        table_box.pack(fill="both", expand=True, padx=12, pady=(0, 8))
        tree = ttk.Treeview(table_box, columns=("base","bruto","dup","pend","final"), show="headings", height=9)
        for key, label, width in [
            ("base","Base",230),("bruto","Bruto",110),("dup","Duplicatas",120),("pend","Pendentes",110),("final","Final",110)
        ]:
            tree.heading(key, text=label); tree.column(key, width=width, anchor="center" if key != "base" else "w")
        tree.pack(fill="both", expand=True)

        footer_var = tk.StringVar()
        ttk.Label(win, textvariable=footer_var, padding=(12,4)).pack(fill="x")

        def refresh():
            m = self.project_metrics()
            for key in ["bruto","duplicatas_total","final","pendentes","bases","segunda_confirmados","segunda_pendentes"]:
                card_vars[key].set(str(m[key]))
            card_vars["salvo"].set(self.translate_literal("SALVO" if self.current_project_path else "NÃO SALVO"))
            tree.delete(*tree.get_children())
            for _, row in m["mini"].iterrows():
                tree.insert("", "end", values=(row["Base"], int(row["Bruto"]), int(row["Duplicatas removidas"]), int(row["Pendentes"]), int(row["Final"])))
            footer_var.set(self.translate_literal(
                f"Ordem de comparação: {' → '.join(self.comparison_order) if self.comparison_order else '-'}  |  "
                f"Limiar fuzzy: {int(self.threshold_var.get())}%  |  Versão {VERSION}"
            ))

        actions = ttk.Frame(win, padding=(12, 4, 12, 12))
        actions.pack(fill="x")
        ttk.Button(actions, text="ATUALIZAR", command=refresh).pack(side="left")
        ttk.Button(actions, text="PRISMA", command=self.open_prisma_window).pack(side="left", padx=6)
        ttk.Button(actions, text="BIBLIOMETRIA", command=self.open_bibliometric_merge_window).pack(side="left", padx=6)
        ttk.Button(actions, text="SALVAR PROJETO", command=self.save_project).pack(side="right")
        refresh()

    def export_audit_report(self, folder):
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        xlsx = folder / "AUDITORIA_REPRODUTIBILIDADE.xlsx"
        txt = folder / "AUDITORIA_REPRODUTIBILIDADE.txt"
        metrics = self.project_metrics()

        summary_rows = [
            ("Programa", APP_TITLE), ("Versão", VERSION),
            ("Desenvolvedor", AUTHOR_NAME),
            ("Licença", LICENSE_NAME),
            ("Repositório", PROJECT_URL),
            ("Gerado em", datetime.now().isoformat(timespec="seconds")),
            ("Projeto", self.current_project_path or "não salvo"),
            ("Limiar fuzzy (%)", int(self.threshold_var.get())),
            ("Ordem de comparação", " → ".join(self.comparison_order)),
            ("Registros brutos", metrics["bruto"]),
            ("Duplicatas automáticas", metrics["duplicatas_auto"]),
            ("Duplicatas confirmadas manualmente", metrics["duplicatas_manual"]),
            ("Duplicatas removidas - total", metrics["duplicatas_total"]),
            ("Pendentes de decisão", metrics["pendentes"]),
            ("Registros após deduplicação", metrics["final"]),
            ("2ª triagem executada", "SIM" if self.second_screen_applied else "NÃO"),
            ("Revisões/meta confirmadas", metrics["segunda_confirmados"]),
            ("2ª triagem pendentes", metrics["segunda_pendentes"]),
        ]
        summary_df = pd.DataFrame(summary_rows, columns=["Parâmetro", "Valor"])
        files_df = self.audit_files_df()
        events_df = pd.DataFrame(self.audit_events)
        if events_df.empty:
            events_df = pd.DataFrame(columns=["data_hora","acao","detalhes","registros_brutos","registros_finais","pendentes_duplicata","versao_programa"])

        with pd.ExcelWriter(xlsx, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="RESUMO", index=False)
            metrics["mini"].to_excel(writer, sheet_name="BASES", index=False)
            files_df.to_excel(writer, sheet_name="ARQUIVOS", index=False)
            events_df.to_excel(writer, sheet_name="EVENTOS", index=False)
            if self.comparison_df is not None and not self.comparison_df.empty:
                self.comparison_df.to_excel(writer, sheet_name="DECISOES_DUPLICATAS", index=False)
            if self.second_screen_applied and self.review_filter_df is not None and not self.review_filter_df.empty:
                self.review_filter_df.to_excel(writer, sheet_name="SEGUNDA_TRIAGEM", index=False)
            from openpyxl.styles import PatternFill, Font
            blue = PatternFill(fill_type="solid", fgColor="0B4F8A")
            for ws in writer.book.worksheets:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.fill = blue; cell.font = Font(bold=True, color="FFFFFF")
                # Larguras fixas evitam varrer centenas de milhares de células só para medir texto.
                for letter, width in {
                    "A": 28, "B": 34, "C": 55, "D": 28, "E": 24, "F": 18,
                    "G": 18, "H": 22, "I": 42, "J": 30, "K": 30, "L": 30,
                    "M": 30, "N": 30, "O": 30, "P": 30, "Q": 30, "R": 30,
                    "S": 30, "T": 30,
                }.items():
                    ws.column_dimensions[letter].width = width

        lines = ["AUDITORIA E REPRODUTIBILIDADE", "="*72]
        lines += [f"{k}: {v}" for k,v in summary_rows]
        lines += ["", "ARQUIVOS DE ORIGEM"]
        if files_df.empty:
            lines.append("Nenhum arquivo registrado.")
        else:
            for _, r in files_df.iterrows():
                lines.append(f"- {r['Base']} | {r['Arquivo']} | formato={r['Formato_detectado']} | SHA256={r['SHA256'] or 'indisponível'}")
        lines += ["", "HISTÓRICO DE EVENTOS"]
        for _, r in events_df.iterrows():
            lines.append(f"- {r.get('data_hora','')} | {r.get('acao','')} | {r.get('detalhes','')}")
        txt.write_text("\n".join(lines), encoding="utf-8")
        return [xlsx, txt]

    def export_audit_dialog(self):
        if self.all_df.empty:
            self._showwarning("Auditoria", "Processe pelo menos uma base antes de exportar a auditoria.")
            return
        folder = self._askdirectory(title="Pasta para salvar a auditoria")
        if not folder:
            return
        try:
            created = self.export_audit_report(folder)
            self.log_audit_event("RELATÓRIO DE AUDITORIA EXPORTADO", str(folder))
        except Exception as exc:
            self._showerror("Auditoria", str(exc))
            return
        self._showinfo("Auditoria", f"Relatórios criados:\n" + "\n".join(str(p) for p in created))

    def export_prisma_data_excel(self, output_path):
        """Exporta as contagens que alimentam o PRISMA em formato auditável."""
        counts = self.prisma_counts()
        flow = pd.DataFrame([
            ("Registros identificados", counts["identified"]),
            ("Duplicatas removidas", counts["duplicates"]),
            ("Registros após deduplicação", counts["deduplicated"]),
            ("Revisões/meta confirmadas e mantidas", counts["second_confirmed"]),
            ("Revisões/meta pendentes", counts["second_pending"]),
            ("Registros disponíveis para título/resumo", counts["screening"]),
            ("Excluídos em título/resumo", counts["title_abstract_excluded"]),
            ("Relatórios buscados", counts["reports_sought"]),
            ("Relatórios não recuperados", counts["reports_not_retrieved"]),
            ("Textos completos avaliados", counts["reports_assessed"]),
            ("Textos completos excluídos", counts["full_text_excluded"]),
            ("Estudos incluídos", counts["studies_included"]),
        ], columns=["Etapa", "n"])
        per_base = self.base_report_df().copy()
        second_types = pd.DataFrame(
            sorted(counts.get("second_type_counts", {}).items()),
            columns=["Classificação 2ª triagem", "n"]
        ) if counts.get("second_type_counts") else pd.DataFrame(columns=["Classificação 2ª triagem", "n"])
        extras = pd.DataFrame([
            ("Motivos de exclusão em texto completo", counts.get("full_text_reasons", "")),
            ("Gerado em", datetime.now().isoformat(timespec="seconds")),
            ("Versão", VERSION),
        ], columns=["Campo", "Valor"])
        output_path = Path(output_path)
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            flow.to_excel(writer, sheet_name="FLUXO_PRISMA", index=False)
            per_base.to_excel(writer, sheet_name="POR_BASE", index=False)
            second_types.to_excel(writer, sheet_name="SEGUNDA_TRIAGEM", index=False)
            extras.to_excel(writer, sheet_name="INFORMACOES", index=False)
            from openpyxl.styles import PatternFill, Font
            blue = PatternFill(fill_type="solid", fgColor="0B4F8A")
            for ws in writer.book.worksheets:
                ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]: cell.fill = blue; cell.font = Font(bold=True, color="FFFFFF")
                ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 30
        return output_path

    def open_reference_formatter_window(self):
        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("Formatador de referências — ABNT / Vancouver / APA")
        win.geometry("900x690")
        win.minsize(780, 600)

        info = ttk.LabelFrame(win, text="Objetivo", padding=10)
        info.pack(fill="x", padx=12, pady=10)
        ttk.Label(
            info,
            text=(
                "Carregue um DOCX ou TXT com uma referência por parágrafo/linha. O módulo extrai o DOI, "
                "consulta os metadados no Crossref e cria um NOVO DOCX formatado. "
                "Referências sem correspondência segura são preservadas e marcadas para revisão manual."
            ),
            wraplength=840,
            justify="left",
        ).pack(anchor="w")
        ttk.Label(
            info,
            text=(
                "A busca por texto é opcional porque pode encontrar trabalhos parecidos; o programa só aceita automaticamente "
                "correspondências com escore alto. DOI explícito é sempre priorizado."
            ),
            wraplength=840,
            justify="left",
        ).pack(anchor="w", pady=(6, 0))

        file_box = ttk.LabelFrame(win, text="Arquivo de entrada", padding=10)
        file_box.pack(fill="x", padx=12, pady=6)
        input_var = tk.StringVar()
        ttk.Entry(file_box, textvariable=input_var).pack(side="left", fill="x", expand=True, padx=(0, 8))

        def choose_input():
            path = self._askopenfilename(
                title="Selecione o arquivo com as referências",
                filetypes=[("Word", "*.docx"), ("Texto", "*.txt"), ("Todos", "*.*")],
                parent=win,
            )
            if path:
                input_var.set(path)

        ttk.Button(file_box, text="ESCOLHER...", command=choose_input).pack(side="right")

        options = ttk.LabelFrame(win, text="Saída", padding=10)
        options.pack(fill="x", padx=12, pady=6)
        ttk.Label(options, text="Estilo:").grid(row=0, column=0, sticky="w")
        style_var = tk.StringVar(value="TODOS")
        combo = ttk.Combobox(
            options,
            textvariable=style_var,
            values=["TODOS", "ABNT", "Vancouver", "APA 7"],
            state="readonly",
            width=18,
        )
        combo.grid(row=0, column=1, sticky="w", padx=(8, 24))
        text_search_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            options,
            text="Tentar localizar metadados pelo texto quando não houver DOI",
            variable=text_search_var,
        ).grid(row=0, column=2, sticky="w")

        progress_var = tk.DoubleVar(value=0)
        ttk.Progressbar(
            win,
            variable=progress_var,
            maximum=100,
            style="Blue.Horizontal.TProgressbar",
        ).pack(fill="x", padx=12, pady=(12, 4))
        progress_text = tk.StringVar(value="Aguardando")
        ttk.Label(win, textvariable=progress_text).pack(anchor="w", padx=12)

        log = tk.Text(win, height=18, wrap="word", font=("Consolas", 9))
        log.pack(fill="both", expand=True, padx=12, pady=8)
        log.insert("end", "Selecione um arquivo e clique em FORMATAR REFERÊNCIAS.\n")
        log.configure(state="disabled")

        action = ttk.Frame(win, padding=(12, 4, 12, 12))
        action.pack(fill="x")
        run_btn = ttk.Button(action, text="FORMATAR REFERÊNCIAS")
        run_btn.pack(side="right")

        q = queue.Queue()

        def append_log(message):
            log.configure(state="normal")
            log.insert("end", self.translate_literal(message) + "\n")
            log.see("end")
            log.configure(state="disabled")

        def poll():
            finished = False
            while True:
                try:
                    ev = q.get_nowait()
                except queue.Empty:
                    break
                kind = ev[0]
                if kind == "progress":
                    progress_var.set(ev[1])
                    progress_text.set(self.translate_literal(ev[2]))
                    append_log(ev[2])
                elif kind == "log":
                    append_log(ev[1])
                elif kind == "error":
                    finished = True
                    run_btn.state(["!disabled"])
                    self._showerror("Formatador de referências", ev[1], parent=win)
                elif kind == "done":
                    finished = True
                    run_btn.state(["!disabled"])
                    result = ev[1]
                    progress_var.set(100)
                    progress_text.set(self.translate_literal("Concluído"))
                    append_log(f"Formatadas automaticamente: {result['resolved']}")
                    append_log(f"Para revisão manual: {result['unresolved']}")
                    append_log("Arquivos criados:")
                    for item in result["files"]:
                        append_log(f"  - {item}")
                    self.log_audit_event(
                        "FORMATAÇÃO DE REFERÊNCIAS",
                        f"Entrada={result['input']}; resolvidas={result['resolved']}; não resolvidas={result['unresolved']}; estilos={'; '.join(result['styles'])}",
                    )
                    self._showinfo(
                        "Referências formatadas",
                        f"Concluído.\n\nFormatadas: {result['resolved']}\nPara revisão manual: {result['unresolved']}\n\nPasta: {result['folder']}",
                        parent=win,
                    )
            if not finished and run_btn.instate(["disabled"]):
                win.after(100, poll)

        def worker(input_path, output_folder, styles, allow_text_search):
            try:
                refs = extract_reference_paragraphs(input_path)
                refs = [
                    re.sub(r"^\s*(?:\[?\d+\]?\s*[.):-]?\s*)", "", ref).strip()
                    for ref in refs
                    if clean(ref).casefold() not in {"referências", "referencias", "references", "bibliografia"}
                    and len(clean(ref)) >= 12
                ]
                if not refs:
                    raise RuntimeError("Nenhuma referência foi reconhecida no arquivo.")

                metadata_rows = []
                formatted_by_style = {style: [] for style in styles}
                total = len(refs)
                resolved = 0
                unresolved = 0

                for i, ref in enumerate(refs, 1):
                    q.put(("progress", ((i - 1) / max(1, total)) * 86, f"Consultando referência {i}/{total}..."))
                    meta, method, note = crossref_metadata_from_reference(ref, allow_text_search=allow_text_search)
                    doi = ""
                    if meta:
                        resolved += 1
                        doi = normalize_doi(meta.get("DOI", ""))
                        for style in styles:
                            formatted = format_reference_from_crossref(meta, style)
                            formatted_by_style[style].append(formatted or ref)
                    else:
                        unresolved += 1
                        for style in styles:
                            formatted_by_style[style].append(f"[REVISAR MANUALMENTE] {ref}")

                    metadata_rows.append({
                        "Referencia_original": ref,
                        "DOI": doi or extract_doi_from_text(ref),
                        "Metodo": method,
                        "Status": "FORMATADA" if meta else "REVISAR",
                        "Observacao": note,
                    })

                folder = Path(output_folder)
                folder.mkdir(parents=True, exist_ok=True)
                stem = Path(input_path).stem
                created = []

                for style in styles:
                    items = list(formatted_by_style[style])
                    if style.upper().startswith("VAN"):
                        items = [f"{n}. {text}" for n, text in enumerate(items, 1)]
                    else:
                        items = sorted(items, key=lambda x: unicodedata.normalize("NFKD", x).casefold())
                    safe_style = "APA7" if style.upper().startswith("APA") else style.upper().replace(" ", "_")
                    out = folder / f"{stem}_REFERENCIAS_{safe_style}.docx"
                    write_simple_docx(out, f"Referências — {style}", items)
                    created.append(out)

                audit_path = folder / f"{stem}_RELATORIO_FORMATACAO_REFERENCIAS.csv"
                pd.DataFrame(metadata_rows).to_csv(audit_path, index=False, encoding="utf-8-sig")
                created.append(audit_path)

                q.put(("progress", 98, "Salvando arquivos..."))
                q.put(("done", {
                    "input": str(input_path),
                    "folder": str(folder),
                    "resolved": resolved,
                    "unresolved": unresolved,
                    "styles": styles,
                    "files": [str(x) for x in created],
                }))
            except Exception as exc:
                q.put(("error", str(exc)))

        def run():
            input_path = clean(input_var.get())
            if not input_path or not Path(input_path).exists():
                self._showwarning("Formatador", "Selecione um arquivo DOCX ou TXT válido.", parent=win)
                return
            folder = self._askdirectory(title="Pasta para salvar os novos arquivos", parent=win)
            if not folder:
                return
            selected_style = style_var.get()
            styles = ["ABNT", "Vancouver", "APA 7"] if selected_style == "TODOS" else [selected_style]
            run_btn.state(["disabled"])
            progress_var.set(0)
            progress_text.set("Iniciando...")
            append_log("--- Nova formatação ---")
            threading.Thread(
                target=worker,
                args=(input_path, folder, styles, bool(text_search_var.get())),
                name="review-tool-reference-formatter",
                daemon=True,
            ).start()
            win.after(100, poll)

        run_btn.configure(command=run)

    @staticmethod
    def _bibliometric_record_score(record):
        score = 0
        for field, weight in [("titulo",5),("autores",4),("doi",4),("resumo",4),("revista",3),("ano",2)]:
            value = clean(record.get(field, ""))
            if value:
                score += weight + min(len(value), 500) / 1000.0
        return score

    def _bibliometric_merge_dataframe(self, selected_bases, progress_callback=None):
        """Mescla bases para bibliometria com deduplicação conservadora.

        Regra alinhada ao mergeDbSources do bibliometrix: DOI idêntico; na ausência
        de DOI coincidente, título normalizado + ano de publicação idênticos.
        A proveniência combinada é guardada fora do campo DB para não fazer o
        Biblioshiny interpretar combinações de bases como bancos diferentes.
        """
        source = self.all_df[self.all_df["base"].isin(selected_bases)].copy().reset_index(drop=True)
        if source.empty:
            return pd.DataFrame(), pd.DataFrame()

        canonical = []
        provenance = []
        key_to_idx = {}
        duplicate_rows = []
        total = max(1, len(source))

        db_label = {
            "Web of Science": "ISI",
            "Scopus": "SCOPUS",
            "Embase": "EMBASE",
            "PubMed": "PUBMED",
        }

        def merge_values(target, new):
            for field in ["titulo", "autores", "resumo", "revista"]:
                tv, nv = clean(target.get(field, "")), clean(new.get(field, ""))
                if (not tv) or (nv and len(nv) > len(tv)):
                    target[field] = nv
            for field in ["ano", "doi", "id_origem", "uid"]:
                if not clean(target.get(field, "")) and clean(new.get(field, "")):
                    target[field] = clean(new.get(field, ""))

        for i, (_, row) in enumerate(source.iterrows(), 1):
            rec = {field: clean(row.get(field, "")) for field in FIELDS}
            doi = normalize_doi(rec.get("doi", ""))
            title = normalize_title(rec.get("titulo", ""))
            year = normalize_year(rec.get("ano", "")) or clean(rec.get("ano", ""))

            doi_key = f"doi:{doi}" if doi else ""
            # Bibliometrix usa título + ano para a segunda etapa. Evita fundir
            # trabalhos diferentes que por acaso tenham exatamente o mesmo título.
            title_year_key = f"titleyear:{title}|{year}" if len(title) >= 15 and year else ""

            existing = None
            criterion = ""
            if doi_key and doi_key in key_to_idx:
                existing = key_to_idx[doi_key]
                criterion = "DOI idêntico"
            elif title_year_key and title_year_key in key_to_idx:
                existing = key_to_idx[title_year_key]
                criterion = "Título normalizado + ano idênticos"

            if existing is None:
                idx = len(canonical)
                canonical.append(rec)
                provenance.append({
                    "primary_base": rec["base"],
                    "bases": [rec["base"]],
                    "arquivos": [rec["arquivo"]],
                    "ids": [rec["id_origem"]],
                    "uids": [rec["uid"]],
                })
                if doi_key:
                    key_to_idx[doi_key] = idx
                if title_year_key:
                    key_to_idx[title_year_key] = idx
            else:
                target = canonical[existing]
                prov = provenance[existing]
                previous_primary = prov["primary_base"]

                # Usa o registro globalmente mais completo como base, sem perder
                # campos complementares. Se o registro principal mudar, atualiza a
                # base primária usada em DB_Original.
                if self._bibliometric_record_score(rec) > self._bibliometric_record_score(target):
                    old = target.copy()
                    canonical[existing] = rec.copy()
                    target = canonical[existing]
                    merge_values(target, old)
                    prov["primary_base"] = rec["base"]
                else:
                    merge_values(target, rec)

                for name, value in [
                    ("bases", rec["base"]),
                    ("arquivos", rec["arquivo"]),
                    ("ids", rec["id_origem"]),
                    ("uids", rec["uid"]),
                ]:
                    if value and value not in prov[name]:
                        prov[name].append(value)

                if doi_key:
                    key_to_idx[doi_key] = existing
                if title_year_key:
                    key_to_idx[title_year_key] = existing

                duplicate_rows.append({
                    "criterio": criterion,
                    "base_removida": rec["base"],
                    "id_removido": rec["id_origem"],
                    "titulo": rec["titulo"],
                    "ano": year,
                    "doi": rec["doi"],
                    "base_primaria_mantida": prov["primary_base"],
                    "bases_em_que_aparece": "; ".join(prov["bases"]),
                })

            if progress_callback and (i == total or i % max(1, total // 100) == 0):
                progress_callback((i / total) * 75, f"Mesclando e deduplicando: {i}/{total}")

        rows = []
        for idx, rec in enumerate(canonical):
            prov = provenance[idx]
            primary_base = prov.get("primary_base") or rec.get("base", "")
            primary_db = db_label.get(primary_base, primary_base.upper() or "GENERIC")
            rows.append({
                "AU": rec.get("autores", ""),
                "TI": rec.get("titulo", ""),
                "SO": rec.get("revista", ""),
                "PY": normalize_year(rec.get("ano", "")) or rec.get("ano", ""),
                "DI": normalize_doi(rec.get("doi", "")) or rec.get("doi", ""),
                "AB": rec.get("resumo", ""),
                # Tags padrão do bibliometrix. Campos ainda não capturados pelos
                # parsers permanecem vazios; nunca são inventados.
                "DT": "", "DE": "", "ID": "", "C1": "", "RP": "",
                "CR": "", "CR_raw": "", "TC": "",
                # DB deve representar uma única base por registro. A proveniência
                # completa fica em SOURCE_DATABASES, evitando o falso aviso de 7 DBs.
                "DB": primary_db,
                "DB_Original": primary_db,
                "SOURCE_DATABASES": "; ".join(prov["bases"]),
                "UT": rec.get("id_origem", ""),
                "SOURCE_FILES": "; ".join(prov["arquivos"]),
                "SOURCE_IDS": "; ".join(prov["ids"]),
                "N_BASES": len(prov["bases"]),
                "UID_INTERNO": rec.get("uid", ""),
            })
        return pd.DataFrame(rows), pd.DataFrame(duplicate_rows)

    def _standardized_pubmed_dataframe(self):
        df = self.all_df[self.all_df["base"] == "PubMed"].copy()
        if df.empty:
            return pd.DataFrame()
        rows = []
        for _, r in df.iterrows():
            rows.append({
                "AU": clean(r.get("autores","")), "TI": clean(r.get("titulo","")),
                "SO": clean(r.get("revista","")), "PY": normalize_year(r.get("ano","")) or clean(r.get("ano","")),
                "DI": normalize_doi(r.get("doi","")) or clean(r.get("doi","")), "AB": clean(r.get("resumo","")),
                "DT": "", "DE": "", "ID": "", "C1": "", "RP": "", "CR": "", "CR_raw": "", "TC": "",
                "DB": "PUBMED", "DB_Original": "PUBMED", "SOURCE_DATABASES": "PubMed",
                "UT": clean(r.get("id_origem","")),
                "SOURCE_FILES": clean(r.get("arquivo","")), "SOURCE_IDS": clean(r.get("id_origem","")),
                "N_BASES": 1, "UID_INTERNO": clean(r.get("uid","")),
            })
        return pd.DataFrame(rows)

    @staticmethod
    def _bibtex_escape(value):
        value = clean(value)
        value = value.replace("\\", " ")
        value = value.replace("{", "").replace("}", "")
        return value

    def _write_bibtex_dataframe(self, df, output_path):
        """Grava o banco deduplicado em BibTeX genérico, compatível com convert2df(..., dbsource='generic')."""
        output_path = Path(output_path)
        lines = []
        used_keys = set()

        for index, row in df.reset_index(drop=True).iterrows():
            title = clean(row.get("TI", ""))
            authors = clean(row.get("AU", ""))
            year = clean(row.get("PY", ""))
            first = authors.split(";")[0].strip() if authors else "REF"
            surname = re.sub(r"[^A-Za-z0-9]", "", unicodedata.normalize("NFKD", first).encode("ascii", "ignore").decode("ascii"))[:18] or "REF"
            digest = hashlib.sha1((title + year + str(index)).encode("utf-8", errors="ignore")).hexdigest()[:7]
            key = f"{surname}{year}_{digest}"
            if key in used_keys:
                key = f"{key}_{index+1}"
            used_keys.add(key)

            fields = []
            if authors:
                fields.append(("author", " and ".join(a.strip() for a in authors.split(";") if a.strip())))
            if title:
                fields.append(("title", title))
            if clean(row.get("SO", "")):
                fields.append(("journal", clean(row.get("SO", ""))))
            if year:
                fields.append(("year", year))
            if clean(row.get("DI", "")):
                fields.append(("doi", normalize_doi(row.get("DI", "")) or clean(row.get("DI", ""))))
            if clean(row.get("AB", "")):
                fields.append(("abstract", clean(row.get("AB", ""))))
            if clean(row.get("DE", "")):
                fields.append(("keywords", clean(row.get("DE", ""))))
            if clean(row.get("UT", "")):
                fields.append(("note", f"UT={clean(row.get('UT',''))}; DB={clean(row.get('DB',''))}"))

            lines.append(f"@article{{{key},")
            for pos, (name, value) in enumerate(fields):
                comma = "," if pos < len(fields) - 1 else ""
                lines.append(f"  {name} = {{{self._bibtex_escape(value)}}}{comma}")
            lines.append("}\n")

        tmp = output_path.with_suffix(output_path.suffix + ".tmp")
        tmp.write_text("\n".join(lines), encoding="utf-8")
        tmp.replace(output_path)
        return output_path

    @staticmethod
    def _find_rscript():
        candidate = shutil.which("Rscript") or shutil.which("Rscript.exe")
        if candidate:
            return candidate
        if os.name == "nt":
            roots = [
                Path(os.environ.get("ProgramFiles", r"C:\\Program Files")) / "R",
                Path(os.environ.get("ProgramFiles(x86)", r"C:\\Program Files (x86)")) / "R",
            ]
            found = []
            for root in roots:
                if root.exists():
                    found.extend(root.glob("R-*\\bin\\Rscript.exe"))
                    found.extend(root.glob("R-*\\bin\\x64\\Rscript.exe"))
            if found:
                found.sort(reverse=True)
                return str(found[0])
        return ""

    @staticmethod
    def _rdata_builder_script_text():
        return r'''# Arquivo gerado pelo Systematic Review Tool
# Cria um objeto M que o Biblioshiny reconhece como bibliometrixDB.
options(stringsAsFactors = FALSE)
M <- read.csv(
  "BIBLIOMETRIA_MESCLADA.csv",
  stringsAsFactors = FALSE,
  check.names = FALSE,
  fileEncoding = "UTF-8-BOM",
  na.strings = character(0)
)

required <- c("AU","TI","SO","PY","DI","AB","DT","DE","ID","C1","RP","CR","CR_raw","TC","DB","DB_Original","SOURCE_DATABASES","UT")
for (nm in required) {
  if (!nm %in% names(M)) M[[nm]] <- ""
  M[[nm]][is.na(M[[nm]])] <- ""
}

if (!"DB_Original" %in% names(M)) M$DB_Original <- M$DB
# Em coleção já mesclada, o bibliometrix usa DB="ISI" e mantém a origem
# em DB_Original. SOURCE_DATABASES guarda a proveniência completa.
M$DB <- "ISI"
M$PY <- suppressWarnings(as.numeric(M$PY))
M$TC <- suppressWarnings(as.numeric(ifelse(M$TC == "", NA, M$TC)))

first_author <- trimws(sub(";.*$", "", M$AU))
first_author[first_author == ""] <- "NA"
year_text <- ifelse(is.na(M$PY), "NA", as.character(M$PY))
source_text <- ifelse(M$SO == "", "NA", M$SO)
M$SR_FULL <- trimws(paste(first_author, year_text, source_text, sep = ", "))
M$SR <- make.unique(M$SR_FULL, sep = "-")
M$KW_Merged <- trimws(gsub("^;+|;+$", "", paste(M$DE, M$ID, sep = ";")))
rownames(M) <- M$SR
class(M) <- c("bibliometrixDB", "data.frame")

save(M, file = "BIBLIOSHINY_PRONTO.RData", compress = TRUE)
cat("OK:", nrow(M), "registros salvos em BIBLIOSHINY_PRONTO.RData\n")
'''

    def _generate_biblioshiny_rdata(self, folder, progress_callback=None):
        folder = Path(folder)
        script_path = folder / "GERAR_BIBLIOSHINY_RDATA.R"
        script_path.write_text(self._rdata_builder_script_text(), encoding="utf-8")
        rscript = self._find_rscript()
        rdata = folder / "BIBLIOSHINY_PRONTO.RData"

        if not rscript:
            return {
                "created": False,
                "path": str(rdata),
                "script": str(script_path),
                "warning": "Rscript não foi localizado. O script GERAR_BIBLIOSHINY_RDATA.R foi criado para execução manual.",
            }

        if progress_callback:
            progress_callback(94, "Gerando BIBLIOSHINY_PRONTO.RData...")
        try:
            kwargs = {
                "cwd": str(folder),
                "capture_output": True,
                "text": True,
                "timeout": 120,
            }
            if os.name == "nt":
                kwargs["creationflags"] = getattr(subprocess, "CREATE_NO_WINDOW", 0)
            result = subprocess.run([rscript, script_path.name], **kwargs)
            if result.returncode != 0:
                detail = clean(result.stderr) or clean(result.stdout) or f"Rscript retornou código {result.returncode}"
                return {"created": False, "path": str(rdata), "script": str(script_path), "warning": detail}
            if not rdata.exists():
                return {"created": False, "path": str(rdata), "script": str(script_path), "warning": "O R terminou sem erro, mas o arquivo RData não foi encontrado."}
            return {"created": True, "path": str(rdata), "script": str(script_path), "warning": ""}
        except subprocess.TimeoutExpired:
            return {"created": False, "path": str(rdata), "script": str(script_path), "warning": "A geração do RData excedeu 120 segundos e foi interrompida para evitar travamento."}
        except Exception as exc:
            return {"created": False, "path": str(rdata), "script": str(script_path), "warning": str(exc)}

    def _launch_biblioshiny(self, rdata_path, parent=None):
        rscript = self._find_rscript()
        if not rscript:
            self._showerror(
                "Biblioshiny",
                "Rscript não foi localizado. Instale o R e o pacote bibliometrix ou adicione o R ao PATH.",
                parent=parent,
            )
            return False
        rdata_path = Path(rdata_path)
        if not rdata_path.exists():
            self._showerror("Biblioshiny", "O arquivo BIBLIOSHINY_PRONTO.RData ainda não existe.", parent=parent)
            return False

        command = (
            "if (!requireNamespace('bibliometrix', quietly=TRUE)) "
            "stop('Pacote bibliometrix nao instalado. Use install.packages(\"bibliometrix\")'); "
            "bibliometrix::biblioshiny()"
        )
        kwargs = {"cwd": str(rdata_path.parent)}
        if os.name == "nt":
            kwargs["creationflags"] = getattr(subprocess, "CREATE_NEW_CONSOLE", 0)
        subprocess.Popen([rscript, "-e", command], **kwargs)

        try:
            self.clipboard_clear()
            self.clipboard_append(str(rdata_path))
            self.update_idletasks()
        except Exception:
            pass

        if os.name == "nt":
            try:
                subprocess.Popen(["explorer", "/select,", str(rdata_path)])
            except Exception:
                pass

        self._showinfo(
            "Biblioshiny iniciado",
            "O Biblioshiny foi iniciado.\n\n"
            "O arquivo pronto está selecionado no Explorador e o caminho foi copiado. "
            "No Biblioshiny use Data > Import or Load > Load files e escolha BIBLIOSHINY_PRONTO.RData.\n\n"
            "O Biblioshiny oficial não oferece um argumento para pré-carregar um arquivo ao iniciar; "
            "por isso essa última seleção é necessária.",
            parent=parent,
        )
        return True

    def _run_bibliometric_export(self, selected_bases, export_pubmed, folder, progress_callback=None):
        folder = Path(folder)
        folder.mkdir(parents=True, exist_ok=True)
        created = []
        warnings = []
        merged = duplicates = pd.DataFrame()
        rdata_info = {"created": False, "path": "", "warning": ""}

        if selected_bases:
            if progress_callback:
                progress_callback(2, "Preparando bases selecionadas...")
            merged, duplicates = self._bibliometric_merge_dataframe(selected_bases, progress_callback)
            if merged.empty:
                raise RuntimeError("Nenhum registro foi encontrado nas bases selecionadas.")

            xlsx = folder / "BIBLIOMETRIA_MESCLADA.xlsx"
            csv_path = folder / "BIBLIOMETRIA_MESCLADA.csv"
            bib_path = folder / "BIBLIOMETRIA_MESCLADA.bib"

            if progress_callback:
                progress_callback(78, "Salvando CSV mesclado...")
            tmp_csv = csv_path.with_suffix(".csv.tmp")
            merged.to_csv(tmp_csv, index=False, encoding="utf-8-sig")
            tmp_csv.replace(csv_path)
            created.append(csv_path)

            if progress_callback:
                progress_callback(82, "Salvando BibTeX mesclado e deduplicado...")
            self._write_bibtex_dataframe(merged, bib_path)
            created.append(bib_path)

            if progress_callback:
                progress_callback(86, "Salvando Excel de auditoria...")
            summary = pd.DataFrame([
                ("Bases mescladas", "; ".join(selected_bases)),
                ("Registros de entrada", int(self.all_df["base"].isin(selected_bases).sum())),
                ("Duplicatas exatas removidas", int(len(duplicates))),
                ("Registros finais", int(len(merged))),
                ("Critério automático", "DOI idêntico ou título normalizado + ano idênticos"),
                ("Fuzzy", "Não utilizado para exclusão neste módulo"),
                ("Arquivo Biblioshiny", "BIBLIOSHINY_PRONTO.RData (quando Rscript estiver disponível)"),
                ("Gerado em", datetime.now().isoformat(timespec="seconds")),
            ], columns=["Etapa", "Valor"])

            completeness_rows = []
            for tag, description in [
                ("SO", "Journal"), ("PY", "Publication Year"), ("TI", "Title"),
                ("AU", "Author"), ("DI", "DOI"), ("AB", "Abstract"),
                ("C1", "Affiliation"), ("CR", "Cited References"),
                ("RP", "Corresponding Author"), ("DE", "Author Keywords"),
                ("ID", "Keywords Plus / indexed keywords"),
            ]:
                if tag in merged.columns:
                    vals = merged[tag].fillna("").astype(str).str.strip()
                    missing = int((vals == "").sum())
                    pct = round((missing / max(1, len(merged))) * 100, 2)
                    completeness_rows.append((tag, description, missing, pct))
            completeness = pd.DataFrame(
                completeness_rows,
                columns=["Metadata", "Description", "Missing Counts", "Missing %"],
            )

            tmp_xlsx = folder / "BIBLIOMETRIA_MESCLADA.__tmp__.xlsx"
            with pd.ExcelWriter(tmp_xlsx, engine="openpyxl") as writer:
                merged.to_excel(writer, sheet_name="BASE_MESCLADA", index=False)
                summary.to_excel(writer, sheet_name="RESUMO", index=False)
                duplicates.to_excel(writer, sheet_name="DUPLICATAS_REMOVIDAS", index=False)
                completeness.to_excel(writer, sheet_name="COMPLETUDE_METADATA", index=False)
                from openpyxl.styles import PatternFill, Font
                blue = PatternFill(fill_type="solid", fgColor="0B4F8A")
                for ws in writer.book.worksheets:
                    ws.freeze_panes = "A2"
                    ws.auto_filter.ref = ws.dimensions
                    for cell in ws[1]:
                        cell.fill = blue
                        cell.font = Font(bold=True, color="FFFFFF")
                    for letter, width in {
                        "A": 44, "B": 70, "C": 34, "D": 12, "E": 34, "F": 85,
                        "G": 18, "H": 30, "I": 30, "J": 34, "K": 34, "L": 70,
                    }.items():
                        ws.column_dimensions[letter].width = width
            tmp_xlsx.replace(xlsx)
            created.append(xlsx)

            # Alertas objetivos de completude. Não bloqueiam a exportação.
            for tag, label in [("AB", "resumos"), ("C1", "afiliações"), ("DE", "palavras-chave dos autores"), ("ID", "palavras-chave indexadas")]:
                if tag in merged.columns and len(merged):
                    missing = int((merged[tag].fillna("").astype(str).str.strip() == "").sum())
                    if missing == len(merged):
                        warnings.append(f"{tag}: {label} estão ausentes em 100% dos registros; o parser/formato de origem ainda não trouxe esse campo.")

            if progress_callback:
                progress_callback(92, "Preparando arquivo direto para Biblioshiny...")
            rdata_info = self._generate_biblioshiny_rdata(folder, progress_callback)
            script_path = Path(rdata_info.get("script", ""))
            if script_path.exists():
                created.append(script_path)
            if rdata_info.get("created") and Path(rdata_info["path"]).exists():
                created.append(Path(rdata_info["path"]))
            elif rdata_info.get("warning"):
                warnings.append(rdata_info["warning"])

            r_script = folder / "IMPORTAR_BIBLIOMETRIX.R"
            r_script.write_text(
                '# Base padronizada gerada pelo Systematic Review Tool\n'
                'library(bibliometrix)\n'
                'load("BIBLIOSHINY_PRONTO.RData")\n'
                'results <- biblioAnalysis(M, sep = ";")\n'
                'summary(results, k = 10, pause = FALSE)\n',
                encoding="utf-8",
            )
            created.append(r_script)

        pubmed = pd.DataFrame()
        if export_pubmed:
            pubmed = self._standardized_pubmed_dataframe()
            if not pubmed.empty:
                if progress_callback:
                    progress_callback(97, "Salvando PubMed separadamente...")
                px = folder / "PUBMED_PADRONIZADO.xlsx"
                pc = folder / "PUBMED_PADRONIZADO.csv"
                pb = folder / "PUBMED_PADRONIZADO.bib"
                pubmed.to_excel(px, index=False, engine="openpyxl")
                pubmed.to_csv(pc, index=False, encoding="utf-8-sig")
                self._write_bibtex_dataframe(pubmed, pb)
                created += [px, pc, pb]

        if progress_callback:
            progress_callback(100, "Concluído")
        summary_dict = {
            "bases": list(selected_bases),
            "entrada_mesclagem": int(self.all_df["base"].isin(selected_bases).sum()) if selected_bases else 0,
            "duplicatas_removidas": int(len(duplicates)) if not duplicates.empty else 0,
            "final_mesclado": int(len(merged)) if not merged.empty else 0,
            "pubmed_separado": int(len(pubmed)) if not pubmed.empty else 0,
            "pasta": str(folder),
            "arquivos": [str(p) for p in created],
            "bib": str(folder / "BIBLIOMETRIA_MESCLADA.bib") if selected_bases else "",
            "rdata": rdata_info.get("path", ""),
            "rdata_created": bool(rdata_info.get("created")),
            "warnings": warnings,
        }
        self.last_biblio_merge_summary = summary_dict
        return summary_dict

    def open_bibliometric_merge_window(self):
        if self.all_df.empty:
            self._showwarning("Bibliometria", "Processe as bases antes de realizar a mesclagem bibliométrica.")
            return
        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("Mesclagem e arquivo pronto para Biblioshiny")
        win.geometry("860x760")
        win.minsize(760, 650)

        info = ttk.LabelFrame(win, text="Objetivo", padding=10)
        info.pack(fill="x", padx=12, pady=10)
        ttk.Label(info, text=(
            "Mescla Web of Science, Embase e/ou Scopus, remove duplicatas por DOI idêntico ou título + ano idênticos e gera CSV, Excel e BibTeX. "
            "Também tenta criar BIBLIOSHINY_PRONTO.RData para você carregar diretamente no Biblioshiny sem refazer a mesclagem."
        ), wraplength=810, justify="left").pack(anchor="w")
        ttk.Label(info, text=(
            "Importante: o programa não inventa metadados. Palavras-chave, referências citadas, afiliações e citações só estarão "
            "disponíveis se esses campos tiverem sido importados do arquivo original."
        ), wraplength=810, justify="left").pack(anchor="w", pady=(6, 0))

        bases_box = ttk.LabelFrame(win, text="Bases para a mesclagem principal", padding=10)
        bases_box.pack(fill="x", padx=12, pady=6)
        base_vars = {}
        available = set(self.all_df["base"].astype(str).tolist())
        for db in ["Web of Science", "Embase", "Scopus"]:
            var = tk.BooleanVar(value=db in available)
            base_vars[db] = var
            cb = ttk.Checkbutton(
                bases_box,
                text=f"{db}  ({int((self.all_df['base'] == db).sum())} registros)",
                variable=var,
            )
            cb.pack(anchor="w", pady=2)
            if db not in available:
                cb.state(["disabled"])

        pubmed_var = tk.BooleanVar(value="PubMed" in available)
        pbox = ttk.LabelFrame(win, text="PubMed", padding=10)
        pbox.pack(fill="x", padx=12, pady=6)
        pcb = ttk.Checkbutton(
            pbox,
            text=f"Exportar PubMed separadamente ({int((self.all_df['base'] == 'PubMed').sum())} registros)",
            variable=pubmed_var,
        )
        pcb.pack(anchor="w")
        if "PubMed" not in available:
            pcb.state(["disabled"])

        launch_var = tk.BooleanVar(value=False)
        launch_box = ttk.LabelFrame(win, text="Biblioshiny", padding=10)
        launch_box.pack(fill="x", padx=12, pady=6)
        ttk.Checkbutton(
            launch_box,
            text="Abrir o Biblioshiny ao terminar (requer R + pacote bibliometrix)",
            variable=launch_var,
        ).pack(anchor="w")
        ttk.Label(
            launch_box,
            text="O caminho do RData será copiado e o arquivo será selecionado no Explorador para facilitar a última etapa de carga.",
            wraplength=800,
        ).pack(anchor="w", pady=(4, 0))

        progress_var = tk.DoubleVar(value=0)
        ttk.Progressbar(win, variable=progress_var, maximum=100, style="Blue.Horizontal.TProgressbar").pack(fill="x", padx=12, pady=(12, 4))
        progress_text = tk.StringVar(value="Aguardando")
        ttk.Label(win, textvariable=progress_text).pack(anchor="w", padx=12)

        log = tk.Text(win, height=13, wrap="word", font=("Consolas", 9))
        log.pack(fill="both", expand=True, padx=12, pady=8)
        log.insert("end", "Selecione as bases e clique em REALIZAR MESCLAGEM.\n")
        log.configure(state="disabled")

        action = ttk.Frame(win, padding=(12, 4, 12, 12))
        action.pack(fill="x")
        open_btn = ttk.Button(action, text="ABRIR BIBLIOSHINY", state="disabled")
        open_btn.pack(side="left")
        run_btn = ttk.Button(action, text="REALIZAR MESCLAGEM")
        run_btn.pack(side="right")

        q = queue.Queue()

        def append_log(msg):
            log.configure(state="normal")
            log.insert("end", self.translate_literal(msg) + "\n")
            log.see("end")
            log.configure(state="disabled")

        def poll():
            done = False
            while True:
                try:
                    ev = q.get_nowait()
                except queue.Empty:
                    break
                if ev[0] == "progress":
                    progress_var.set(ev[1])
                    progress_text.set(self.translate_literal(ev[2]))
                    append_log(ev[2])
                elif ev[0] == "error":
                    done = True
                    run_btn.state(["!disabled"])
                    self._showerror("Bibliometria", ev[1], parent=win)
                elif ev[0] == "done":
                    done = True
                    run_btn.state(["!disabled"])
                    s = ev[1]
                    append_log(f"Registros de entrada na mesclagem: {s['entrada_mesclagem']}")
                    append_log(f"Duplicatas exatas removidas: {s['duplicatas_removidas']}")
                    append_log(f"Registros finais salvos: {s['final_mesclado']}")
                    if s["bib"]:
                        append_log(f"BibTeX pronto: {s['bib']}")
                    if s["rdata_created"]:
                        append_log(f"RData pronto para Biblioshiny: {s['rdata']}")
                        open_btn.state(["!disabled"])
                        open_btn.configure(command=lambda p=s["rdata"]: self._launch_biblioshiny(p, parent=win))
                    else:
                        append_log("RData não foi criado automaticamente.")
                    for warning in s.get("warnings", []):
                        append_log(f"AVISO: {warning}")
                    if s["pubmed_separado"]:
                        append_log(f"PubMed salvo separadamente: {s['pubmed_separado']} registros")
                    append_log(f"Pasta: {s['pasta']}")
                    self.log_audit_event(
                        "MESCLAGEM BIBLIOMÉTRICA",
                        f"Bases={'; '.join(s['bases'])}; entrada={s['entrada_mesclagem']}; duplicatas={s['duplicatas_removidas']}; final={s['final_mesclado']}; RData={s['rdata_created']}",
                    )
                    if launch_var.get() and s["rdata_created"]:
                        self._launch_biblioshiny(s["rdata"], parent=win)
                    warning_text = "\n".join(s.get("warnings", []))
                    message = (
                        "Mesclagem concluída.\n\n"
                        f"Entrada: {s['entrada_mesclagem']}\n"
                        f"Duplicatas removidas: {s['duplicatas_removidas']}\n"
                        f"Registros finais: {s['final_mesclado']}\n"
                        f"BibTeX: {'SIM' if s['bib'] else 'NÃO'}\n"
                        f"RData Biblioshiny: {'SIM' if s['rdata_created'] else 'NÃO'}\n\n"
                        f"Pasta: {s['pasta']}"
                    )
                    if warning_text:
                        message += "\n\nAviso:\n" + warning_text
                    self._showinfo("Bibliometria", message, parent=win)
            if not done and run_btn.instate(["disabled"]):
                win.after(80, poll)

        def run():
            selected = [db for db, var in base_vars.items() if var.get()]
            export_pubmed = bool(pubmed_var.get())
            if not selected and not export_pubmed:
                self._showwarning("Bibliometria", "Selecione pelo menos uma base ou o PubMed separado.", parent=win)
                return
            folder = self._askdirectory(title="Pasta para salvar a mesclagem", parent=win)
            if not folder:
                return
            run_btn.state(["disabled"])
            open_btn.state(["disabled"])
            progress_var.set(0)
            progress_text.set("Iniciando...")
            append_log("--- Nova execução ---")

            def cb(p, m):
                q.put(("progress", p, m))

            def worker():
                try:
                    q.put(("done", self._run_bibliometric_export(selected, export_pubmed, folder, cb)))
                except Exception as exc:
                    q.put(("error", str(exc)))

            threading.Thread(target=worker, name="review-tool-bibliometry", daemon=True).start()
            win.after(80, poll)

        run_btn.configure(command=run)


    def sync_triage_df(self):
        """Sincroniza a triagem com o banco deduplicado, preservando decisões por UID."""
        if self.final_df is None or self.final_df.empty:
            self.triage_df = pd.DataFrame(columns=TRIAGE_FIELDS)
            return

        previous = {}
        if self.triage_df is not None and not self.triage_df.empty:
            for _, row in self.triage_df.iterrows():
                previous[clean(row.get("uid", ""))] = (
                    clean(row.get("triagem_status", "")) or "NÃO AVALIADO",
                    clean(row.get("motivo_exclusao", "")),
                )

        rows = []
        for _, row in self.final_df.iterrows():
            item = {field: clean(row.get(field, "")) for field in FIELDS}
            status, reason = previous.get(item["uid"], ("NÃO AVALIADO", ""))
            item["triagem_status"] = status
            item["motivo_exclusao"] = reason
            rows.append(item)

        self.triage_df = pd.DataFrame(rows, columns=TRIAGE_FIELDS)

    @staticmethod
    def _df_records(df):
        if df is None or df.empty:
            return []
        return df.fillna("").astype(str).to_dict(orient="records")

    @staticmethod
    def _records_df(records, columns):
        df = pd.DataFrame(records or [])
        for column in columns:
            if column not in df.columns:
                df[column] = ""
        if df.empty:
            return pd.DataFrame(columns=columns)
        return df[columns].fillna("")

    def save_project(self, save_as=False):
        if self.all_df.empty:
            self._showwarning("Projeto", "Processe pelo menos uma base antes de salvar o projeto.")
            return

        path = self.current_project_path
        if save_as or not path:
            path = self._asksaveasfilename(
                title="Salvar projeto",
                defaultextension=".srtproj",
                filetypes=[("Projeto BioReviewPy", "*.srtproj"), ("JSON", "*.json")],
            )
        if not path:
            return

        self.sync_triage_df()
        if self.second_screen_applied:
            self.sync_review_filter_df()
        self.log_audit_event("PROJETO SALVO", str(path))
        payload = {
            "app": APP_TITLE,
            "version": VERSION,
            "project_format": 2,
            "saved_at": datetime.now().isoformat(timespec="seconds"),
            "threshold": int(self.threshold_var.get()),
            "language": self.language_code,
            "file_entries": self.file_entries,
            "all_df": self._df_records(self.all_df),
            "base_final_df": self._df_records(self.base_final_df),
            "final_df": self._df_records(self.final_df),
            "comparison_df": self._df_records(self.comparison_df),
            "comparison_order": list(self.comparison_order),
            "triage_df": self._df_records(self.triage_df),  # compatibilidade com projetos v0.9
            "review_filter_df": self._df_records(self.review_filter_df),
            "second_screen_applied": bool(self.second_screen_applied),
            "prisma_extra": dict(self.prisma_extra),
            "audit_events": list(self.audit_events),
            "last_biblio_merge_summary": dict(self.last_biblio_merge_summary),
        }

        try:
            serialized = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
            # .srtproj v2: ZIP compacto contendo project.json.
            # Projetos JSON antigos continuam compatíveis no open_project().
            if str(path).lower().endswith(".json"):
                Path(path).write_text(serialized, encoding="utf-8")
            else:
                with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
                    zf.writestr("project.json", serialized.encode("utf-8"))
        except Exception as exc:
            self._showerror("Projeto", f"Não foi possível salvar o projeto.\n\n{exc}")
            return

        self.current_project_path = str(path)
        self.status_var.set(self.translate_literal(f"Projeto salvo: {Path(path).name}"))
        self._showinfo("Projeto", "Projeto salvo com sucesso.")

    def open_project(self):
        path = self._askopenfilename(
            title="Abrir projeto",
            filetypes=[("Projeto BioReviewPy", "*.srtproj *.json"), ("Todos", "*.*")],
        )
        if not path:
            return

        try:
            if zipfile.is_zipfile(path):
                with zipfile.ZipFile(path, "r") as zf:
                    name = "project.json" if "project.json" in zf.namelist() else zf.namelist()[0]
                    payload = json.loads(zf.read(name).decode("utf-8"))
            else:
                # Compatibilidade com .srtproj/.json das versões anteriores.
                payload = json.loads(Path(path).read_text(encoding="utf-8"))
        except Exception as exc:
            self._showerror("Projeto", f"Arquivo de projeto inválido.\n\n{exc}")
            return

        try:
            self.file_entries = payload.get("file_entries", []) or []
            self.all_df = self._records_df(payload.get("all_df", []), FIELDS)
            self.base_final_df = self._records_df(payload.get("base_final_df", []), FIELDS)
            self.final_df = self._records_df(payload.get("final_df", []), FIELDS)
            self.comparison_df = self._records_df(payload.get("comparison_df", []), COMPARISON_FIELDS)
            self.comparison_order = payload.get("comparison_order", []) or []
            self.triage_df = self._records_df(payload.get("triage_df", []), TRIAGE_FIELDS)
            self.review_filter_df = self._records_df(payload.get("review_filter_df", []), SECOND_SCREEN_FIELDS)
            self.second_screen_applied = bool(payload.get("second_screen_applied", False))
            self.prisma_extra = payload.get("prisma_extra", {}) or {}
            self.audit_events = payload.get("audit_events", []) or []
            self.last_biblio_merge_summary = payload.get("last_biblio_merge_summary", {}) or {}

            saved_language = payload.get("language", self.language_code)
            if saved_language in set(LANGUAGE_OPTIONS.values()) and saved_language != self.language_code:
                self.language_var.set(self._language_name_from_code(saved_language))
                self.change_language()

            for key, default in {
                "records_excluded_title_abstract": 0,
                "reports_not_retrieved": 0,
                "full_text_excluded": 0,
                "studies_included": 0,
                "full_text_reasons": "",
            }.items():
                self.prisma_extra.setdefault(key, default)

            self.threshold_var.set(int(payload.get("threshold", 95)))
            self.base_data = {}
            if not self.all_df.empty:
                for database, group in self.all_df.groupby("base", sort=False):
                    self.base_data[database] = group.reset_index(drop=True).copy()

            self.files_tree.delete(*self.files_tree.get_children())
            for entry in self.file_entries:
                file_path = entry.get("path", "")
                database = entry.get("database", "Outra")
                style = entry.get("style", "Projeto salvo")
                self.files_tree.insert(
                    "", "end",
                    values=(self.display_database_name(database), file_path, Path(file_path).suffix.lower(), style)
                )

            counts = self.all_df.groupby("base").size().sort_values(ascending=False) if not self.all_df.empty else pd.Series(dtype=int)
            if not counts.empty:
                self.reference_var.set(
                    f"{self.display_database_name(counts.index[0])} ({int(counts.iloc[0])} {self.t('records')})"
                )
            else:
                self.reference_var.set("-")

            self.sync_triage_df()
            if self.second_screen_applied:
                self.sync_review_filter_df()
            self.fill_comparison_tree()
            self.update_summary()
            self.current_project_path = str(path)
            self.log_audit_event("PROJETO ABERTO", str(path))
            self.status_var.set(self.translate_literal(f"Projeto aberto: {Path(path).name}"))
            loaded_text = {"pt":"Projeto carregado", "en":"Project loaded", "es":"Proyecto cargado"}.get(self.language_code, "Projeto carregado")
            self.set_progress(100 if not self.all_df.empty else 0, loaded_text if not self.all_df.empty else self.t("waiting_processing"))
        except Exception as exc:
            self._showerror("Projeto", f"Não foi possível restaurar o projeto.\n\n{exc}")
            return

        self._showinfo("Projeto", "Projeto restaurado com sucesso.")

    def show_evidence_details(self):
        if self.comparison_df.empty:
            return
        selected = self.comparison_tree.selection()
        if not selected:
            self._showinfo("Evidência", "Selecione uma comparação na grade.")
            return
        tags = self.comparison_tree.item(selected[0], "tags")
        if not tags:
            return
        comp_id = tags[0]
        subset = self.comparison_df[self.comparison_df["comparacao_id"] == comp_id]
        if subset.empty:
            return
        row = subset.iloc[0]
        reference = {
            "titulo": row.get("titulo_referencia", ""),
            "doi": row.get("doi_referencia", ""),
            "autores": row.get("autores_referencia", ""),
            "ano": row.get("ano_referencia", ""),
        }
        candidate = {
            "titulo": row.get("titulo_comparado", ""),
            "doi": row.get("doi_comparado", ""),
            "autores": row.get("autores_comparado", ""),
            "ano": row.get("ano_comparado", ""),
        }
        ev = build_evidence(reference, candidate, row.get("similaridade", 0) or 0)
        doi_text = "idêntico" if ev["doi_same"] else (
            "diferente" if normalize_doi(reference["doi"]) and normalize_doi(candidate["doi"]) else "ausente em pelo menos um registro"
        )
        message = (
            f"Comparação: {comp_id}\n"
            f"Status: {clean(row.get('status', ''))}\n"
            f"Critério: {clean(row.get('criterio', ''))}\n"
            f"Evidência compacta: {clean(row.get('evidencia', ''))}\n\n"
            f"Título: {ev['title_similarity']:.1f}% de similaridade\n"
            f"DOI: {doi_text}\n"
            f"Primeiro autor: {ev['author_state']}\n"
            f"Ano: {ev['year_state']}\n\n"
            f"JÁ EXISTENTE\n{clean(row.get('titulo_referencia', ''))}\n\n"
            f"NOVO REGISTRO\n{clean(row.get('titulo_comparado', ''))}"
        )
        self._showinfo("Evidências da comparação", message)

    def _build_review_filter_dataframe(self, previous_df=None):
        """Monta a prévia classificatória da 2ª triagem; nenhum registro é excluído."""
        if self.final_df is None or self.final_df.empty:
            return pd.DataFrame(columns=SECOND_SCREEN_FIELDS)

        previous = {}
        if previous_df is not None and not previous_df.empty:
            for _, row in previous_df.iterrows():
                uid = clean(row.get("uid", ""))
                if uid:
                    previous[uid] = (
                        clean(row.get("decisao_2triagem", "")),
                        clean(row.get("origem_decisao_2triagem", "")),
                        clean(row.get("tipo_detectado", "")),
                    )

        rows = []
        for _, row in self.final_df.iterrows():
            classification = classify_review_meta_record(row)
            if classification is None:
                continue

            item = {field: clean(row.get(field, "")) for field in FIELDS}
            item.update({
                "tipo_detectado": classification["tipo_detectado"],
                "fonte_deteccao": classification["fonte_deteccao"],
                "evidencia_2triagem": classification["evidencia_2triagem"],
            })

            automatic = classification["decisao_automatica"]
            old_decision, old_origin, old_type = previous.get(item["uid"], ("", "", ""))
            old_decision = normalize_second_screen_decision(old_decision, old_type or item["tipo_detectado"])

            if old_origin == "MANUAL" and old_decision in {
                "CONFIRMADO", "REVISAR", "NÃO É REVISÃO/META"
            }:
                item["decisao_2triagem"] = old_decision
                item["origem_decisao_2triagem"] = "MANUAL"
                if old_type and old_decision == "CONFIRMADO":
                    item["tipo_detectado"] = old_type
            else:
                item["decisao_2triagem"] = automatic
                item["origem_decisao_2triagem"] = "AUTOMÁTICA"
            rows.append(item)

        return pd.DataFrame(rows, columns=SECOND_SCREEN_FIELDS)

    def sync_review_filter_df(self):
        """Recalcula a 2ª triagem já confirmada e preserva classificações manuais."""
        if not getattr(self, "second_screen_applied", False) or self.final_df is None or self.final_df.empty:
            self.review_filter_df = pd.DataFrame(columns=SECOND_SCREEN_FIELDS)
            return

        self.review_filter_df = self._build_review_filter_dataframe(self.review_filter_df)

    def post_second_screen_df(self):
        """Retorna o banco deduplicado. A 2ª triagem apenas marca; nunca remove registros."""
        if self.final_df is None or self.final_df.empty:
            return pd.DataFrame(columns=FIELDS)
        if getattr(self, "second_screen_applied", False):
            self.sync_review_filter_df()
        return self.final_df.copy().reset_index(drop=True)

    def second_screen_stats(self):
        """Contagens classificatórias da 2ª triagem, sem alterar o tamanho do banco."""
        stats = {
            "detected": 0,
            "confirmed": 0,
            "pending": 0,
            "dismissed": 0,
            "type_counts": {},
            "base_counts": {},
        }
        if not getattr(self, "second_screen_applied", False):
            return stats
        self.sync_review_filter_df()
        if self.review_filter_df.empty:
            return stats

        df = self.review_filter_df.copy()
        df["_decision"] = [
            normalize_second_screen_decision(d, t)
            for d, t in zip(df["decisao_2triagem"], df["tipo_detectado"])
        ]
        stats["detected"] = int(len(df))
        stats["confirmed"] = int((df["_decision"] == "CONFIRMADO").sum())
        stats["pending"] = int((df["_decision"] == "REVISAR").sum())
        stats["dismissed"] = int((df["_decision"] == "NÃO É REVISÃO/META").sum())

        confirmed_df = df[df["_decision"] == "CONFIRMADO"]
        if not confirmed_df.empty:
            stats["type_counts"] = {
                clean(k): int(v)
                for k, v in confirmed_df.groupby("tipo_detectado").size().to_dict().items()
            }
            stats["base_counts"] = {
                clean(k): int(v)
                for k, v in confirmed_df.groupby("base").size().to_dict().items()
            }
        return stats

    def _second_screen_annotation_maps(self):
        """Mapas de anotação por UID/DOI/título para propagar a marcação às bases originais."""
        maps = {"uid": {}, "doi": {}, "title": {}}
        if not getattr(self, "second_screen_applied", False):
            return maps
        self.sync_review_filter_df()
        if self.review_filter_df.empty:
            return maps

        for _, row in self.review_filter_df.iterrows():
            decision = normalize_second_screen_decision(
                row.get("decisao_2triagem", ""), row.get("tipo_detectado", "")
            )
            annotation = {
                "status": decision,
                "classificacao": clean(row.get("tipo_detectado", "")),
                "fonte": clean(row.get("fonte_deteccao", "")),
                "evidencia": clean(row.get("evidencia_2triagem", "")),
                "origem": clean(row.get("origem_decisao_2triagem", "")),
            }
            uid = clean(row.get("uid", ""))
            doi = normalize_doi(row.get("doi", ""))
            title = normalize_title(row.get("titulo", ""))
            if uid:
                maps["uid"][uid] = annotation
            if doi and doi not in maps["doi"]:
                maps["doi"][doi] = annotation
            if title and title not in maps["title"]:
                maps["title"][title] = annotation
        return maps

    @staticmethod
    def _second_screen_annotation_for_record(record, maps):
        uid = clean(record.get("uid", ""))
        if uid and uid in maps.get("uid", {}):
            return maps["uid"][uid]
        doi = normalize_doi(record.get("doi", ""))
        if doi and doi in maps.get("doi", {}):
            return maps["doi"][doi]
        title = normalize_title(record.get("titulo", ""))
        if title and title in maps.get("title", {}):
            return maps["title"][title]
        return {"status": "", "classificacao": "", "fonte": "", "evidencia": "", "origem": ""}

    def _annotate_export_dataframe(self, df):
        """Anexa as colunas classificatórias da 2ª triagem sem remover linhas."""
        out = df.copy()
        maps = self._second_screen_annotation_maps()
        annotations = [
            self._second_screen_annotation_for_record(row, maps)
            for _, row in out.iterrows()
        ]
        out["status_2triagem_export"] = [a["status"] for a in annotations]
        out["classificacao_2triagem_export"] = [a["classificacao"] for a in annotations]
        out["fonte_2triagem_export"] = [a["fonte"] for a in annotations]
        out["evidencia_2triagem_export"] = [a["evidencia"] for a in annotations]
        out["origem_2triagem_export"] = [a["origem"] for a in annotations]
        return out

    def open_second_screen_window(self):
        if self.final_df.empty:
            self._showwarning(
                "2ª triagem",
                "Processe e deduplique as bases antes de executar a 2ª triagem."
            )
            return

        pending_duplicates = 0
        if self.comparison_df is not None and not self.comparison_df.empty:
            pending_duplicates = int((self.comparison_df["status"] == "REVISAR").sum())
        if pending_duplicates:
            self._showwarning(
                "2ª triagem",
                f"Ainda existem {pending_duplicates} possível(is) duplicata(s) aguardando decisão.\n\n"
                "Resolva primeiro a conferência de duplicatas; depois execute a 2ª triagem."
            )
            return

        previous_confirmed = (
            self.review_filter_df.copy()
            if getattr(self, "second_screen_applied", False)
            else pd.DataFrame(columns=SECOND_SCREEN_FIELDS)
        )
        working_df = self._build_review_filter_dataframe(previous_confirmed)

        # Textos específicos desta janela. Os valores internos permanecem em
        # português para manter compatibilidade com projetos salvos anteriormente;
        # apenas a apresentação ao usuário é traduzida.
        def ss(pt, en, es):
            return {"pt": pt, "en": en, "es": es}.get(self.language_code, pt)

        second_value_map = {
            "CONFIRMADO": {"en": "CONFIRMED", "es": "CONFIRMADO"},
            "REVISAR": {"en": "REVIEW", "es": "REVISAR"},
            "NÃO É REVISÃO/META": {"en": "NOT A REVIEW/META-ANALYSIS", "es": "NO ES REVISIÓN/METAANÁLISIS"},
            "REVISÃO SISTEMÁTICA": {"en": "SYSTEMATIC REVIEW", "es": "REVISIÓN SISTEMÁTICA"},
            "META-ANÁLISE": {"en": "META-ANALYSIS", "es": "METAANÁLISIS"},
            "REVISÃO SISTEMÁTICA + META-ANÁLISE": {"en": "SYSTEMATIC REVIEW + META-ANALYSIS", "es": "REVISIÓN SISTEMÁTICA + METAANÁLISIS"},
            "REVISÃO DE REVISÕES": {"en": "OVERVIEW OF REVIEWS", "es": "REVISIÓN DE REVISIONES"},
            "POSSÍVEL REVISÃO SISTEMÁTICA/META-ANÁLISE": {"en": "POSSIBLE SYSTEMATIC REVIEW/META-ANALYSIS", "es": "POSIBLE REVISIÓN SISTEMÁTICA/METAANÁLISIS"},
            "Título": {"en": "Title", "es": "Título"},
            "Resumo": {"en": "Abstract", "es": "Resumen"},
            "AUTOMÁTICA": {"en": "AUTOMATIC", "es": "AUTOMÁTICA"},
            "MANUAL": {"en": "MANUAL", "es": "MANUAL"},
        }

        def ss_value(value):
            value = clean(value)
            item = second_value_map.get(value, {})
            return item.get(self.language_code, value)

        win = tk.Toplevel(self)
        suffix = ss(
            " — CONFIRMADA (edite e confirme novamente para alterar)",
            " — CONFIRMED (edit and confirm again to change)",
            " — CONFIRMADA (edite y confirme nuevamente para cambiar)",
        ) if self.second_screen_applied else ""
        win.title(ss(
            "2ª triagem — Marcação de revisões sistemáticas e meta-análises",
            "Second screening — Systematic review and meta-analysis classification",
            "Segundo cribado — Clasificación de revisiones sistemáticas y metaanálisis",
        ) + suffix)
        win.geometry("1420x880")
        win.minsize(1050, 700)

        info = ttk.LabelFrame(
            win,
            text=ss(
                "Prévia classificatória da 2ª triagem",
                "Second-screening classification preview",
                "Vista previa de clasificación del segundo cribado",
            ),
            padding=10,
        )
        info.pack(fill="x", padx=10, pady=10)
        ttk.Label(
            info,
            text=ss(
                "Esta etapa NÃO exclui nem apaga artigos. Ela apenas identifica e marca revisões sistemáticas, "
                "meta-análises e revisões de revisões. Termo explícito no TÍTULO → classificação sugerida como "
                "CONFIRMADA. Termo encontrado somente no RESUMO → fica como REVISAR. Você pode corrigir o tipo "
                "manualmente. Ao confirmar, as marcações serão levadas para as planilhas de cada base e para o banco final.",
                "This step does NOT exclude or delete articles. It only identifies and labels systematic reviews, "
                "meta-analyses, and overviews of reviews. An explicit term in the TITLE → classification is suggested as "
                "CONFIRMED. A term found only in the ABSTRACT → remains flagged for REVIEW. You can correct the type "
                "manually. After confirmation, the labels are carried into the database-specific spreadsheets and the final dataset.",
                "Esta etapa NO excluye ni elimina artículos. Solo identifica y marca revisiones sistemáticas, "
                "metaanálisis y revisiones de revisiones. Un término explícito en el TÍTULO → se sugiere como "
                "CONFIRMADO. Un término encontrado solo en el RESUMEN → queda marcado para REVISAR. Puede corregir el tipo "
                "manualmente. Al confirmar, las marcas se incorporan a las hojas de cada base y al conjunto final.",
            ),
            wraplength=1320,
            justify="left",
        ).pack(anchor="w")
        ttk.Label(
            info,
            text=ss(
                "✓ Nenhum registro será retirado do banco final nesta etapa.",
                "✓ No record will be removed from the final dataset at this stage.",
                "✓ Ningún registro se eliminará del conjunto final en esta etapa.",
            ),
            font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", pady=(6, 0))

        top = ttk.Frame(win, padding=(10, 0, 10, 4))
        top.pack(fill="x")
        ttk.Label(top, text=ss("Buscar:", "Search:", "Buscar:")).pack(side="left")
        search_var = tk.StringVar()
        search_entry = ttk.Entry(top, textvariable=search_var, width=48)
        search_entry.pack(side="left", padx=6)
        counter_var = tk.StringVar()
        ttk.Label(top, textvariable=counter_var, font=("Segoe UI", 9, "bold")).pack(side="right")

        base_counter_var = tk.StringVar()
        ttk.Label(
            win,
            textvariable=base_counter_var,
            font=("Segoe UI", 9),
            padding=(10, 0, 10, 5),
        ).pack(fill="x")

        frame = ttk.Frame(win, padding=(10, 0, 10, 4))
        frame.pack(fill="both", expand=True)
        cols = ("status", "tipo", "fonte", "base", "ano", "titulo")
        tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode="extended", height=16)
        settings = {
            "status": (ss("Status", "Status", "Estado"), 130),
            "tipo": (ss("Classificação", "Classification", "Clasificación"), 285),
            "fonte": (ss("Detectado em", "Detected in", "Detectado en"), 95),
            "base": (ss("Base", "Database", "Base"), 130),
            "ano": (ss("Ano", "Year", "Año"), 65),
            "titulo": (ss("Título", "Title", "Título"), 650),
        }
        for col in cols:
            tree.heading(col, text=settings[col][0])
            tree.column(col, width=settings[col][1], minwidth=60)
        ybar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=ybar.set)
        tree.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")

        buttons = ttk.Frame(win, padding=(10, 4))
        buttons.pack(fill="x")
        detail = tk.Text(win, height=10, wrap="word", font=("Segoe UI", 10), padx=8, pady=8)
        detail.pack(fill="both", expand=False, padx=10, pady=(0, 4))
        detail.configure(state="disabled")

        final_buttons = ttk.Frame(win, padding=(10, 4, 10, 10))
        final_buttons.pack(fill="x")

        classification_internal_choices = [
            "REVISÃO SISTEMÁTICA",
            "META-ANÁLISE",
            "REVISÃO SISTEMÁTICA + META-ANÁLISE",
            "REVISÃO DE REVISÕES",
        ]
        classification_display_choices = [ss_value(v) for v in classification_internal_choices]
        display_to_internal = dict(zip(classification_display_choices, classification_internal_choices))
        classification_var = tk.StringVar(value=classification_display_choices[0])

        def selected_uids():
            out = []
            for item in tree.selection():
                tags = tree.item(item, "tags")
                if tags:
                    out.append(tags[0])
            return out

        def normalized_decisions():
            if working_df.empty:
                return pd.Series(dtype=str)
            return pd.Series(
                [
                    normalize_second_screen_decision(d, t)
                    for d, t in zip(working_df["decisao_2triagem"], working_df["tipo_detectado"])
                ],
                index=working_df.index,
            )

        def update_counter():
            detected = len(working_df)
            decisions = normalized_decisions()
            confirmed = int((decisions == "CONFIRMADO").sum()) if detected else 0
            pending = int((decisions == "REVISAR").sum()) if detected else 0
            dismissed = int((decisions == "NÃO É REVISÃO/META").sum()) if detected else 0
            counter_var.set(ss(
                f"Detectados: {detected}   |   Confirmados: {confirmed}   |   Revisar: {pending}   |   Falsos positivos: {dismissed}   |   Banco final: {len(self.final_df)}",
                f"Detected: {detected}   |   Confirmed: {confirmed}   |   Review: {pending}   |   False positives: {dismissed}   |   Final dataset: {len(self.final_df)}",
                f"Detectados: {detected}   |   Confirmados: {confirmed}   |   Revisar: {pending}   |   Falsos positivos: {dismissed}   |   Conjunto final: {len(self.final_df)}",
            ))

            if detected:
                base_counts = working_df.groupby("base").size().sort_values(ascending=False).to_dict()
                base_text = "   |   ".join(
                    f"{self.display_database_name(db)}: {int(n)}" for db, n in base_counts.items()
                )
                base_counter_var.set(ss(
                    f"Detectados por base: {base_text}",
                    f"Detected by database: {base_text}",
                    f"Detectados por base: {base_text}",
                ))
            else:
                base_counter_var.set(ss(
                    "Detectados por base: nenhum caso encontrado",
                    "Detected by database: no cases found",
                    "Detectados por base: no se encontraron casos",
                ))

        def show_detail(_event=None):
            uids = selected_uids()
            if not uids:
                return
            row = working_df[working_df["uid"] == uids[0]]
            if row.empty:
                return
            r = row.iloc[0]
            decision = normalize_second_screen_decision(
                r.get("decisao_2triagem", ""), r.get("tipo_detectado", "")
            )
            abstract_text = clean(r.get("resumo", "")) or ss(
                "[Resumo não disponível neste arquivo]",
                "[Abstract not available in this file]",
                "[Resumen no disponible en este archivo]",
            )
            text = (
                f"{ss('STATUS', 'STATUS', 'ESTADO')}: {ss_value(decision)} ({ss_value(r.get('origem_decisao_2triagem', ''))})\n"
                f"{ss('CLASSIFICAÇÃO', 'CLASSIFICATION', 'CLASIFICACIÓN')}: {ss_value(r.get('tipo_detectado', ''))}\n"
                f"{ss('DETECTADO EM', 'DETECTED IN', 'DETECTADO EN')}: {ss_value(r.get('fonte_deteccao', ''))}\n"
                f"{ss('EVIDÊNCIA', 'EVIDENCE', 'EVIDENCIA')}: {clean(r.get('evidencia_2triagem', ''))}\n\n"
                f"{ss('TÍTULO', 'TITLE', 'TÍTULO')}\n{clean(r.get('titulo', ''))}\n\n"
                f"{ss('RESUMO', 'ABSTRACT', 'RESUMEN')}\n{abstract_text}"
            )
            detail.configure(state="normal")
            detail.delete("1.0", "end")
            detail.insert("1.0", text)
            detail.configure(state="disabled")

        def refresh(*_args):
            query = normalize_title(search_var.get())
            tree.delete(*tree.get_children())
            for _, row in working_df.iterrows():
                haystack = normalize_title(" ".join([
                    clean(row.get("titulo", "")), clean(row.get("autores", "")),
                    clean(row.get("tipo_detectado", "")), clean(row.get("evidencia_2triagem", "")),
                    clean(row.get("base", "")),
                ]))
                if query and query not in haystack:
                    continue
                decision = normalize_second_screen_decision(
                    row.get("decisao_2triagem", ""), row.get("tipo_detectado", "")
                )
                if decision == "CONFIRMADO":
                    color_tag = "confirmed"
                elif decision == "REVISAR":
                    color_tag = "pending"
                else:
                    color_tag = "dismissed"
                tree.insert(
                    "", "end",
                    values=(
                        ss_value(decision),
                        ss_value(row.get("tipo_detectado", "")),
                        ss_value(row.get("fonte_deteccao", "")),
                        self.display_database_name(row.get("base", "")),
                        row.get("ano", ""),
                        row.get("titulo", ""),
                    ),
                    tags=(row.get("uid", ""), color_tag),
                )
            tree.tag_configure("confirmed", background="#D9EAD3")
            tree.tag_configure("pending", background="#FFF2A8")
            tree.tag_configure("dismissed", background="#E7E6E6")
            update_counter()

        def apply_classification():
            uids = selected_uids()
            if not uids:
                self._showinfo("2ª triagem", "Selecione um ou mais registros.", parent=win)
                return
            selected_type = display_to_internal.get(clean(classification_var.get()), "")
            if selected_type not in classification_internal_choices:
                self._showwarning("2ª triagem", "Escolha uma classificação válida.", parent=win)
                return
            mask = working_df["uid"].isin(uids)
            working_df.loc[mask, "tipo_detectado"] = selected_type
            working_df.loc[mask, "decisao_2triagem"] = "CONFIRMADO"
            working_df.loc[mask, "origem_decisao_2triagem"] = "MANUAL"
            refresh()
            show_detail()

        def confirm_detected_type():
            uids = selected_uids()
            if not uids:
                self._showinfo("2ª triagem", "Selecione um ou mais registros.", parent=win)
                return
            for uid in uids:
                mask = working_df["uid"] == uid
                if not mask.any():
                    continue
                current_type = clean(working_df.loc[mask, "tipo_detectado"].iloc[0])
                if current_type.startswith("POSSÍVEL"):
                    continue
                working_df.loc[mask, "decisao_2triagem"] = "CONFIRMADO"
                working_df.loc[mask, "origem_decisao_2triagem"] = "MANUAL"
            refresh()
            show_detail()

        def mark_not_review():
            uids = selected_uids()
            if not uids:
                self._showinfo("2ª triagem", "Selecione um ou mais registros.", parent=win)
                return
            mask = working_df["uid"].isin(uids)
            working_df.loc[mask, "decisao_2triagem"] = "NÃO É REVISÃO/META"
            working_df.loc[mask, "origem_decisao_2triagem"] = "MANUAL"
            refresh()
            show_detail()

        def restore_auto():
            uids = selected_uids()
            if not uids:
                self._showinfo("2ª triagem", "Selecione um ou mais registros.", parent=win)
                return
            automatic_df = self._build_review_filter_dataframe(None)
            automatic_map = {
                clean(r.get("uid", "")): r
                for _, r in automatic_df.iterrows()
            }
            for uid in uids:
                r = automatic_map.get(uid)
                if r is None:
                    continue
                mask = working_df["uid"] == uid
                working_df.loc[mask, "tipo_detectado"] = clean(r.get("tipo_detectado", ""))
                working_df.loc[mask, "fonte_deteccao"] = clean(r.get("fonte_deteccao", ""))
                working_df.loc[mask, "evidencia_2triagem"] = clean(r.get("evidencia_2triagem", ""))
                working_df.loc[mask, "decisao_2triagem"] = clean(r.get("decisao_2triagem", ""))
                working_df.loc[mask, "origem_decisao_2triagem"] = "AUTOMÁTICA"
            refresh()
            show_detail()

        def confirm_second_screen():
            decisions = normalized_decisions()
            confirmed = int((decisions == "CONFIRMADO").sum()) if not working_df.empty else 0
            pending = int((decisions == "REVISAR").sum()) if not working_df.empty else 0
            dismissed = int((decisions == "NÃO É REVISÃO/META").sum()) if not working_df.empty else 0
            ok = self._askyesno(
                ss("Confirmar marcações da 2ª triagem", "Confirm second-screening marks", "Confirmar marcas del segundo cribado"),
                ss(
                    "Salvar definitivamente estas marcações?\n\n"
                    f"Registros detectados: {len(working_df)}\n"
                    f"Classificados como revisão/meta: {confirmed}\n"
                    f"Ainda para revisar: {pending}\n"
                    f"Marcados como falso positivo: {dismissed}\n"
                    f"Registros no banco final: {len(self.final_df)}\n\n"
                    "Nenhum artigo será excluído. As classificações serão incluídas nas exportações.",
                    "Save these labels permanently?\n\n"
                    f"Detected records: {len(working_df)}\n"
                    f"Classified as review/meta-analysis: {confirmed}\n"
                    f"Still to review: {pending}\n"
                    f"Marked as false positives: {dismissed}\n"
                    f"Records in the final dataset: {len(self.final_df)}\n\n"
                    "No article will be excluded. Classifications will be included in the exports.",
                    "¿Guardar definitivamente estas marcas?\n\n"
                    f"Registros detectados: {len(working_df)}\n"
                    f"Clasificados como revisión/metaanálisis: {confirmed}\n"
                    f"Aún por revisar: {pending}\n"
                    f"Marcados como falsos positivos: {dismissed}\n"
                    f"Registros en el conjunto final: {len(self.final_df)}\n\n"
                    "No se excluirá ningún artículo. Las clasificaciones se incluirán en las exportaciones.",
                ),
                parent=win,
            )
            if not ok:
                return

            self.review_filter_df = working_df.copy().reset_index(drop=True)
            # Normaliza possíveis estados antigos antes de salvar.
            for idx, row in self.review_filter_df.iterrows():
                self.review_filter_df.at[idx, "decisao_2triagem"] = normalize_second_screen_decision(
                    row.get("decisao_2triagem", ""), row.get("tipo_detectado", "")
                )
            self.second_screen_applied = True
            stats_now = self.second_screen_stats()
            self.log_audit_event(
                "2ª TRIAGEM CONFIRMADA",
                f"Confirmados={stats_now['confirmed']}; pendentes={stats_now['pending']}; descartados={stats_now['dismissed']}; nenhum registro removido",
            )
            self.update_summary()
            self.status_var.set(
                ss(
                    f"2ª triagem confirmada: {confirmed} revisão(ões)/meta-análise(s) marcada(s); nenhum registro removido.",
                    f"Second screening confirmed: {confirmed} review(s)/meta-analysis(es) marked; no records removed.",
                    f"Segundo cribado confirmado: {confirmed} revisión(es)/metaanálisis marcado(s); ningún registro eliminado.",
                )
            )
            self._showinfo(
                ss("2ª triagem confirmada", "Second screening confirmed", "Segundo cribado confirmado"),
                ss(
                    f"Marcações salvas com sucesso.\n\nConfirmados: {confirmed}\nPara revisar: {pending}\n"
                    f"Falsos positivos: {dismissed}\nRegistros mantidos no banco: {len(self.final_df)}",
                    f"Labels saved successfully.\n\nConfirmed: {confirmed}\nTo review: {pending}\n"
                    f"False positives: {dismissed}\nRecords retained in the dataset: {len(self.final_df)}",
                    f"Marcas guardadas correctamente.\n\nConfirmados: {confirmed}\nPara revisar: {pending}\n"
                    f"Falsos positivos: {dismissed}\nRegistros mantenidos en el conjunto: {len(self.final_df)}",
                ),
                parent=win,
            )
            win.destroy()

        def close_without_confirming():
            if self._askyesno(
                ss("Fechar sem confirmar", "Close without confirming", "Cerrar sin confirmar"),
                ss(
                    "Fechar esta janela sem salvar as alterações desta conferência?\n\n"
                    "As marcações confirmadas anteriormente, se houver, permanecerão inalteradas.",
                    "Close this window without saving the changes from this review?\n\n"
                    "Previously confirmed labels, if any, will remain unchanged.",
                    "¿Cerrar esta ventana sin guardar los cambios de esta revisión?\n\n"
                    "Las marcas confirmadas anteriormente, si las hubiera, permanecerán sin cambios.",
                ),
                parent=win,
            ):
                win.destroy()

        ttk.Button(
            buttons,
            text=ss("CONFIRMAR TIPO DETECTADO", "CONFIRM DETECTED TYPE", "CONFIRMAR TIPO DETECTADO"),
            command=confirm_detected_type,
        ).pack(side="left", padx=3)
        ttk.Label(
            buttons,
            text=ss("Classificar como:", "Classify as:", "Clasificar como:"),
        ).pack(side="left", padx=(12, 3))
        ttk.Combobox(
            buttons,
            textvariable=classification_var,
            values=classification_display_choices,
            state="readonly",
            width=38,
        ).pack(side="left", padx=3)
        ttk.Button(
            buttons,
            text=ss("APLICAR CLASSIFICAÇÃO", "APPLY CLASSIFICATION", "APLICAR CLASIFICACIÓN"),
            command=apply_classification,
        ).pack(side="left", padx=3)
        ttk.Button(
            buttons,
            text=ss("NÃO É REVISÃO/META", "NOT A REVIEW/META-ANALYSIS", "NO ES REVISIÓN/METAANÁLISIS"),
            command=mark_not_review,
        ).pack(side="left", padx=3)
        ttk.Button(
            buttons,
            text=ss("RESTAURAR AUTOMÁTICO", "RESTORE AUTOMATIC", "RESTAURAR AUTOMÁTICO"),
            command=restore_auto,
        ).pack(side="left", padx=3)
        ttk.Label(
            buttons,
            text=ss(
                "Verde = confirmado | Amarelo = revisar | Cinza = falso positivo",
                "Green = confirmed | Yellow = review | Gray = false positive",
                "Verde = confirmado | Amarillo = revisar | Gris = falso positivo",
            ),
        ).pack(side="right")

        ttk.Button(
            final_buttons,
            text=ss(
                "CONFIRMAR MARCAÇÕES DA 2ª TRIAGEM",
                "CONFIRM SECOND-SCREENING MARKS",
                "CONFIRMAR MARCAS DEL SEGUNDO CRIBADO",
            ),
            command=confirm_second_screen,
        ).pack(side="right", padx=(8, 0), ipadx=18, ipady=5)
        ttk.Button(
            final_buttons,
            text=ss("FECHAR SEM CONFIRMAR", "CLOSE WITHOUT CONFIRMING", "CERRAR SIN CONFIRMAR"),
            command=close_without_confirming,
        ).pack(side="right")
        ttk.Label(
            final_buttons,
            text=ss(
                "As marcações serão exportadas; nenhum artigo será excluído.",
                "Labels will be exported; no article will be excluded.",
                "Las marcas se exportarán; ningún artículo será excluido.",
            ),
            font=("Segoe UI", 9, "bold"),
        ).pack(side="left")

        tree.bind("<<TreeviewSelect>>", show_detail)
        search_var.trace_add("write", refresh)
        win.protocol("WM_DELETE_WINDOW", close_without_confirming)
        refresh()
        if tree.get_children():
            tree.selection_set(tree.get_children()[0])
            show_detail()
        else:
            self._showinfo(
                ss("2ª triagem", "Second screening", "Segundo cribado"),
                ss(
                    "Nenhuma revisão sistemática ou meta-análise explícita foi detectada. "
                    "Você pode confirmar esta etapa para registrar zero marcações.",
                    "No explicit systematic review or meta-analysis was detected. "
                    "You can confirm this step to record zero labels.",
                    "No se detectó ninguna revisión sistemática o metaanálisis explícito. "
                    "Puede confirmar esta etapa para registrar cero marcas.",
                ),
                parent=win,
            )
        search_entry.focus_set()

    def _legacy_open_triage_window_v09(self):
        if self.final_df.empty:
            self._showwarning("Triagem", "Processe e deduplique as bases antes de iniciar a triagem.")
            return

        self.sync_triage_df()
        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("Triagem de título e resumo")
        win.geometry("1280x820")
        win.minsize(950, 650)

        top = ttk.Frame(win, padding=8)
        top.pack(fill="x")
        ttk.Label(top, text="Buscar:").pack(side="left")
        search_var = tk.StringVar()
        search_entry = ttk.Entry(top, textvariable=search_var, width=48)
        search_entry.pack(side="left", padx=(6, 12))
        counter_var = tk.StringVar()
        ttk.Label(top, textvariable=counter_var, font=("Segoe UI", 9, "bold")).pack(side="right")

        table_frame = ttk.Frame(win, padding=(8, 0, 8, 4))
        table_frame.pack(fill="both", expand=True)
        tree = ttk.Treeview(
            table_frame,
            columns=("status", "base", "ano", "titulo"),
            show="headings",
            selectmode="browse",
            height=15,
        )
        for col, title, width in [
            ("status", "Triagem", 125),
            ("base", "Base", 135),
            ("ano", "Ano", 70),
            ("titulo", "Título", 780),
        ]:
            tree.heading(col, text=title)
            tree.column(col, width=width, minwidth=60)
        ybar = ttk.Scrollbar(table_frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=ybar.set)
        tree.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")

        action = ttk.Frame(win, padding=(8, 4))
        action.pack(fill="x")

        detail = tk.Text(win, height=13, wrap="word", font=("Segoe UI", 10), padx=8, pady=8)
        detail.pack(fill="both", expand=False, padx=8, pady=(0, 8))
        detail.configure(state="disabled")

        def current_uid():
            selection = tree.selection()
            if not selection:
                return ""
            tags = tree.item(selection[0], "tags")
            return tags[0] if tags else ""

        def update_counter():
            series = self.triage_df["triagem_status"] if not self.triage_df.empty else pd.Series(dtype=str)
            counts = series.value_counts().to_dict()
            counter_var.set(
                f"Incluídos: {counts.get('INCLUIR', 0)}   |   Excluídos: {counts.get('EXCLUIR', 0)}   |   "
                f"Talvez: {counts.get('TALVEZ', 0)}   |   Pendentes: {counts.get('NÃO AVALIADO', 0)}"
            )

        def show_detail(_event=None):
            uid = current_uid()
            if not uid:
                return
            subset = self.triage_df[self.triage_df["uid"] == uid]
            if subset.empty:
                return
            row = subset.iloc[0]
            content = (
                f"STATUS: {clean(row.get('triagem_status', ''))}\n"
                f"MOTIVO DE EXCLUSÃO: {clean(row.get('motivo_exclusao', '')) or '-'}\n\n"
                f"TÍTULO\n{clean(row.get('titulo', ''))}\n\n"
                f"AUTORES\n{clean(row.get('autores', ''))}\n\n"
                f"RESUMO\n{clean(row.get('resumo', '')) or '[Resumo não disponível neste formato de exportação]'}\n\n"
                f"DOI: {clean(row.get('doi', '')) or '-'}   |   ID: {clean(row.get('id_origem', '')) or '-'}"
            )
            detail.configure(state="normal")
            detail.delete("1.0", "end")
            detail.insert("1.0", content)
            detail.configure(state="disabled")

        def refresh(*_args):
            query = normalize_title(search_var.get())
            selected_uid = current_uid()
            tree.delete(*tree.get_children())
            for _, row in self.triage_df.iterrows():
                haystack = normalize_title(
                    " ".join([
                        clean(row.get("titulo", "")), clean(row.get("autores", "")),
                        clean(row.get("doi", "")), clean(row.get("id_origem", "")),
                    ])
                )
                if query and query not in haystack:
                    continue
                status = clean(row.get("triagem_status", "")) or "NÃO AVALIADO"
                tag_color = {
                    "INCLUIR": "include",
                    "EXCLUIR": "exclude",
                    "TALVEZ": "maybe",
                }.get(status, "pending")
                item = tree.insert(
                    "", "end",
                    values=(status, row.get("base", ""), row.get("ano", ""), row.get("titulo", "")),
                    tags=(row.get("uid", ""), tag_color),
                )
                if selected_uid and row.get("uid", "") == selected_uid:
                    tree.selection_set(item)
            tree.tag_configure("include", background="#D9EAD3")
            tree.tag_configure("exclude", background="#F4CCCC")
            tree.tag_configure("maybe", background="#FFF2A8")
            tree.tag_configure("pending", background="#FFFFFF")
            update_counter()
            show_detail()

        def set_triage(status):
            uid = current_uid()
            if not uid:
                self._showinfo("Triagem", "Selecione um artigo.", parent=win)
                return
            reason = ""
            if status == "EXCLUIR":
                current = self.triage_df.loc[self.triage_df["uid"] == uid, "motivo_exclusao"].iloc[0]
                reason = self._askstring(
                    "Motivo da exclusão",
                    "Informe o motivo da exclusão (opcional):",
                    initialvalue=current,
                    parent=win,
                )
                if reason is None:
                    return
            mask = self.triage_df["uid"] == uid
            self.triage_df.loc[mask, "triagem_status"] = status
            self.triage_df.loc[mask, "motivo_exclusao"] = clean(reason) if status == "EXCLUIR" else ""
            refresh()

        ttk.Button(action, text="✓ INCLUIR", command=lambda: set_triage("INCLUIR")).pack(side="left", padx=3)
        ttk.Button(action, text="✗ EXCLUIR", command=lambda: set_triage("EXCLUIR")).pack(side="left", padx=3)
        ttk.Button(action, text="? TALVEZ", command=lambda: set_triage("TALVEZ")).pack(side="left", padx=3)
        ttk.Button(action, text="LIMPAR DECISÃO", command=lambda: set_triage("NÃO AVALIADO")).pack(side="left", padx=3)
        ttk.Label(action, text="As decisões ficam salvas no projeto e entram no PRISMA automático.").pack(side="right")

        tree.bind("<<TreeviewSelect>>", show_detail)
        search_var.trace_add("write", refresh)
        refresh()
        if tree.get_children():
            tree.selection_set(tree.get_children()[0])
            show_detail()
        search_entry.focus_set()

    def open_search_window(self):
        if self.all_df.empty:
            self._showwarning("Buscar artigo", "Nenhuma base foi processada.")
            return

        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("Buscar artigo nas bases")
        win.geometry("1250x620")

        top = ttk.Frame(win, padding=8)
        top.pack(fill="x")
        ttk.Label(top, text="Título, DOI, PMID/ID ou autor:").pack(side="left")
        query_var = tk.StringVar()
        entry = ttk.Entry(top, textvariable=query_var, width=58)
        entry.pack(side="left", padx=8)

        frame = ttk.Frame(win, padding=(8, 0, 8, 8))
        frame.pack(fill="both", expand=True)
        cols = ("base", "situacao", "ano", "titulo", "doi", "id")
        tree = ttk.Treeview(frame, columns=cols, show="headings")
        settings = {
            "base": ("Base", 130), "situacao": ("Situação", 260), "ano": ("Ano", 65),
            "titulo": ("Título", 540), "doi": ("DOI", 200), "id": ("ID", 120),
        }
        for c in cols:
            tree.heading(c, text=settings[c][0])
            tree.column(c, width=settings[c][1], minwidth=60)
        ybar = ttk.Scrollbar(frame, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=ybar.set)
        tree.pack(side="left", fill="both", expand=True)
        ybar.pack(side="right", fill="y")

        final_uids = set(self.final_df["uid"].tolist()) if not self.final_df.empty else set()
        second_maps = self._second_screen_annotation_maps()

        def search(*_args):
            query = normalize_title(query_var.get())
            tree.delete(*tree.get_children())
            if not query:
                return
            for _, row in self.all_df.iterrows():
                haystack = normalize_title(" ".join([
                    clean(row.get("titulo", "")), clean(row.get("autores", "")),
                    clean(row.get("doi", "")), clean(row.get("id_origem", "")),
                ]))
                if query not in haystack:
                    continue
                uid = clean(row.get("uid", ""))
                if uid not in final_uids:
                    situation = "Removido como duplicata"
                elif getattr(self, "second_screen_applied", False):
                    ann = self._second_screen_annotation_for_record(row, second_maps)
                    if ann["status"] == "CONFIRMADO":
                        situation = f"Banco final • 2ª triagem: {ann['classificacao']}"
                    elif ann["status"] == "REVISAR":
                        situation = f"Banco final • 2ª triagem: REVISAR ({ann['classificacao']})"
                    elif ann["status"] == "NÃO É REVISÃO/META":
                        situation = "Banco final • 2ª triagem: falso positivo"
                    else:
                        situation = "Banco final após deduplicação"
                else:
                    situation = "Banco após deduplicação"
                tree.insert("", "end", values=(
                    row.get("base", ""), situation, row.get("ano", ""), row.get("titulo", ""),
                    row.get("doi", ""), row.get("id_origem", ""),
                ))

        ttk.Button(top, text="BUSCAR", command=search).pack(side="left")
        entry.bind("<Return>", search)
        entry.focus_set()

    def prisma_counts(self):
        identified = int(len(self.all_df))
        deduplicated = int(len(self.final_df))
        duplicates = max(0, identified - deduplicated)
        base_counts = self.all_df.groupby("base").size().to_dict() if not self.all_df.empty else {}

        stats = self.second_screen_stats() if getattr(self, "second_screen_applied", False) else {
            "detected": 0, "confirmed": 0, "pending": 0, "dismissed": 0,
            "type_counts": {}, "base_counts": {},
        }

        # A 2ª triagem é apenas classificatória; portanto o total disponível
        # para triagem continua sendo exatamente o banco deduplicado.
        screening = deduplicated

        title_abstract_excluded = max(0, int(self.prisma_extra.get("records_excluded_title_abstract", 0) or 0))
        title_abstract_excluded = min(title_abstract_excluded, screening)
        reports_sought = max(0, screening - title_abstract_excluded)
        not_retrieved = max(0, int(self.prisma_extra.get("reports_not_retrieved", 0) or 0))
        not_retrieved = min(not_retrieved, reports_sought)
        reports_assessed = max(0, reports_sought - not_retrieved)
        full_excluded = max(0, int(self.prisma_extra.get("full_text_excluded", 0) or 0))
        full_excluded = min(full_excluded, reports_assessed)
        studies_included = max(0, int(self.prisma_extra.get("studies_included", 0) or 0))

        later_steps_entered = any([
            title_abstract_excluded, not_retrieved, full_excluded, studies_included,
            clean(self.prisma_extra.get("full_text_reasons", "")),
        ])

        return {
            "identified": identified,
            "duplicates": duplicates,
            "deduplicated": deduplicated,
            "second_screen_applied": bool(getattr(self, "second_screen_applied", False)),
            "second_detected": stats["detected"],
            "second_confirmed": stats["confirmed"],
            "second_pending": stats["pending"],
            "second_dismissed": stats["dismissed"],
            "second_type_counts": stats["type_counts"],
            "second_base_counts": stats["base_counts"],
            "second_excluded": 0,
            "second_possible": stats["pending"],
            "screening": screening,
            "base_counts": base_counts,
            "title_abstract_excluded": title_abstract_excluded,
            "reports_sought": reports_sought,
            "reports_not_retrieved": not_retrieved,
            "reports_assessed": reports_assessed,
            "full_text_excluded": full_excluded,
            "studies_included": studies_included,
            "later_steps_entered": later_steps_entered,
            "full_text_reasons": clean(self.prisma_extra.get("full_text_reasons", "")),
        }

    def export_prisma_png(self, output_path):
        counts = self.prisma_counts()
        try:
            from PIL import Image, ImageDraw, ImageFont
        except Exception as exc:
            raise RuntimeError(
                "Para gerar o PRISMA em PNG, instale a biblioteca Pillow (pip install pillow)."
            ) from exc

        full = counts["later_steps_entered"]
        width, height = 1500, 1940 if full else 1380
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)

        def font(size, bold=False):
            candidates = [
                "arialbd.ttf" if bold else "arial.ttf",
                "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",
            ]
            for name in candidates:
                try:
                    return ImageFont.truetype(name, size)
                except Exception:
                    pass
            return ImageFont.load_default()

        title_font = font(34, True)
        box_font = font(23, False)
        small_font = font(18, False)

        draw.text((width // 2, 35), "Fluxograma PRISMA 2020 — revisão sistemática", anchor="ma", font=title_font, fill="black")
        draw.text((width // 2, 82), "Deduplicação e marcações da 2ª triagem contabilizadas automaticamente", anchor="ma", font=small_font, fill="black")

        main_x1, main_x2 = 150, 900
        side_x1, side_x2 = 1000, 1420
        box_h = 150

        def wrap(text, max_chars=58):
            words = text.split()
            lines, current = [], []
            for word in words:
                candidate = " ".join(current + [word])
                if len(candidate) > max_chars and current:
                    lines.append(" ".join(current))
                    current = [word]
                else:
                    current.append(word)
            if current:
                lines.append(" ".join(current))
            return "\n".join(lines)

        def box(x1, y1, x2, y2, text, header=None):
            draw.rounded_rectangle((x1, y1, x2, y2), radius=14, outline="black", width=3, fill="white")
            if header:
                draw.text(((x1 + x2)//2, y1 + 16), header, anchor="ma", font=font(20, True), fill="black")
                body_y = y1 + 52
            else:
                body_y = y1 + 24
            max_chars = max(22, int((x2 - x1) / 13))
            draw.multiline_text(((x1 + x2)//2, body_y), wrap(text, max_chars=max_chars), anchor="ma", align="center", font=box_font, fill="black", spacing=7)

        def arrow(x, y1, y2):
            draw.line((x, y1, x, y2 - 12), fill="black", width=4)
            draw.polygon([(x, y2), (x - 10, y2 - 16), (x + 10, y2 - 16)], fill="black")

        base_lines = "\n".join(f"{db}: n={n}" for db, n in sorted(counts["base_counts"].items(), key=lambda x: -x[1]))
        y = 145
        box(main_x1, y, main_x2, y + 205, f"Registros identificados nas bases (n={counts['identified']})\n{base_lines}", "IDENTIFICAÇÃO")
        box(side_x1, y + 22, side_x2, y + 180, f"Duplicatas removidas (n={counts['duplicates']})")
        arrow((main_x1 + main_x2)//2, y + 205, y + 250)

        y = 395
        box(main_x1, y, main_x2, y + box_h, f"Registros após deduplicação (n={counts['deduplicated']})", "PRÉ-TRIAGEM")
        if counts["second_screen_applied"]:
            side_text = (
                f"Revisões sistemáticas/meta-análises identificadas e mantidas (n={counts['second_confirmed']})"
            )
            if counts["second_pending"]:
                side_text += f"\nPendentes de confirmação: n={counts['second_pending']}"
        else:
            side_text = "2ª triagem classificatória de revisões/meta-análises ainda não executada"
        box(side_x1, y - 8, side_x2, y + 175, side_text)
        arrow((main_x1 + main_x2)//2, y + box_h, y + box_h + 45)

        y = 595
        box(main_x1, y, main_x2, y + 175, f"Registros disponíveis para triagem de título/resumo (n={counts['screening']})", "TRIAGEM")

        if full:
            box(side_x1, y, side_x2, y + 175, f"Registros excluídos na triagem externa de título/resumo (n={counts['title_abstract_excluded']})")
            arrow((main_x1 + main_x2)//2, y + 175, y + 220)
            y = 815
            box(main_x1, y, main_x2, y + box_h, f"Relatórios buscados para recuperação (n={counts['reports_sought']})", "ELEGIBILIDADE")
            box(side_x1, y, side_x2, y + box_h, f"Relatórios não recuperados (n={counts['reports_not_retrieved']})")
            arrow((main_x1 + main_x2)//2, y + box_h, y + box_h + 45)

            y = 1015
            box(main_x1, y, main_x2, y + box_h, f"Relatórios avaliados em texto completo (n={counts['reports_assessed']})")
            reason = counts["full_text_reasons"] or "Motivos não informados"
            box(side_x1, y - 10, side_x2, y + 185, f"Relatórios excluídos após texto completo (n={counts['full_text_excluded']})\n{reason}")
            arrow((main_x1 + main_x2)//2, y + box_h, y + box_h + 45)

            y = 1220
            box(main_x1, y, main_x2, y + box_h, f"Estudos incluídos na revisão (n={counts['studies_included']})", "INCLUSÃO")
        else:
            draw.text(
                (width // 2, 850),
                "Fluxo parcial: a triagem de título/resumo será realizada fora do programa.",
                anchor="ma", font=font(22, True), fill="black"
            )

        draw.text((width // 2, height - 55), f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}", anchor="ma", font=small_font, fill="black")
        image.save(output_path, format="PNG")
        return Path(output_path)

    def open_prisma_window(self):
        if self.all_df.empty:
            self._showwarning("PRISMA", "Nenhuma base foi processada.")
            return

        win = tk.Toplevel(self)
        self.after_idle(lambda w=win: self._localize_widget_tree(w))
        win.title("PRISMA automático")
        win.geometry("800x700")

        summary = ttk.LabelFrame(win, text="Contagens automáticas", padding=12)
        summary.pack(fill="x", padx=10, pady=10)
        summary_text = tk.StringVar()

        def update_summary_label():
            c = self.prisma_counts()
            second = (
                f"Revisões/meta-análises marcadas na 2ª triagem: {c['second_confirmed']} "
                f"(pendentes: {c['second_pending']}; nenhum registro removido)"
                if c["second_screen_applied"]
                else "2ª triagem classificatória de revisões/meta-análises: ainda não executada"
            )
            summary_text.set(self.translate_literal(
                f"Identificados: {c['identified']}\n"
                f"Duplicatas removidas: {c['duplicates']}\n"
                f"Após deduplicação: {c['deduplicated']}\n"
                f"{second}\n"
                f"Disponíveis para triagem de título/resumo: {c['screening']}"
            ))
        ttk.Label(summary, textvariable=summary_text, justify="left", font=("Segoe UI", 10)).pack(anchor="w")
        update_summary_label()

        later = ttk.LabelFrame(win, text="Etapas posteriores — números agregados opcionais", padding=12)
        later.pack(fill="x", padx=10, pady=(0, 10))
        vars_int = {
            "records_excluded_title_abstract": tk.StringVar(value=str(self.prisma_extra.get("records_excluded_title_abstract", 0))),
            "reports_not_retrieved": tk.StringVar(value=str(self.prisma_extra.get("reports_not_retrieved", 0))),
            "full_text_excluded": tk.StringVar(value=str(self.prisma_extra.get("full_text_excluded", 0))),
            "studies_included": tk.StringVar(value=str(self.prisma_extra.get("studies_included", 0))),
        }
        labels = [
            ("records_excluded_title_abstract", "Excluídos na triagem externa de título/resumo:"),
            ("reports_not_retrieved", "Relatórios não recuperados:"),
            ("full_text_excluded", "Relatórios excluídos após texto completo:"),
            ("studies_included", "Estudos incluídos na revisão:"),
        ]
        for r, (key, label) in enumerate(labels):
            ttk.Label(later, text=label).grid(row=r, column=0, sticky="w", pady=4)
            ttk.Entry(later, textvariable=vars_int[key], width=10).grid(row=r, column=1, sticky="w", padx=8)
        ttk.Label(later, text="Motivos das exclusões em texto completo:").grid(row=4, column=0, sticky="nw", pady=4)
        reasons = tk.Text(later, height=5, width=58, wrap="word")
        reasons.grid(row=4, column=1, sticky="ew", padx=8, pady=4)
        reasons.insert("1.0", clean(self.prisma_extra.get("full_text_reasons", "")))
        later.columnconfigure(1, weight=1)

        def save_values(show=False):
            try:
                for key, var in vars_int.items():
                    value = int(var.get().strip() or "0")
                    if value < 0:
                        raise ValueError
                    self.prisma_extra[key] = value
                self.prisma_extra["full_text_reasons"] = clean(reasons.get("1.0", "end"))
            except Exception:
                self._showerror("PRISMA", "Use apenas números inteiros iguais ou maiores que zero.", parent=win)
                return False
            update_summary_label()
            if show:
                self._showinfo("PRISMA", "Contagens agregadas atualizadas.", parent=win)
            return True

        buttons = ttk.Frame(win, padding=10)
        buttons.pack(fill="x")
        ttk.Button(buttons, text="ATUALIZAR DADOS", command=lambda: save_values(True)).pack(side="left", padx=4)

        def generate():
            if not save_values(False):
                return
            path = self._asksaveasfilename(
                parent=win,
                title="Salvar fluxograma PRISMA",
                defaultextension=".png",
                initialfile="PRISMA_FLUXOGRAMA.png",
                filetypes=[("Imagem PNG", "*.png")],
            )
            if not path:
                return
            try:
                self.export_prisma_png(path)
            except Exception as exc:
                self._showerror("PRISMA", str(exc), parent=win)
                return
            self._showinfo("PRISMA", f"Fluxograma gerado em:\n{path}", parent=win)

        def export_excel():
            if not save_values(False):
                return
            path = self._asksaveasfilename(
                parent=win,
                title="Salvar dados do PRISMA",
                defaultextension=".xlsx",
                initialfile="PRISMA_DADOS.xlsx",
                filetypes=[("Excel", "*.xlsx")],
            )
            if not path:
                return
            try:
                self.export_prisma_data_excel(path)
                self.log_audit_event("PRISMA DADOS EXPORTADOS", path)
            except Exception as exc:
                self._showerror("PRISMA", str(exc), parent=win)
                return
            self._showinfo("PRISMA", f"Dados salvos em:\n{path}", parent=win)

        ttk.Button(buttons, text="GERAR IMAGEM PNG", command=generate).pack(side="left", padx=4)
        ttk.Button(buttons, text="EXPORTAR DADOS EXCEL", command=export_excel).pack(side="left", padx=4)
        ttk.Label(
            win,
            text=(
                "A triagem individual de título/resumo foi removida. O programa contabiliza automaticamente "
                "a deduplicação e a marcação classificatória de revisões/meta-análises. Nenhum registro é removido na 2ª triagem. Os campos acima servem apenas para "
                "inserir números agregados das etapas feitas fora do programa."
            ),
            wraplength=740,
            justify="left",
        ).pack(fill="x", padx=14, pady=8)

    # --------------------------------------------------------
    # Exportação
    # --------------------------------------------------------

    def export_clean_base_excel(self, database, df, folder):
        output = Path(folder) / f"{safe_filename(database)}_LIMPA.xlsx"

        # Mantém todos os registros da base e acrescenta a marcação da 2ª triagem.
        annotated = self._annotate_export_dataframe(df)
        clean_df = annotated[[
            "autores", "titulo", "revista", "ano", "doi", "id_origem",
            "status_2triagem_export", "classificacao_2triagem_export",
            "fonte_2triagem_export", "evidencia_2triagem_export", "origem_2triagem_export",
        ]].copy()

        clean_df.columns = [
            "Autores", "Título do estudo", "Revista", "Ano", "DOI", "ID original",
            "2ª triagem - status", "2ª triagem - classificação", "Detectado em",
            "Evidência da 2ª triagem", "Origem da classificação",
        ]

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            clean_df.to_excel(writer, sheet_name="REFERENCIAS", index=False)
            ws = writer.book["REFERENCIAS"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            widths = {
                "A": 48, "B": 70, "C": 38, "D": 10, "E": 32, "F": 16,
                "G": 22, "H": 38, "I": 18, "J": 38, "K": 20,
            }
            for col, width in widths.items():
                ws.column_dimensions[col].width = width

            from openpyxl.styles import PatternFill, Font
            blue_fill = PatternFill(fill_type="solid", fgColor="0B4F8A")
            green_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
            yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2A8")
            gray_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")

            for cell in ws[1]:
                cell.fill = blue_fill
                cell.font = Font(bold=True, color="FFFFFF")

            # Coluna G = status classificatório. A cor serve só como sinalização.
            for row_number in range(2, ws.max_row + 1):
                status = clean(ws.cell(row=row_number, column=7).value)
                fill = None
                if status == "CONFIRMADO":
                    fill = green_fill
                elif status == "REVISAR":
                    fill = yellow_fill
                elif status == "NÃO É REVISÃO/META":
                    fill = gray_fill
                if fill is not None:
                    for cell in ws[row_number]:
                        cell.fill = fill

        return output

    def export_comparison_excel(self, folder):
        output = (
            Path(folder)
            / "COMPARACAO_DUPLICATAS.xlsx"
        )

        comparison = self.comparison_df.copy().rename(columns={
            "comparacao_id": "ID comparação",
            "status": "Status",
            "criterio": "Critério",
            "similaridade": "Similaridade (%)",
            "evidencia": "Evidência compacta",
            "uid_referencia": "UID já existente",
            "base_referencia": "Já existente - Base",
            "id_referencia": "Já existente - ID",
            "autores_referencia": "Já existente - Autores",
            "ano_referencia": "Já existente - Ano",
            "titulo_referencia": "Já existente - Título",
            "doi_referencia": "Já existente - DOI",
            "uid_comparado": "UID novo registro",
            "base_comparada": "Novo registro - Base",
            "id_comparado": "Novo registro - ID",
            "autores_comparado": "Novo registro - Autores",
            "ano_comparado": "Novo registro - Ano",
            "titulo_comparado": "Novo registro - Título",
            "doi_comparado": "Novo registro - DOI",
            "decisao_manual": "Decisão manual",
        })

        with pd.ExcelWriter(
            output,
            engine="openpyxl"
        ) as writer:
            comparison.to_excel(
                writer,
                sheet_name="COMPARACOES",
                index=False
            )

            self.base_report_df().to_excel(
                writer,
                sheet_name="RESUMO_BASES",
                index=False
            )

            ws = writer.book["COMPARACOES"]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            widths = {
                "A": 15, "B": 28, "C": 20, "D": 16, "E": 22,
                "F": 20, "G": 20, "H": 16, "I": 42, "J": 12,
                "K": 70, "L": 30, "M": 20, "N": 20, "O": 16,
                "P": 42, "Q": 12, "R": 70, "S": 30, "T": 20,
            }

            for col, width in widths.items():
                ws.column_dimensions[col].width = width

            for cell in ws[1]:
                cell.font = cell.font.copy(
                    bold=True
                )

            # Mesmas cores da tela para auditoria manual.
            from openpyxl.styles import PatternFill

            fill_yellow = PatternFill(
                fill_type="solid",
                fgColor="FFF2A8"
            )
            fill_green = PatternFill(
                fill_type="solid",
                fgColor="D9EAD3"
            )
            fill_gray = PatternFill(
                fill_type="solid",
                fgColor="E7E6E6"
            )

            # Coluna B = Status
            for row_number in range(2, ws.max_row + 1):
                status = ws.cell(
                    row=row_number,
                    column=2
                ).value or ""

                if status == "REVISAR":
                    fill = fill_yellow
                elif status in {
                    "DUPLICATA CONFIRMADA",
                    "DUPLICATA CONFIRMADA MANUAL"
                }:
                    fill = fill_green
                elif status == "NÃO DUPLICATA":
                    fill = fill_gray
                else:
                    fill = None

                if fill is not None:
                    for cell in ws[row_number]:
                        cell.fill = fill

            summary_ws = writer.book[
                "RESUMO_BASES"
            ]
            summary_ws.freeze_panes = "A2"
            summary_ws.auto_filter.ref = (
                summary_ws.dimensions
            )

            for cell in summary_ws[1]:
                cell.font = cell.font.copy(
                    bold=True
                )

            for col, width in {
                "A": 28,
                "B": 12,
                "C": 24,
                "D": 14,
                "E": 12,
            }.items():
                summary_ws.column_dimensions[
                    col
                ].width = width

        return output

    def export_final_excel(self, folder):
        output = Path(folder) / "BASE_FINAL_PARA_TRIAGEM.xlsx"

        final_df = self._annotate_export_dataframe(self.post_second_screen_df())
        final = final_df[[
            "base", "autores", "titulo", "resumo", "revista", "ano", "doi", "id_origem",
            "status_2triagem_export", "classificacao_2triagem_export",
            "fonte_2triagem_export", "evidencia_2triagem_export", "origem_2triagem_export",
        ]].copy()
        final.columns = [
            "Base de origem", "Autores", "Título do estudo", "Resumo", "Revista",
            "Ano", "DOI", "ID original", "2ª triagem - status",
            "2ª triagem - classificação", "Detectado em", "Evidência da 2ª triagem",
            "Origem da classificação",
        ]

        marked = final[final["2ª triagem - status"].astype(str).str.strip() != ""].copy()

        # Legenda simples para que o usuário identifique imediatamente as cores.
        legend_rows = []
        for db in [
            "PubMed", "Web of Science", "Scopus", "Embase",
            "LILACS/BVS", "Cochrane Library", "SciELO", "Outra",
        ]:
            legend_rows.append({"Tipo": "Base de origem", "Identificação": db})
        legend_rows.extend([
            {"Tipo": "Revisão/meta", "Identificação": "CONFIRMADO"},
            {"Tipo": "Revisão/meta", "Identificação": "REVISAR"},
            {"Tipo": "Revisão/meta", "Identificação": "NÃO É REVISÃO/META"},
        ])
        legend_df = pd.DataFrame(legend_rows)

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            final.to_excel(writer, sheet_name="BASE_FINAL", index=False)
            if not marked.empty:
                marked.to_excel(writer, sheet_name="2_TRIAGEM_REVISOES_META", index=False)
            legend_df.to_excel(writer, sheet_name="LEGENDA_CORES", index=False)

            from openpyxl.styles import PatternFill, Font, Alignment
            blue_fill = PatternFill(fill_type="solid", fgColor="0B4F8A")
            green_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
            yellow_fill = PatternFill(fill_type="solid", fgColor="FFF2A8")
            gray_fill = PatternFill(fill_type="solid", fgColor="E7E6E6")

            database_fills = {
                db: PatternFill(fill_type="solid", fgColor=color)
                for db, color in DATABASE_EXPORT_COLORS.items()
            }

            # Formata as duas planilhas com registros.
            data_sheet_names = ["BASE_FINAL"]
            if "2_TRIAGEM_REVISOES_META" in writer.book.sheetnames:
                data_sheet_names.append("2_TRIAGEM_REVISOES_META")

            for sheet_name in data_sheet_names:
                sheet = writer.book[sheet_name]
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions

                for cell in sheet[1]:
                    cell.fill = blue_fill
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(vertical="center")

                for col, width in {
                    "A":22,"B":48,"C":75,"D":90,"E":38,"F":10,"G":32,"H":16,
                    "I":22,"J":38,"K":18,"L":38,"M":20,
                }.items():
                    sheet.column_dimensions[col].width = width

                for row_number in range(2, sheet.max_row + 1):
                    database = clean(sheet.cell(row=row_number, column=1).value)
                    base_fill = database_fills.get(database, database_fills.get("Outra"))

                    # Por padrão, toda a linha recebe a cor da base de origem.
                    # Quando houver marcação de revisão/meta, o bloco I:M será
                    # sobrescrito pela cor classificatória para continuar evidente.
                    if base_fill is not None:
                        for col_number in range(1, 14):
                            sheet.cell(row=row_number, column=col_number).fill = base_fill

                    # Coluna A ganha negrito para facilitar localizar as bases.
                    sheet.cell(row=row_number, column=1).font = Font(bold=True)

                    # Coluna I = status da 2ª triagem: preserva a sinalização já existente.
                    status = clean(sheet.cell(row=row_number, column=9).value)
                    status_fill = None
                    if status == "CONFIRMADO":
                        status_fill = green_fill
                    elif status == "REVISAR":
                        status_fill = yellow_fill
                    elif status == "NÃO É REVISÃO/META":
                        status_fill = gray_fill
                    classification = clean(sheet.cell(row=row_number, column=10).value)
                    class_fill = status_fill
                    if class_fill is None and classification:
                        if classification.upper().startswith("POSSÍVEL"):
                            class_fill = yellow_fill
                        else:
                            class_fill = green_fill

                    # I:M é todo o bloco classificatório. Assim a base continua
                    # visualmente identificável e revisão/meta permanece destacada.
                    if class_fill is not None:
                        for col_number in range(9, 14):
                            sheet.cell(row=row_number, column=col_number).fill = class_fill
                        sheet.cell(row=row_number, column=9).font = Font(bold=True)
                        sheet.cell(row=row_number, column=10).font = Font(bold=True)

            # Formata a legenda.
            legend_ws = writer.book["LEGENDA_CORES"]
            legend_ws.freeze_panes = "A2"
            legend_ws.column_dimensions["A"].width = 20
            legend_ws.column_dimensions["B"].width = 38
            for cell in legend_ws[1]:
                cell.fill = blue_fill
                cell.font = Font(bold=True, color="FFFFFF")

            for row_number in range(2, legend_ws.max_row + 1):
                kind = clean(legend_ws.cell(row=row_number, column=1).value)
                label = clean(legend_ws.cell(row=row_number, column=2).value)
                fill = None
                if kind == "Base de origem":
                    fill = database_fills.get(label, database_fills.get("Outra"))
                elif label == "CONFIRMADO":
                    fill = green_fill
                elif label == "REVISAR":
                    fill = yellow_fill
                elif label == "NÃO É REVISÃO/META":
                    fill = gray_fill
                if fill is not None:
                    for cell in legend_ws[row_number]:
                        cell.fill = fill

        return output

    def export_report_txt(self, folder):
        output = (
            Path(folder)
            / "RELATORIO_COMPARACAO.txt"
        )

        mini = self.base_report_df()

        lines = [
            f"{APP_TITLE} v{VERSION}",
            "=" * 70,
            f"Desenvolvido por: {AUTHOR_NAME}",
            f"E-mail: {AUTHOR_EMAIL}",
            COPYRIGHT_NOTICE,
            f"Licença: {LICENSE_NAME}",
            f"Repositório: {PROJECT_URL}",
            "",
            "RESUMO POR BASE",
            "",
            (
                f"{'Base':<24}"
                f"{'Bruto':>10}"
                f"{'Duplicatas':>14}"
                f"{'Pendentes':>12}"
                f"{'Final':>10}"
            ),
            "-" * 70,
        ]

        for _, row in mini.iterrows():
            lines.append(
                f"{str(row['Base']):<24}"
                f"{int(row['Bruto']):>10}"
                f"{int(row['Duplicatas removidas']):>14}"
                f"{int(row['Pendentes']):>12}"
                f"{int(row['Final']):>10}"
            )

        total_bruto = int(
            mini["Bruto"].sum()
        ) if not mini.empty else 0

        total_duplicates = int(
            mini["Duplicatas removidas"].sum()
        ) if not mini.empty else 0

        total_pending = int(
            mini["Pendentes"].sum()
        ) if not mini.empty else 0

        total_final = int(
            mini["Final"].sum()
        ) if not mini.empty else 0

        lines.extend([
            "-" * 70,
            (
                f"{'TOTAL':<24}"
                f"{total_bruto:>10}"
                f"{total_duplicates:>14}"
                f"{total_pending:>12}"
                f"{total_final:>10}"
            ),
            "",
            f"Base principal: {self.reference_var.get()}",
            f"Ordem de comparação: {' → '.join(self.comparison_order)}",
            "",
            "INTERPRETAÇÃO DO RELATÓRIO",
            (
                "- 'Duplicatas' = registros removidos daquela base porque "
                "já existiam em uma base processada anteriormente."
            ),
            (
                "- A base principal mantém sua própria cópia; por isso normalmente "
                "não terá registros removidos nesta etapa."
            ),
            (
                "- 'Pendentes' = títulos semelhantes que ainda aguardam "
                "confirmação manual."
            ),
            (
                "- 'Final' = quantos registros daquela base permanecerão "
                "no banco para triagem no estado atual."
            ),
            "",
            "CRITÉRIOS",
            "- DOI idêntico = duplicata confirmada automaticamente.",
            "- Título normalizado idêntico = duplicata confirmada automaticamente.",
            (
                f"- Título com similaridade >= {self.threshold_var.get()}% "
                "= enviado para conferência manual."
            ),
            "- Uma base nunca é comparada com ela mesma.",
            (
                "- Registros de bases anteriores permanecem no histórico "
                "de comparação, mesmo quando já foram removidos do banco final."
            ),
            "",
            "2ª TRIAGEM — REVISÕES SISTEMÁTICAS / META-ANÁLISES",
            (
                f"- Status: {'executada' if self.second_screen_applied else 'não executada'}."
            ),
            (
                f"- Confirmados como revisão/meta: {self.second_screen_stats()['confirmed'] if self.second_screen_applied else 0}."
            ),
            (
                f"- Pendentes para revisar: {self.second_screen_stats()['pending'] if self.second_screen_applied else 0}."
            ),
            "- Removidos nesta etapa: 0 (a 2ª triagem é somente classificatória).",
            (
                f"- Registros disponíveis para triagem de título/resumo: {len(self.post_second_screen_df())}."
            ),
            "",
            "ARQUIVOS EXPORTADOS",
            "- Excel limpo separado de cada base, com colunas de marcação da 2ª triagem.",
            "- COMPARACAO_DUPLICATAS.xlsx (inclui aba RESUMO_BASES).",
            "- BASE_FINAL_PARA_TRIAGEM.xlsx (todos os registros deduplicados são mantidos; inclui classificação de revisão/meta e aba específica da 2ª triagem).",
            "- PRISMA_FLUXOGRAMA.png (deduplicação automática + marcação classificatória informativa da 2ª triagem).",
            "- PRISMA_DADOS.xlsx (contagens auditáveis do fluxo e resumo por base).",
            "- AUDITORIA_REPRODUTIBILIDADE.xlsx/.txt (arquivos, parâmetros, decisões e histórico).",
            "- Menu Bibliometria: mesclagem conservadora de WoS/Embase/Scopus e exportação separada de PubMed.",
        ])

        output.write_text(
            "\n".join(lines),
            encoding="utf-8"
        )

        return output

    def export_all(self):
        if self.all_df.empty:
            self._showwarning(
                "Sem dados",
                "Processe os arquivos antes de exportar."
            )
            return

        folder = self._askdirectory(
            title="Escolha a pasta para salvar os resultados"
        )

        if not folder:
            return

        created = []

        try:
            for database, df in self.base_data.items():
                created.append(
                    self.export_clean_base_excel(
                        database,
                        df,
                        folder
                    )
                )

            created.append(
                self.export_comparison_excel(folder)
            )

            created.append(
                self.export_final_excel(folder)
            )

            created.append(
                self.export_report_txt(folder)
            )

            created.append(
                self.export_prisma_png(Path(folder) / "PRISMA_FLUXOGRAMA.png")
            )
            created.append(
                self.export_prisma_data_excel(Path(folder) / "PRISMA_DADOS.xlsx")
            )
            created.extend(self.export_audit_report(folder))
            self.log_audit_event("EXPORTAÇÃO COMPLETA", str(folder))

        except Exception as exc:
            self._showerror(
                "Erro de exportação",
                str(exc)
            )
            return

        self._showinfo(
            "Exportação concluída",
            f"{len(created)} arquivo(s) foram criados em:\n{folder}"
        )

        self.status_var.set(
            self.translate_literal(f"Exportado para: {folder}")
        )


if __name__ == "__main__":
    SimpleReviewApp().mainloop()
